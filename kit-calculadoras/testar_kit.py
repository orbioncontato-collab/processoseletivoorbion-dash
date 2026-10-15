# -*- coding: utf-8 -*-
"""Auditoria automática: cria variações do arquivo, recalcula no LibreOffice e confere resultados.

Uso: python testar_kit.py arquivo.xlsx registro.json pasta_temp [alturas_saida.json]
Gera também as alturas de linha necessárias para os textos dinâmicos.
"""
import json
import math
import os
import shutil
import sys

from openpyxl import load_workbook

from lo_recalc import main as recalc

BASE, REG, TMP = sys.argv[1], sys.argv[2], sys.argv[3]
ALT_OUT = sys.argv[4] if len(sys.argv) > 4 else None
reg = json.load(open(REG, encoding="utf-8"))
os.makedirs(TMP, exist_ok=True)

ABAS = {1: "02 — Vazão e Hidráulica", 2: "03 — Carga Poluidora", 3: "04 — Eficiência de Tratamento",
        4: "05 — Balanço de Massa", 5: "06 — Diluição de Soluções", 6: "07 — Dosagem Química"}


def T(n, t):
    return reg["calcs"][str(n)]["tools"][t]


# Cada caso: nome, lista de alterações (calc, ferramenta, chave|'scen'|'sel', valor), verificações
# verificação: (calc, ferramenta, alvo, esperado) alvo: 'main', 'sec0'.., 'status~', 'interp~'
M = "Meu cálculo"
CASOS = [
    ("exemplo01", [], [
        (1, 0, "main", 180), (1, 0, "sec1", 50), (1, 0, "sec0", 0.05), (1, 0, "sec2", 4320),
        (1, 1, "sec0", 0.0706858), (1, 1, "main", 0.1060288), (1, 2, "main", 1.5915494), (1, 3, "main", 10),
        (1, 3, "sec0", 600), (2, 0, "main", 72), (2, 0, "sec0", 72000), (2, 1, "main", 60), (2, 2, "main", 400),
        (3, 0, "main", 80), (3, 1, "main", 80), (3, 1, "sec0", 100), (3, 1, "sec1", 20),
        (4, 0, "main", 140), (4, 0, "sec0", 200), (4, 1, "main", 7530), (4, 1, "sec3", "✔ Fecha"),
        (4, 2, "main", 80), (5, 0, "main", 100), (5, 0, "sec0", 400),
        (6, 0, "sec0", 2), (6, 0, "main", 4), (6, 0, "sec1", 3.3333333),
        (5, 0, "interp~", "Utilize 100 mL da solução estoque e complete com solvente até atingir o volume final de 500 mL."),
    ]),
    ("exemplo02", [(n, t, "scen", "Exemplo 02") for n, k in [(1, 4), (2, 3), (3, 2), (4, 3), (5, 1), (6, 1)]
                   for t in range(k)], [
        (1, 0, "main", 36), (1, 0, "sec1", 10), (1, 0, "sec2", 864), (1, 1, "main", 0.01413717), (1, 1, "sec1", 14.137167),
        (1, 2, "main", 2.5464791), (1, 3, "main", 6), (1, 3, "sec0", 360), (1, 3, "sec1", 0.25),
        (2, 0, "main", 108), (2, 1, "main", 125), (2, 2, "main", 750),
        (3, 0, "main", 70), (3, 1, "main", 90.5), (3, 1, "sec1", 28.5),
        (4, 0, "main", 4.7619048), (4, 0, "sec0", 9072), (4, 1, "main", 13025), (4, 2, "main", -5),
        (4, 2, "interp~", "Acúmulo negativo"),
        (5, 0, "main", 20), (6, 0, "sec0", 4.32), (6, 0, "main", 36), (6, 0, "sec1", 30),
    ]),
    ("meu_vazio", [(n, t, "scen", M) for n, k in [(1, 4), (2, 3), (3, 2), (4, 3), (5, 1), (6, 1)]
                   for t in range(k)], [
        (1, 0, "main", "—"), (1, 0, "status~", "Informe o volume para continuar."),
        (1, 1, "status~", "Informe o diâmetro"), (2, 0, "status~", "Informe a vazão"),
        (3, 0, "status~", "Informe a concentração de entrada"), (4, 1, "status~", "Informe a vazão de entrada"),
        (5, 0, "status~", "Informe a concentração do estoque"), (6, 0, "status~", "Informe a vazão"),
        (6, 0, "main", "—"), (4, 2, "main", "—"),
    ]),
    ("unidades", [
        (1, 0, "scen", M), (1, 0, "V", 360000), (1, 0, "uV", "L"), (1, 0, "t", 120), (1, 0, "ut", "min"),
        (1, 1, "scen", M), (1, 1, "D", 300), (1, 1, "uD", "mm"), (1, 1, "v", 1.5),
        (1, 2, "scen", M), (1, 2, "Q", 1), (1, 2, "uQ", "m³/s"), (1, 2, "D", 0.2), (1, 2, "uD", "m"),
        (1, 3, "scen", M), (1, 3, "V", 500000), (1, 3, "uV", "L"), (1, 3, "Q", 1200), (1, 3, "uQ", "m³/dia"),
        (2, 0, "scen", M), (2, 0, "P", "DQO"), (2, 0, "Q", 10), (2, 0, "uQ", "L/s"), (2, 0, "C", 500),
        (6, 0, "scen", M), (6, 0, "Q", 20), (6, 0, "uQ", "m³/h"), (6, 0, "D", 10), (6, 0, "pct", 100),
    ], [
        (1, 0, "sec1", 50), (1, 1, "main", 0.1060288), (1, 2, "main", 31.830989),
        (1, 2, "interp~", "acima de 3 m/s"), (1, 3, "main", 10), (2, 0, "main", 432), (2, 0, "status~", "✔"),
        (6, 0, "sec0", 4.8), (6, 0, "main", 4.8), (6, 0, "sec1", "sem densidade"),
    ]),
    ("zero_negativo_texto", [
        (1, 0, "scen", M), (1, 0, "V", 0), (1, 0, "t", 2),
        (1, 1, "scen", M), (1, 1, "D", -0.3), (1, 1, "v", 1.5),
        (1, 3, "scen", M), (1, 3, "V", "abc"), (1, 3, "Q", 50),
        (3, 0, "scen", M), (3, 0, "Ce", 0), (3, 0, "Cs", 10),
        (2, 2, "scen", M), (2, 2, "L", 10), (2, 2, "C", 0),
        (6, 0, "scen", M), (6, 0, "Q", 500), (6, 0, "D", 4), (6, 0, "pct", 0),
    ], [
        (1, 0, "status~", "O volume deve ser maior que zero."), (1, 0, "main", "—"),
        (1, 1, "status~", "O diâmetro deve ser maior que zero."),
        (1, 3, "status~", "precisa ser um número"),
        (3, 0, "status~", "A concentração de entrada deve ser maior que zero."),
        (2, 2, "status~", "A concentração deve ser maior que zero."),
        (6, 0, "status~", "Confira a porcentagem do componente ativo"), (6, 0, "main", "—"),
    ]),
    ("eficiencia_sinais", [
        (3, 0, "scen", M), (3, 0, "P", "SST"), (3, 0, "Ce", 50), (3, 0, "Cs", 100),
        (3, 1, "scen", M), (3, 1, "P", "DBO"), (3, 1, "Qe", 100), (3, 1, "Ce", 100), (3, 1, "Qs", 100), (3, 1, "Cs", 100),
        (4, 2, "scen", M), (4, 2, "E", 10), (4, 2, "Sd", 10), (4, 2, "G", 0), (4, 2, "Cn", 0),
    ], [
        (3, 0, "main", -100), (3, 0, "interp~", "Houve aumento de SST"), (3, 1, "main", 0),
        (3, 1, "interp~", "Não houve redução"), (4, 2, "main", 0), (4, 2, "interp~", "Acúmulo nulo"),
    ]),
    ("decantador", [
        (4, 0, "scen", M), (4, 0, "Q1", 100), (4, 0, "C1", 80), (4, 0, "Q2", 100), (4, 0, "C2", 80),
        (4, 1, "scen", M), (4, 1, "Qe", 500), (4, 1, "Ce", 180), (4, 1, "Qa", 480), (4, 1, "Ca", 30), (4, 1, "Ql", 10),
    ], [
        (4, 0, "main", 80), (4, 0, "interp~", "mesma concentração"),
        (4, 1, "main", 7560), (4, 1, "sec3", "⚠ Não fecha"), (4, 1, "interp~", "Diferença de 10 m³/dia"),
    ]),
    ("decantador_negativo", [
        (4, 1, "scen", M), (4, 1, "Qe", 100), (4, 1, "Ce", 10), (4, 1, "Qa", 90), (4, 1, "Ca", 50), (4, 1, "Ql", 10),
    ], [(4, 1, "main", -350), (4, 1, "interp~", "O resultado é negativo")]),
    ("diluicao_alertas", [
        (5, 0, "scen", M), (5, 0, "modo", "V₁ (volume de estoque)"), (5, 0, "C1", 100), (5, 0, "C2", 200),
        (5, 0, "V2", 500),
    ], [(5, 0, "status~", "A concentração desejada é maior que a solução estoque."), (5, 0, "main", "—")]),
    ("diluicao_unidade", [
        (5, 0, "scen", M), (5, 0, "modo", "V₁ (volume de estoque)"), (5, 0, "C1", 1), (5, 0, "uC1", "mol/L"),
        (5, 0, "C2", 200), (5, 0, "uC2", "mg/L"), (5, 0, "V2", 500),
    ], [(5, 0, "status~", "precisam ser do mesmo tipo")]),
    ("diluicao_v2", [
        (5, 0, "scen", M), (5, 0, "modo", "V₂ (volume final)"), (5, 0, "C1", 1), (5, 0, "uC1", "g/L"),
        (5, 0, "V1", 10), (5, 0, "uV1", "mL"), (5, 0, "C2", 100), (5, 0, "uC2", "mg/L"), (5, 0, "uV2", "mL"),
    ], [(5, 0, "main", 100), (5, 0, "sec0", 90), (5, 0, "sec1", 10)]),
    ("diluicao_c1", [
        (5, 0, "scen", M), (5, 0, "modo", "C₁ (concentração do estoque)"), (5, 0, "uC1", "mg/L"), (5, 0, "V1", 25),
        (5, 0, "uV1", "mL"), (5, 0, "C2", 50), (5, 0, "uC2", "mg/L"), (5, 0, "V2", 1), (5, 0, "uV2", "L"),
    ], [(5, 0, "main", 2000), (5, 0, "sec0", 0.975)]),
    ("diluicao_v1maior", [
        (5, 0, "scen", M), (5, 0, "modo", "C₂ (concentração final)"), (5, 0, "C1", 1000), (5, 0, "V1", 600),
        (5, 0, "V2", 500),
    ], [(5, 0, "status~", "não pode ser maior que o volume final")]),
    ("dosagem_alertas", [
        (6, 0, "scen", M), (6, 0, "Q", 500), (6, 0, "D", 4), (6, 0, "pct", 0.5), (6, 0, "dens", 3.5),
    ], [(6, 0, "main", 400), (6, 0, "interp~", "digite 50 (e não 0,5)"), (6, 0, "interp~", "Confira a densidade")]),
    ("seletor", [("sel", 1, "Ferramenta 3: Velocidade do escoamento")], [("summary", 1, None, "1,5915 m/s")]),
]


def aplicar(wb, alts):
    for a in alts:
        if a[0] == "sel":
            _, n, val = a
            wb[ABAS[n]][reg["calcs"][str(n)]["tool_sel"]] = val
            continue
        n, t, k, val = a
        ws = wb[ABAS[n]]
        cell = T(n, t)["scen"] if k == "scen" else T(n, t)["my"][k]
        ws[cell] = val


arquivos = []
for nome, alts, _ in CASOS:
    f = os.path.join(TMP, f"caso_{nome}.xlsx")
    if alts:
        wb = load_workbook(BASE)
        aplicar(wb, alts)
        wb.save(f)
    else:
        shutil.copy(BASE, f)
    arquivos.append(f)
recalc(arquivos)

falhas, total, erros_formula = [], 0, 0
textos = {}
for (nome, alts, checks), f in zip(CASOS, arquivos):
    wb = load_workbook(f, data_only=True)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value[:1] == "#" and c.value.rstrip("!?/0").upper() in (
                        "#VALUE", "#DIV", "#REF", "#NAME", "#N/A", "#NUM", "#NULL"):
                    erros_formula += 1
                    falhas.append(f"[{nome}] erro bruto {ws.title}!{c.coordinate} = {c.value}")
    for d in reg["dyn"]:
        v = wb[d[0]][d[2]].value
        if isinstance(v, str):
            key = f"{d[0]}|{d[1]}"
            textos.setdefault(key, []).append((v, d[3], d[4]))
    for chk in checks:
        total += 1
        if chk[0] == "summary":
            _, n, _, esp = chk
            got = wb[ABAS[n]][reg["calcs"][str(n)]["tool_summary"]].value
            ok = got == esp
        else:
            n, t, alvo, esp = chk
            info = T(n, t)
            if alvo == "main":
                cell = info["main"]
            elif alvo.startswith("sec"):
                cell = info["sec"][int(alvo[3:])]
            elif alvo == "status~":
                cell = info["status"]
            else:
                cell = info["interp"]
            got = wb[ABAS[n]][cell].value
            if alvo.endswith("~"):
                ok = isinstance(got, str) and esp in got
            elif isinstance(esp, (int, float)):
                ok = isinstance(got, (int, float)) and math.isclose(got, esp, rel_tol=1e-6, abs_tol=1e-9)
            else:
                ok = got == esp
        if not ok:
            falhas.append(f"[{nome}] {chk} -> obtido {got!r}")

print(f"Verificações: {total - len([x for x in falhas if 'erro bruto' not in x])}/{total} corretas")
print(f"Erros brutos de fórmula em todas as variações: {erros_formula}")
for x in falhas:
    print("FALHA", x)

if ALT_OUT:
    alt = {}
    for key, lst in textos.items():
        best = 0
        for v, larg, size in lst:
            cpl = larg * 1.08 * (10.0 / size)
            linhas = sum(max(1, math.ceil(len(p) / cpl)) for p in v.split("\n"))
            best = max(best, linhas)
        alt[key] = max(20, best * 13.0 * size / 10.0 + 8)
    json.dump(alt, open(ALT_OUT, "w"), indent=0)
    print("alturas:", len(alt))
sys.exit(1 if falhas else 0)
