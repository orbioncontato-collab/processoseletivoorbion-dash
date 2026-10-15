# -*- coding: utf-8 -*-
"""Núcleo do gerador: estilos, utilitários de célula e mini linguagem de fórmulas."""
import math
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.hyperlink import Hyperlink

FONTE = "Arial"

# Paleta
NAVY = "1F3864"
NAVY2 = "2F5597"
AZUL = "2E75B6"
AZUL_L = "DDEBF7"
AZUL_XL = "EEF4FB"
VERDE_D = "2E7D32"
VERDE_T = "1E5631"
VERDE_L = "E2F0D9"
VERDE_OK = "C6EFCE"
AMARELO = "FFF2CC"
AMARELO_B = "BF9000"
AMARELO_S = "FFE699"
CINZA_L = "F2F2F2"
CINZA_XL = "F7F7F7"
CINZA_B = "D9D9D9"
CINZA_T = "595959"
CINZA_M = "8C8C8C"
VERM_L = "FCE4E4"
VERM_T = "C00000"
ARDOSIA = "44546A"
BRANCO = "FFFFFF"
TXT = "262626"

FMT_POS = '#,##0.00###'
FMT_SIG = '#,##0.00;-#,##0.00;0.00'
FMT_PCT = '0.0'

CENARIOS = ["Exemplo 01", "Exemplo 02", "Meu cálculo"]


def tint(hexcor, f):
    r, g, b = (int(hexcor[i:i + 2], 16) for i in (0, 2, 4))
    r, g, b = (round(x + (255 - x) * f) for x in (r, g, b))
    return f"{r:02X}{g:02X}{b:02X}"


def F(size=10, bold=False, color=TXT, italic=False, underline=None):
    return Font(name=FONTE, size=size, bold=bold, color=color, italic=italic, underline=underline)


def FL(c):
    return PatternFill("solid", fgColor=c)


def BD(c=CINZA_B, s="thin"):
    sd = Side(style=s, color=c)
    return Border(left=sd, right=sd, top=sd, bottom=sd)


def AL(h="left", v="center", wrap=True, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


AL_L = AL("left", indent=1)
AL_LT = AL("left", "top", indent=1)
AL_C = AL("center")
AL_R = AL("right", indent=1)

DESBLOQ = Protection(locked=False)


def q(nome):
    return "'" + nome.replace("'", "''") + "'"


def box(ws, rng, v=None, font=None, fill=None, al=None, border=None, fmt=None, unlock=False):
    """Escreve (e mescla, se for intervalo) aplicando estilo em todas as células."""
    if ":" in rng:
        ws.merge_cells(rng)
        cells = [c for row in ws[rng] for c in row]
    else:
        cells = [ws[rng]]
    for c in cells:
        if font is not None:
            c.font = font
        if fill is not None:
            c.fill = fill
        if al is not None:
            c.alignment = al
        if border is not None:
            c.border = border
    first = cells[0]
    if v is not None:
        first.value = v
    if fmt is not None:
        first.number_format = fmt
    if unlock:
        first.protection = DESBLOQ
    return first


def link(cell, sheet, ref="A1"):
    # display é obrigatório: sem ele o Google Planilhas mostra o endereço (#gid=...) no lugar do texto
    texto = cell.value if isinstance(cell.value, str) else str(cell.value or "")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{q(sheet)}!{ref}", display=texto,
                               tooltip=f"Ir para {sheet}")


def altura_texto(texto, largura, size=10, bold=False, minimo=18, extra=6):
    """Estima a altura de linha (pt) para texto com quebra automática."""
    cpl = largura * (1.0 if bold else 1.08) * (10.0 / size)
    linhas = 0
    for par in str(texto).split("\n"):
        linhas += max(1, math.ceil(len(par) / max(cpl, 1)))
    return max(minimo, linhas * 13.0 * (size / 10.0) + extra)


# ---------------------------------------------------------------- fórmulas

def N(k):
    """Número como texto. O construtor cria células auxiliares [k__t] com FIXED (respeita o idioma)."""
    return f"[{k}__t]"


def fmt_texto(xr):
    """Formata um número já arredondado (xr) com o mínimo de casas decimais necessárias."""
    ks = "7"
    for d in range(6, -1, -1):
        ks = f"IF(ROUND({xr},{d})={xr},{d},{ks})"
    return f'=IF(ISNUMBER({xr}),FIXED({xr},{ks}),"?")'


def lit(texto):
    """Literal de texto em fórmula, quebrado em partes < 255 caracteres."""
    if texto == "":
        return '""'
    partes = [texto[i:i + 200] for i in range(0, len(texto), 200)]
    return "&".join('"' + p.replace('"', '""') + '"' for p in partes)


def S(texto):
    """Mini linguagem: {k} = número da chave k; {@k} = valor bruto; {/} = quebra de linha."""
    out = []
    for p in re.split(r"(\{[^{}]*\})", texto):
        if not p:
            continue
        if p.startswith("{") and p.endswith("}"):
            k = p[1:-1]
            if k == "/":
                out.append("CHAR(10)")
            elif k.startswith("@"):
                out.append(f"[{k[1:]}]")
            elif k.startswith("="):
                out.append("(" + k[1:] + ")")
            else:
                out.append(N(k))
        else:
            out.append(lit(p))
    return "&".join(out) if out else '""'


def resolve(tpl, mp):
    return re.sub(r"\[(\w+)\]", lambda m: mp[m.group(1)], tpl)


def arr(opts):
    return "{" + ",".join('"' + o + '"' for o in opts) + "}"


def fator(chave, opts, facs):
    return f"CHOOSE(MATCH([{chave}],{arr(opts)},0),{','.join(str(x) for x in facs)})"


def _fem(art):
    return art.startswith("a ")


def _cap(art):
    return art[0].upper() + art[1:]


def _skip(lst, skip):
    if not skip:
        return lst
    return [(f"AND(NOT({skip}),{c})", m) for c, m in lst]


def chk_pos(k, art, skip=None):
    return _skip([
        (f'[{k}]=""', f"Informe {art} para continuar."),
        (f"NOT(ISNUMBER([{k}]))", f"{_cap(art)} precisa ser um número. Confira o valor digitado."),
        (f"[{k}]<=0", f"{_cap(art)} deve ser maior que zero."),
    ], skip)


def chk_nn(k, art, skip=None):
    neg = "negativa" if _fem(art) else "negativo"
    return _skip([
        (f'[{k}]=""', f"Informe {art} para continuar (digite 0 se for o caso)."),
        (f"NOT(ISNUMBER([{k}]))", f"{_cap(art)} precisa ser um número. Confira o valor digitado."),
        (f"[{k}]<0", f"{_cap(art)} não pode ser {neg}."),
    ], skip)


def chk_un(fk, what):
    return [(f'[{fk}]=""', f"Confira a unidade selecionada para {what}.")]


def aninhar(checks, fim='""'):
    f = fim
    for cond, msg in reversed(checks):
        f = f"IF({cond},{lit(msg)},{f})"
    return f
