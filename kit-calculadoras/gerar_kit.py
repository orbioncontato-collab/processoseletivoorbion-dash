# -*- coding: utf-8 -*-
"""Gera KIT_6_CALCULADORAS_ENGENHARIA_AMBIENTAL.xlsx (compatível com Excel e Google Planilhas).

Uso: python gerar_kit.py saida.xlsx [alturas.json] [registro.json]
"""
import json
import re
import sys

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from nucleo import *  # noqa: F401,F403
from nucleo import (FONTE, AL, AL_C, AL_L, AL_LT, AL_R, AMARELO, AMARELO_B, AMARELO_S, ARDOSIA, AZUL, AZUL_L, AZUL_XL,
                    BD, BRANCO, CENARIOS, CINZA_B, CINZA_L, CINZA_M, CINZA_T, CINZA_XL, DESBLOQ, FL, NAVY, NAVY2,
                    TXT, VERDE_D, VERDE_L, VERDE_OK, VERDE_T, VERM_L, VERM_T, F, aninhar, altura_texto, arr, box,
                    fmt_texto, link, lit, q, resolve, tint)
from specs import CALCS

SHEETS = ["00 — Painel Inicial", "01 — Comece Aqui", "02 — Vazão e Hidráulica", "03 — Carga Poluidora",
          "04 — Eficiência de Tratamento", "05 — Balanço de Massa", "06 — Diluição de Soluções",
          "07 — Dosagem Química", "08 — Exemplos Resolvidos"]
PAINEL, GUIA, BIBLIO = SHEETS[0], SHEETS[1], SHEETS[8]
ACC = {1: "2E75B6", 2: "C55A11", 3: "548235", 4: "7030A0", 5: "00808C", 6: "B8860B"}
WC = {"A": 2, "B": 26, "C": 12, "D": 12, "E": 14, "F": 10, "G": 13, "H": 36, "I": 2, "J": 4, "K": 34, "L": 18}
LCH = 97  # largura (caracteres) do bloco mesclado C:H
LBH = WC["B"] + LCH  # B:H


class Ctx:
    """Contexto de uma aba: ponteiro de linha, área técnica (K:L) e registro de alturas dinâmicas."""

    def __init__(self, ws, alturas, dyn):
        self.ws = ws
        self.name = ws.title
        self.r = 1
        self.h = 3
        self.alturas = alturas
        self.dyn = dyn
        self.prefix = ""
        box(ws, "K1:L1", "ÁREA TÉCNICA · cálculos auxiliares (não editar)", F(9, True, BRANCO), FL(CINZA_M), AL_L)
        box(ws, "K2", "Descrição", F(8, True, CINZA_T), FL(CINZA_L), AL_L)
        box(ws, "L2", "Valor", F(8, True, CINZA_T), FL(CINZA_L), AL_C)

    def helper(self, label, formula):
        r = self.h
        self.h += 1
        box(self.ws, f"K{r}", f"{self.prefix}{label}", F(8, color=CINZA_T), None, AL("left", wrap=False))
        box(self.ws, f"L{r}", formula, F(8, color=CINZA_T), None, AL("left", wrap=False))
        return f"$L${r}"

    def res(self, tpl, mp):
        for k in sorted(set(re.findall(r"\[(\w+?)__t\]", tpl))):
            if k + "__t" not in mp:
                b = mp[k]
                xr = self.helper(f"{k} (arredondado)",
                                 f'=IFERROR(ROUND({b},IF({b}=0,0,4-INT(LOG10(ABS({b}))))),"")')
                mp[k + "__t"] = self.helper(f"{k} (texto)", fmt_texto(xr))
        return resolve(tpl, mp)

    def altura(self, r, h):
        self.ws.row_dimensions[r].height = self.alturas.get(f"{self.name}|{r}", h)

    def dinamica(self, r, cell, largura, padrao=30, size=10):
        self.dyn.append([self.name, r, cell, largura, size])
        self.altura(r, padrao)


def colunas(ws, larguras):
    for c, w in larguras.items():
        ws.column_dimensions[c].width = w


def proteger(ws):
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


def dv_add(ws, cell, **kw):
    dv = DataValidation(allow_blank=True, showErrorMessage=True, showInputMessage=True, **kw)
    ws.add_data_validation(dv)
    dv.add(cell)


def cf(ws, rng, formula, fill=None, font=None):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=fill, font=font, stopIfTrue=False))


def botao(ws, rng, texto, sheet, ref="A1", cor=AZUL, size=10):
    c = box(ws, rng, texto, F(size, True, BRANCO), FL(cor), AL_C, BD(BRANCO))
    link(c, sheet, ref)
    return c


def secao(ws, r, texto, cor=NAVY, size=11, h=24, rng=None):
    box(ws, rng or f"B{r}:H{r}", texto, F(size, True, BRANCO), FL(cor), AL_L)
    ws.row_dimensions[r].height = h


def subsecao(ws, r, texto):
    box(ws, f"B{r}:H{r}", texto, F(9, True, NAVY), FL("DCE6F1"), AL_L)
    ws.row_dimensions[r].height = 20


def linha_texto(cx, r, rotulo, valor, *, fill=CINZA_XL, rfill=CINZA_L, font=None, dyn=False, padrao=30, h=None):
    ws = cx.ws
    box(ws, f"B{r}", rotulo, F(9, True, CINZA_T), FL(rfill), AL_L, BD())
    c = box(ws, f"C{r}:H{r}", valor, font or F(10), FL(fill), AL_L, BD())
    if dyn:
        cx.dinamica(r, f"C{r}", LCH, padrao, (font.sz if font else 10))
    else:
        txt = valor if isinstance(valor, str) else ""
        sz = font.sz if font else 10
        cx.altura(r, h or altura_texto(txt, LCH, sz, bool(font and font.b)))
    return c


def espaco(ws, r, h=8):
    ws.row_dimensions[r].height = h


# ====================================================================== cabeçalho de aba
def cabecalho(cx, titulo, subtitulo, cor, idx_aba):
    ws = cx.ws
    box(ws, "B1:H1", titulo, F(17, True, BRANCO), FL(NAVY), AL_L)
    ws.row_dimensions[1].height = 40
    box(ws, "B2:H2", subtitulo, F(9, color="D9E2F3", italic=True), FL(NAVY), AL_L)
    ws.row_dimensions[2].height = 18
    for col in "AI":
        ws[f"{col}1"].fill = FL(cor)
        ws[f"{col}2"].fill = FL(cor)
    botao(ws, "B3", "🏠  Painel", PAINEL, cor=AZUL)
    botao(ws, "C3:D3", "📘  Guia", GUIA, cor=AZUL)
    botao(ws, "E3:F3", "📚  Exemplos", BIBLIO, cor=AZUL)
    ant = SHEETS[idx_aba - 1] if idx_aba > 0 else SHEETS[0]
    prox = SHEETS[idx_aba + 1] if idx_aba + 1 < len(SHEETS) else SHEETS[0]
    botao(ws, "G3", "◀ Anterior", ant, cor=NAVY2, size=9)
    botao(ws, "H3", "Próxima aba ▶", prox, cor=NAVY2, size=9)
    ws.row_dimensions[3].height = 26
    ws.freeze_panes = "A4"
    espaco(ws, 4, 10)


def configurar_aba(ws, cor, zoom=100):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.tabColor = cor
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


# ====================================================================== bloco de ferramenta
def bloco(cx, T, num_txt, sel=None, k_sel=None, acc=AZUL):
    ws = cx.ws
    mp = {}
    r = cx.r
    info = dict(title_row=r, my={}, inuse={}, ex={})
    # título
    box(ws, f"B{r}:G{r}", num_txt + T["title"].upper(), F(13, True, BRANCO), FL(AZUL), AL_L)
    c = box(ws, f"H{r}", "▲ Voltar ao topo", F(9, True, BRANCO, underline="single"), FL(AZUL), AL_R)
    link(c, cx.name, "A1")
    ws.row_dimensions[r].height = 30
    if sel:
        cf(ws, f"B{r}:H{r}", f"{sel}={k_sel}", fill=FL(VERDE_D))
    r += 1
    linha = box(ws, f"B{r}:H{r}", T["desc"], F(10, color=CINZA_T, italic=True), FL(CINZA_XL), AL_L)
    cx.altura(r, altura_texto(T["desc"], LBH, 10))
    r += 1
    espaco(ws, r, 6)
    r += 1
    # B · cenário
    subsecao(ws, r, "B · ESCOLHA DO CENÁRIO")
    r += 1
    box(ws, f"B{r}", "Cenário  ▼", F(10, True, NAVY), FL(BRANCO), AL_L, BD())
    sc = box(ws, f"C{r}:D{r}", "Exemplo 01", F(11, True, TXT), FL(AMARELO), AL_C, BD(AMARELO_B, "medium"), unlock=True)
    dv_add(ws, f"C{r}", type="list", formula1='"' + ",".join(CENARIOS) + '"', errorTitle="Opção inválida",
           error="Escolha Exemplo 01, Exemplo 02 ou Meu cálculo.", promptTitle="Escolha o cenário",
           prompt="Exemplo 01 e Exemplo 02 usam dados prontos. Meu cálculo usa os valores que você digitar na "
                  "coluna amarela MEU CÁLCULO.")
    box(ws, f"E{r}:H{r}", "Exemplo 01 e Exemplo 02 usam dados prontos. Meu cálculo usa os valores que você digitar "
                          "na coluna amarela.", F(9, color=CINZA_T, italic=True), None, AL_L)
    ws.row_dimensions[r].height = 30
    info["scen"] = f"C{r}"
    mp["idx"] = cx.helper("índice do cenário (1, 2 ou 3)",
                          f'=IFERROR(MATCH($C${r},{arr(CENARIOS)},0),0)')
    idx = mp["idx"]
    r += 1
    espaco(ws, r, 6)
    r += 1
    # C · dados de entrada
    subsecao(ws, r, "C · DADOS DE ENTRADA")
    r += 1
    heads = [("B", "Variável"), ("C", "Exemplo 01"), ("D", "Exemplo 02"), ("E", "MEU CÁLCULO"),
             ("F", "Unidade"), ("G", "Valor em uso"), ("H", "Orientação de preenchimento")]
    for col, t in heads:
        cor = AMARELO_B if col == "E" else (VERDE_D if col == "G" else ARDOSIA)
        box(ws, f"{col}{r}", t, F(9, True, BRANCO), FL(cor), AL_C, BD(BRANCO))
    ws.row_dimensions[r].height = 22
    r += 1
    first_in = r
    rows = {}
    for inp in T["inputs"]:
        k = inp["key"]
        rows[k] = r
        kind = inp["kind"]
        box(ws, f"B{r}", inp["label"], F(10, True, NAVY), FL(BRANCO), AL_L, BD())
        for col, v in zip("CD", inp["ex"]):
            box(ws, f"{col}{r}", v, F(10, color=CINZA_T), FL(CINZA_L), AL_C, BD())
        my = inp.get("my")
        box(ws, f"E{r}", my, F(10, True), FL(AMARELO), AL_C, BD(AMARELO_B), unlock=True)
        base = f'IF({idx}=1,C{r},IF({idx}=2,D{r},IF({idx}=3,IF(E{r}="","",E{r}),"")))'
        if inp.get("target"):
            base = f'IF({resolve(T["mode_expr"], mp)}={inp["target"]},"calculado ▼",{base})'
        box(ws, f"G{r}", "=" + base, F(10, True, VERDE_T), FL(AZUL_L), AL_C, BD())
        mp[k] = f"$G${r}"
        info["my"][k] = f"E{r}"
        info["inuse"][k] = f"G{r}"
        info["ex"][k] = (f"$C${r}", f"$D${r}")
        box(ws, f"H{r}", inp["hint"], F(9, color=CINZA_T, italic=True), FL(BRANCO), AL_L, BD())
        cx.altura(r, max(22, altura_texto(inp["hint"], WC["H"], 9)))
        prompt = inp["hint"][:250]
        title = inp["label"][:32]
        if kind == "list":
            dv_add(ws, f"E{r}", type="list", formula1='"' + ",".join(inp["opts"]) + '"', errorTitle="Opção inválida",
                   error="Escolha uma opção da lista.", promptTitle=title, prompt=prompt)
        elif kind == "pct":
            dv_add(ws, f"E{r}", type="decimal", operator="between", formula1="0", formula2="100",
                   errorTitle="Porcentagem inválida", error="Digite a porcentagem de 0 a 100. Para 50%, digite 50.",
                   promptTitle=title, prompt=prompt)
        elif kind == "text":
            dv_add(ws, f"E{r}", type="textLength", operator="lessThanOrEqual", formula1="60",
                   errorTitle="Texto longo", error="Use até 60 caracteres.", promptTitle=title, prompt=prompt)
        else:
            dv_add(ws, f"E{r}", type="decimal", operator="greaterThanOrEqual", formula1="0",
                   errorTitle="Valor inválido",
                   error="Digite um número maior ou igual a zero, sem a unidade. Use vírgula para decimais (ex.: 0,5).",
                   promptTitle=title, prompt=prompt)
        r += 1
    last_in = r - 1
    # unidades (coluna F)
    for inp in T["inputs"]:
        rr = rows[inp["key"]]
        u = inp.get("unit")
        if inp["kind"] == "list":
            u = "lista ▼"
        if u and u.startswith("@"):
            v = "=" + mp[u[1:]]
        else:
            v = u or ""
        box(ws, f"F{rr}", v, F(9, True, CINZA_T), FL(BRANCO), AL_C, BD())
    # destaque da coluna ativa
    cf(ws, f"C{first_in}:C{last_in}", f"{idx}=1", fill=FL(AZUL_L), font=F(10, True, NAVY))
    cf(ws, f"D{first_in}:D{last_in}", f"{idx}=2", fill=FL(AZUL_L), font=F(10, True, NAVY))
    cf(ws, f"E{first_in}:E{last_in}", f"{idx}=3", fill=FL(AMARELO_S))
    # cálculos auxiliares
    cx.prefix = f"[{T['short']}] "
    for k, tpl in T["calcs"]:
        mp[k] = cx.helper(k, "=IFERROR(" + resolve(tpl, mp) + ',"")')
    for inp in T["inputs"]:
        if inp.get("target"):
            rr = rows[inp["key"]]
            cf(ws, f"B{rr}:H{rr}", f"{mp['m']}={inp['target']}", fill=FL("EDEDED"), font=F(10, color=CINZA_M, italic=True))
    checks = [("[idx]=0", "Escolha um cenário na lista: Exemplo 01, Exemplo 02 ou Meu cálculo.")] + T["checks"]
    err = cx.helper("mensagem de validação",
                    "=IFERROR(" + resolve(aninhar(checks), mp) + ',"Resultado indisponível: existem dados inválidos.")')
    ok = cx.helper("dados válidos?", f'={err}=""')
    mp["err"], mp["ok"] = err, ok
    info["ok"], info["err"] = ok, err
    espaco(ws, r, 8)
    r += 1

    # D · resultado
    def unidade(u):
        if u.startswith("@"):
            return "=" + mp[u[1:]]
        if u.startswith("="):
            return "=" + cx.res(u[1:], mp)
        return u

    secao(ws, r, "D · RESULTADO AUTOMÁTICO", VERDE_D, 11, 24)
    r += 1
    lab, key, uni, fmt = T["main"]
    box(ws, f"B{r}:D{r}", ("=" + cx.res(lab[1:], mp)) if lab.startswith("=") else lab, F(10, True, VERDE_T),
        FL(VERDE_L), AL_L, BD(VERDE_L))
    box(ws, f"E{r}:H{r}", "OUTROS RESULTADOS", F(9, True, VERDE_T), FL(VERDE_L), AL_L, BD(VERDE_L))
    ws.row_dimensions[r].height = 22
    r += 1
    vr = r
    mv = box(ws, f"B{r}:C{r + 3}", f'=IF({ok},{mp[key]},"—")', F(24, True, VERDE_T), FL(BRANCO), AL_C,
             BD(VERDE_D, "medium"), fmt)
    box(ws, f"D{r}:D{r + 3}", unidade(uni), F(13, True, VERDE_T), FL(BRANCO), AL_C, BD(VERDE_D, "medium"))
    cf(ws, f"B{r}:C{r + 3}", f"NOT({ok})", font=F(24, True, CINZA_M))
    info["main"] = f"B{r}"
    info["sec"] = []
    for i in range(4):
        rr = r + i
        if i < len(T["sec"]):
            sl, sk, su, sf = T["sec"][i]
            sk = sk.lstrip("=")
            box(ws, f"E{rr}:F{rr}", sl, F(9, color=TXT), FL(VERDE_L), AL_L, BD(BRANCO))
            box(ws, f"G{rr}", f'=IF({ok},{mp[sk]},"—")', F(11, True, VERDE_T), FL(BRANCO), AL_C, BD(VERDE_L), sf)
            box(ws, f"H{rr}", unidade(su), F(9, True, CINZA_T), FL(VERDE_L), AL_L, BD(BRANCO))
            info["sec"].append(f"G{rr}")
        else:
            box(ws, f"E{rr}:H{rr}", None, None, FL(VERDE_L), None, None)
        ws.row_dimensions[rr].height = 21
    r += 4
    box(ws, f"B{r}", "Interpretação", F(9, True, VERDE_T), FL(VERDE_L), AL_L, BD(BRANCO))
    box(ws, f"C{r}:H{r}", f'=IF({ok},' + cx.res(T["interp"], mp) +
        ',"Resultado indisponível: corrija o item indicado em Status dos dados.")', F(10), FL(BRANCO), AL_L, BD(VERDE_L))
    cx.dinamica(r, f"C{r}", LCH, 45)
    info["interp"] = f"C{r}"
    r += 1
    box(ws, f"B{r}", "Status dos dados", F(9, True, VERDE_T), FL(VERDE_L), AL_L, BD(BRANCO))
    box(ws, f"C{r}:H{r}", f'=IF({ok},"✔ Dados válidos. Resultado calculado automaticamente.","⚠ "&{err})',
        F(10, True), FL(BRANCO), AL_L, BD(VERDE_L))
    cf(ws, f"C{r}:H{r}", f"{ok}", fill=FL(VERDE_OK), font=F(10, True, VERDE_T))
    cf(ws, f"C{r}:H{r}", f"NOT({ok})", fill=FL(VERM_L), font=F(10, True, VERM_T))
    cx.dinamica(r, f"C{r}", LCH, 22)
    info["status"] = f"C{r}"
    r += 1
    espaco(ws, r, 10)
    r += 1

    # esquema visual (diluição)
    if T.get("esquema"):
        secao(ws, r, "ESQUEMA DE PREPARO DA SOLUÇÃO", "00808C", 11, 24)
        r += 1
        faixas = [f"B{r}", f"C{r}:D{r}", f"E{r}:G{r}", f"H{r}"]
        tons = ["006D77", "00808C", "2A9D8F", VERDE_D]
        for (tit, _), rng, cor in zip(T["esquema"], faixas, tons):
            box(ws, rng, tit, F(10, True, BRANCO), FL(cor), AL_C, BD(BRANCO))
        ws.row_dimensions[r].height = 26
        r += 1
        faixas = [f"B{r}", f"C{r}:D{r}", f"E{r}:G{r}", f"H{r}"]
        for (_, tpl), rng in zip(T["esquema"], faixas):
            box(ws, rng, f'=IF({ok},' + cx.res(tpl, mp) + ',"—")', F(10, True, "004D55"), FL("E0F2F1"), AL_C,
                BD(BRANCO))
        ws.row_dimensions[r].height = 58
        r += 1
        espaco(ws, r, 10)
        r += 1

    # E · entenda
    secao(ws, r, "E · ENTENDA O CÁLCULO", ARDOSIA, 11, 24)
    r += 1
    for rot, tipo, conteudo in T["entenda"]:
        if tipo == "formula":
            linha_texto(cx, r, rot, conteudo, fill=AZUL_XL, font=F(12, True, NAVY), h=max(26, altura_texto(conteudo, LCH, 12, True)))
        elif tipo == "static":
            linha_texto(cx, r, rot, conteudo)
        elif tipo == "live":
            linha_texto(cx, r, rot, "=" + cx.res(valores_tpl(T), mp), dyn=True, padrao=22)
        else:
            linha_texto(cx, r, rot, f'=IF({ok},' + cx.res(conteudo, mp) +
                        ',"Aguardando dados válidos para montar a conta.")', dyn=True, padrao=30)
        r += 1
    espaco(ws, r, 10)
    r += 1

    # G · gráfico
    if T.get("chart"):
        ch = T["chart"]
        secao(ws, r, "G · COMPARAÇÃO VISUAL", ARDOSIA, 11, 24)
        r += 1
        top = r
        box(ws, f"E{r}:H{r}", "Dados do gráfico (atualizam sozinhos)", F(9, True, CINZA_T), FL(CINZA_L), AL_L)
        r += 1
        d0 = r
        for lab_i, key_i, uni_i in ch["items"]:
            box(ws, f"E{r}:F{r}", lab_i, F(9), FL(CINZA_XL), AL_L, BD())
            box(ws, f"G{r}", f'=IF({ok},{mp[key_i]},"")', F(10, True, VERDE_T), FL(BRANCO), AL_C, BD(), FMT_POS_LOCAL)
            box(ws, f"H{r}", unidade(uni_i), F(9, color=CINZA_T), FL(CINZA_XL), AL_L, BD())
            r += 1
        d1 = r - 1
        chart = BarChart()
        chart.type = "col"
        chart.title = ch["title"]
        chart.legend = None
        chart.add_data(Reference(ws, min_col=7, min_row=d0, max_row=d1), titles_from_data=False)
        chart.set_categories(Reference(ws, min_col=5, min_row=d0, max_row=d1))
        s = chart.series[0]
        s.graphicalProperties.solidFill = ch["colors"][0]
        s.graphicalProperties.line.solidFill = ch["colors"][0]
        for i, cor in enumerate(ch["colors"]):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = cor
            pt.graphicalProperties.line.solidFill = cor
            s.dPt.append(pt)
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        s.dLbls.showSerName = False
        s.dLbls.showCatName = False
        s.dLbls.showLegendKey = False
        chart.y_axis.title = ch["ytitle"]
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.gapWidth = 70
        chart.width = 9.2
        chart.height = 6.6
        ws.add_chart(chart, f"B{top}")
        alvo = top + 14
        while r < alvo:
            r += 1
        for rr in range(top, alvo):
            if ws.row_dimensions[rr].height is None:
                ws.row_dimensions[rr].height = 15
        box(ws, f"E{d1 + 2}:H{d1 + 4}", "O gráfico usa os mesmos resultados do card verde. Mude um dado e veja as "
                                        "barras mudarem.", F(9, color=CINZA_T, italic=True), None, AL_LT)
        espaco(ws, r, 10)
        r += 1

    # F · exemplo e orientação
    secao(ws, r, "F · EXEMPLO E ORIENTAÇÃO", ARDOSIA, 11, 24)
    r += 1
    g = T["guia"]
    ex_txt = (f"Exemplo 01: {T['ex_info'][0][1]}\nExemplo 02: {T['ex_info'][1][1]}\n"
              "Escolha o exemplo na lista Cenário para ver a resolução completa.")
    linha_texto(cx, r, "Exemplos desta ferramenta", ex_txt)
    r += 1
    linha_texto(cx, r, "Como interpretar", g["interpretar"])
    r += 1
    linha_texto(cx, r, "⚠ Erro comum", g["erro"], fill=VERM_L, rfill="F8CBAD")
    r += 1
    linha_texto(cx, r, "💡 Dica prática", g["dica"], fill=AZUL_XL, rfill=AZUL_L)
    r += 1
    if g.get("atencao"):
        linha_texto(cx, r, "Atenção", g["atencao"], fill="FFF4E5", rfill="FBE5D6")
        r += 1
    espaco(ws, r, 14)
    r += 1
    espaco(ws, r, 14)
    r += 1
    cx.r = r
    cx.prefix = ""
    # resumo para o seletor do topo
    unit_main = T["main"][2]
    if unit_main.startswith("@"):
        ue = "[" + unit_main[1:] + "]"
    elif unit_main.startswith("="):
        ue = "[runit]"  # unidade dinâmica (calculadora de diluição)
    else:
        ue = lit(unit_main)
    info["summary"] = cx.helper(f"[{T['short']}] resumo", "=" + cx.res(f'IF([ok],{N(key)}&" "&{ue},"—")', mp))
    info["mp"] = mp
    return info


FMT_POS_LOCAL = "#,##0.00###"


def valores_tpl(T):
    if T.get("values"):
        return T["values"]
    partes = []
    for inp in T["inputs"]:
        k, kind = inp["key"], inp["kind"]
        if kind == "list":
            if inp.get("show"):
                partes.append(lit(inp["label"] + ": ") + f"&[{k}]")
            continue
        if kind == "text":
            partes.append(lit(inp["sym"] + ": ") + f'&IF([{k}]="","não informado",[{k}])')
            continue
        u = inp["unit"]
        ue = f"[{u[1:]}]" if u.startswith("@") else lit(u)
        if kind == "dens":
            partes.append(lit(inp["sym"] + " = ") + f'&IF([{k}]="","não informada",{N(k)}&" "&{ue})')
        else:
            partes.append(lit(inp["sym"] + " = ") + f"&{N(k)}&" + lit(" ") + f"&{ue}")
    return '&"   |   "&'.join(partes)


# ====================================================================== calculadora
def calculadora(wb, C, alturas, dyn, reg):
    n = C["n"]
    nome = SHEETS[n + 1]
    ws = wb.create_sheet(nome)
    acc = ACC[n]
    configurar_aba(ws, acc)
    colunas(ws, WC)
    cx = Ctx(ws, alturas, dyn)
    cabecalho(cx, f"{C['emoji']}  CALCULADORA {n:02d}  ·  {C['nome'].upper()}",
              "KIT DE 6 CALCULADORAS AUTOMÁTICAS DE ENGENHARIA AMBIENTAL", acc, n + 1)
    r = 5
    tools = C["tools"]
    passos = ("COMO USAR:   ① Escolha o cenário na lista amarela   ② Para usar seus dados, escolha Meu cálculo e digite "
              "na coluna amarela MEU CÁLCULO   ③ Veja o resultado no card verde   ④ Confira a conta em "
              "ENTENDA O CÁLCULO")
    if len(tools) > 1:
        passos = (
            "COMO USAR:   ① Escolha a ferramenta abaixo   ② Escolha o cenário na lista amarela   ③ Para usar seus "
            "dados, escolha Meu cálculo e digite na coluna amarela   ④ Veja o resultado no card verde   ⑤ Confira a "
            "conta em ENTENDA O CÁLCULO")
    box(ws, f"B{r}:H{r}", passos, F(10, True, NAVY), FL(AZUL_L), AL_L, BD(AZUL_L))
    cx.altura(r, altura_texto(passos, LBH, 10, True))
    r += 1
    espaco(ws, r, 8)
    r += 1
    secao(ws, r, "A · IDENTIFICAÇÃO", NAVY, 11, 24)
    r += 1
    for rot, txt in (("O que calcula", C["what"]), ("Quando usar", C["when"]), ("Exemplo prático", C["example"])):
        linha_texto(cx, r, rot, txt)
        r += 1
    espaco(ws, r, 10)
    r += 1
    secao(ws, r, "FERRAMENTAS DESTA CALCULADORA" if len(tools) > 1 else "COMO ESTA CALCULADORA FUNCIONA", NAVY, 11, 24)
    r += 1
    idx_rows = []
    for i, T in enumerate(tools):
        idx_rows.append(r)
        rot = f"▶ Ferramenta {i + 1}" if len(tools) > 1 else "▶ Ir para o cálculo"
        box(ws, f"B{r}", rot, F(10, True, BRANCO), FL(AZUL), AL_C, BD(BRANCO))
        box(ws, f"C{r}:H{r}", f"{T['title']}: {T['desc']}", F(9, color=TXT), FL(CINZA_XL), AL_L, BD())
        cx.altura(r, max(24, altura_texto(T["title"] + ": " + T["desc"], LCH, 9)))
        r += 1
    sel = None
    resumo_rows = None
    if len(tools) > 1:
        espaco(ws, r, 6)
        r += 1
        opts = [f"Ferramenta {i + 1}: {T['short']}" for i, T in enumerate(tools)]
        box(ws, f"B{r}", "Ferramenta escolhida ▼", F(10, True, NAVY), FL(BRANCO), AL_L, BD())
        box(ws, f"C{r}:E{r}", opts[0], F(10, True), FL(AMARELO), AL_C, BD(AMARELO_B, "medium"), unlock=True)
        dv_add(ws, f"C{r}", type="list", formula1='"' + ",".join(opts) + '"', errorTitle="Opção inválida",
               error="Escolha uma ferramenta da lista.", promptTitle="Escolha a ferramenta",
               prompt="A ferramenta escolhida fica destacada em verde e o resultado dela aparece aqui ao lado.")
        box(ws, f"F{r}:H{r}", "A ferramenta escolhida fica com o título em verde. Use os botões azuis acima para ir "
                              "até ela.", F(9, color=CINZA_T, italic=True), None, AL_L)
        ws.row_dimensions[r].height = 30
        sel = cx.helper("ferramenta escolhida (número)", f"=IFERROR(MATCH($C${r},{arr(opts)},0),0)")
        reg["tool_sel"] = f"C{r}"
        r += 1
        resumo_rows = (r, r + 1)
        r += 2
    espaco(ws, r, 14)
    r += 1
    cx.r = r
    infos = []
    for i, T in enumerate(tools):
        num_txt = f"FERRAMENTA {i + 1}  ·  " if len(tools) > 1 else ""
        info = bloco(cx, T, num_txt, sel, i + 1, acc)
        infos.append(info)
        link(ws[f"B{idx_rows[i]}"], nome, f"B{info['title_row']}")
    if resumo_rows:
        r1, r2 = resumo_rows
        box(ws, f"B{r1}", "Resultado atual", F(10, True, VERDE_T), FL(VERDE_L), AL_L, BD())
        lst_s = ",".join(i["summary"] for i in infos)
        box(ws, f"C{r1}:H{r1}", f'=IFERROR(CHOOSE({sel},{lst_s}),"—")', F(14, True, VERDE_T), FL(BRANCO), AL_L,
            BD(VERDE_D))
        ws.row_dimensions[r1].height = 28
        box(ws, f"B{r2}", "Status", F(10, True, VERDE_T), FL(VERDE_L), AL_L, BD())
        lst_st = ",".join(i["status"] for i in infos)
        box(ws, f"C{r2}:H{r2}", f'=IFERROR(CHOOSE({sel},{lst_st}),"Escolha uma ferramenta.")', F(10, True),
            FL(BRANCO), AL_L, BD())
        lst_ok = ",".join(i["ok"] for i in infos)
        cf(ws, f"C{r2}:H{r2}", f"IFERROR(CHOOSE({sel},{lst_ok}),FALSE)", fill=FL(VERDE_OK), font=F(10, True, VERDE_T))
        cf(ws, f"C{r2}:H{r2}", f"NOT(IFERROR(CHOOSE({sel},{lst_ok}),FALSE))", fill=FL(VERM_L), font=F(10, True, VERM_T))
        cx.dinamica(r2, f"C{r2}", LCH, 22)
        reg["tool_summary"] = f"C{r1}"
    # rodapé
    r = cx.r
    box(ws, f"B{r}:H{r}", "Lembrete: os resultados são cálculos teóricos de apoio ao estudo. Confira unidades e "
                          "hipóteses antes de usar em trabalhos, relatórios ou no estágio.",
        F(9, color=CINZA_T, italic=True), FL(CINZA_L), AL_L)
    cx.altura(r, 30)
    ws.print_area = f"A1:I{r}"
    proteger(ws)
    reg["tools"] = [{k: v for k, v in i.items() if k != "mp"} for i in infos]
    return ws, infos


# ====================================================================== painel
def painel(wb, alturas, dyn):
    ws = wb.active
    ws.title = PAINEL
    configurar_aba(ws, NAVY)
    ws.sheet_view.showRowColHeaders = False
    colunas(ws, {"A": 3, "B": 34, "C": 3, "D": 34, "E": 3, "F": 34, "G": 3})
    for r in range(1, 60):
        for col in "ABCDEFG":
            ws[f"{col}{r}"].fill = FL(BRANCO)
    espaco(ws, 1, 12)
    box(ws, "B2:F2", "KIT DE 6 CALCULADORAS AUTOMÁTICAS DE ENGENHARIA AMBIENTAL", F(18, True, BRANCO), FL(NAVY), AL_C)
    ws.row_dimensions[2].height = 52
    box(ws, "B3:F3", "Resolva, confira e entenda cálculos de hidráulica, carga poluidora, eficiência, balanço de "
                     "massa, diluição e dosagem química. Tudo automático, com a fórmula explicada passo a passo.",
        F(11, color="D9E2F3"), FL(NAVY), AL_C)
    ws.row_dimensions[3].height = 40
    box(ws, "B4:F4", None, None, FL("70AD47"))
    ws.row_dimensions[4].height = 5
    espaco(ws, 5, 14)
    botao(ws, "B6", "📘  COMECE AQUI (GUIA)", GUIA, cor=AZUL, size=11)
    botao(ws, "D6", "📚  BIBLIOTECA DE EXEMPLOS", BIBLIO, cor=AZUL, size=11)
    botao(ws, "F6", "▶  COMO USAR EM 5 PASSOS", GUIA, "B20", cor=VERDE_D, size=11)
    ws.row_dimensions[6].height = 34
    espaco(ws, 7, 16)
    box(ws, "B8:F8", "ESCOLHA SUA CALCULADORA", F(13, True, NAVY), None, AL("left"))
    ws.row_dimensions[8].height = 26
    cols = ["B", "D", "F"]
    base = 9
    for i, C in enumerate(CALCS):
        col = cols[i % 3]
        r0 = base + (i // 3) * 7
        acc = ACC[C["n"]]
        nome = SHEETS[C["n"] + 1]
        box(ws, f"{col}{r0}", f"{C['emoji']}   CALCULADORA {C['n']:02d}", F(11, True, BRANCO),
            FL(acc), AL_L)
        box(ws, f"{col}{r0 + 1}", C["nome"], F(13, True, NAVY), FL(CINZA_XL), AL_L)
        box(ws, f"{col}{r0 + 2}", C["card"], F(10, color=CINZA_T), FL(CINZA_XL), AL_LT)
        box(ws, f"{col}{r0 + 3}", C["card2"] + f"  ·  aba {C['n'] + 1:02d}", F(9, True, acc, italic=True), FL(CINZA_XL), AL_L)
        botao(ws, f"{col}{r0 + 4}", "Abrir calculadora  ▶", nome, cor=AZUL)
        for rr in range(r0, r0 + 5):
            c = ws[f"{col}{rr}"]
            c.border = BD(CINZA_B)
        ws.row_dimensions[r0].height = 30
        ws.row_dimensions[r0 + 1].height = 36
        ws.row_dimensions[r0 + 2].height = 70
        ws.row_dimensions[r0 + 3].height = 30
        ws.row_dimensions[r0 + 4].height = 28
        ws.row_dimensions[r0 + 5].height = 16
        ws.row_dimensions[r0 + 6].height = 6
    r = base + 14
    box(ws, f"B{r}:D{r}", "COMO COMEÇAR", F(12, True, BRANCO), FL(NAVY), AL_L)
    box(ws, f"F{r}", "LEGENDA DAS CORES", F(12, True, BRANCO), FL(NAVY), AL_L)
    ws.row_dimensions[r].height = 26
    passos = ["1.  Escolha uma calculadora e clique em Abrir calculadora.",
              "2.  Na lista Cenário, escolha Exemplo 01, Exemplo 02 ou Meu cálculo.",
              "3.  Em Meu cálculo, preencha os campos amarelos da coluna MEU CÁLCULO.",
              "4.  Confira o resultado automático destacado em verde.",
              "5.  Leia a explicação em Entenda o cálculo e verifique as unidades."]
    legenda = [("Amarelo-claro: campo para você digitar ou escolher", AMARELO, TXT),
               ("Verde: resultado automático e indicadores positivos", VERDE_OK, VERDE_T),
               ("Azul-claro: valor em uso, que entrou na conta", AZUL_L, NAVY),
               ("Azul: botão de navegação (clique)", AZUL, BRANCO),
               ("Cinza-claro: explicações e exemplos cadastrados", CINZA_L, CINZA_T),
               ("Vermelho suave: avisos e dados inválidos", VERM_L, VERM_T)]
    for i in range(6):
        rr = r + 1 + i
        if i < 5:
            box(ws, f"B{rr}:D{rr}", passos[i], F(10, color=TXT), FL(CINZA_XL), AL_L, BD())
        else:
            box(ws, f"B{rr}:D{rr}", "Dúvidas? Abra o guia completo na aba 01: Comece Aqui.", F(10, True, AZUL),
                FL(CINZA_XL), AL_L, BD())
            link(ws[f"B{rr}"], GUIA)
        t, fc, tc = legenda[i]
        box(ws, f"F{rr}", t, F(9, True, tc), FL(fc), AL_L, BD(BRANCO))
        ws.row_dimensions[rr].height = 24
    r += 8
    nota = ("Importante: os resultados são cálculos teóricos de apoio ao estudo. Projetos, laudos e decisões "
            "operacionais exigem normas técnicas, legislação aplicável e responsável técnico. Guarde uma cópia "
            "original deste arquivo e salve cada trabalho com outro nome.")
    box(ws, f"B{r}:F{r}", nota, F(9, color=CINZA_T, italic=True), FL(CINZA_L), AL_L)
    ws.row_dimensions[r].height = 40
    r += 1
    box(ws, f"B{r}:F{r}", "Sem macros · Sem internet · Excel e Google Planilhas · Abas protegidas sem senha contra "
                          "alterações acidentais", F(8, color=CINZA_M), None, AL_C)
    ws.row_dimensions[r].height = 20
    ws.print_area = f"A1:G{r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 1
    proteger(ws)


# ====================================================================== guia
def guia(wb, alturas, dyn):
    ws = wb.create_sheet(GUIA)
    configurar_aba(ws, VERDE_D)
    colunas(ws, {"A": 2, "B": 30, "C": 30, "D": 52, "E": 2, "F": 2, "G": 2, "H": 2})
    L3 = 112
    cx = Ctx.__new__(Ctx)
    cx.ws, cx.name, cx.alturas, cx.dyn = ws, GUIA, alturas, dyn
    box(ws, "B1:D1", "📘  COMECE AQUI: GUIA COMPLETO", F(17, True, BRANCO), FL(NAVY), AL_L)
    ws.row_dimensions[1].height = 40
    box(ws, "B2:D2", "Leia esta página uma vez e você saberá usar todo o kit.", F(9, color="D9E2F3", italic=True),
        FL(NAVY), AL_L)
    ws.row_dimensions[2].height = 18
    botao(ws, "B3", "🏠  Painel Inicial", PAINEL)
    botao(ws, "C3", "📚  Exemplos Resolvidos", BIBLIO)
    botao(ws, "D3", "Primeira calculadora ▶", SHEETS[2], cor=NAVY2)
    ws.row_dimensions[3].height = 26
    ws.freeze_panes = "A4"
    r = 5

    def sec(t):
        nonlocal r
        espaco(ws, r, 12)
        r += 1
        box(ws, f"B{r}:D{r}", t, F(12, True, BRANCO), FL(NAVY), AL_L)
        ws.row_dimensions[r].height = 26
        r += 1

    def par(t, fill=CINZA_XL, font=None):
        nonlocal r
        box(ws, f"B{r}:D{r}", t, font or F(10), FL(fill), AL_L, BD())
        ws.row_dimensions[r].height = altura_texto(t, L3, font.sz if font else 10, bool(font and font.b), 20, 8)
        r += 1

    def tabela(cab, linhas, links=None, cores=None):
        nonlocal r
        for col, t in zip("BCD", cab):
            box(ws, f"{col}{r}", t, F(10, True, BRANCO), FL(AZUL), AL_L, BD(BRANCO))
        ws.row_dimensions[r].height = 22
        r += 1
        for i, ln in enumerate(linhas):
            fill = CINZA_XL if i % 2 == 0 else BRANCO
            h = 20
            for j, (col, t) in enumerate(zip("BCD", ln)):
                f_ = F(10, j == 0, NAVY if j == 0 else TXT)
                fl = fill
                if cores and j == 0 and cores[i]:
                    fl, tc = cores[i]
                    f_ = F(10, True, tc)
                box(ws, f"{col}{r}", t, f_, FL(fl), AL_L, BD())
                h = max(h, altura_texto(t, {"B": 30, "C": 30, "D": 52}[col], 10, j == 0, 20, 8))
            if links and links[i]:
                link(ws[f"{links[i][0]}{r}"], links[i][1])
                ws[f"{links[i][0]}{r}"].font = F(10, True, AZUL, underline="single")
            ws.row_dimensions[r].height = h
            r += 1

    sec("1   O QUE É O KIT")
    par("Este kit reúne 6 calculadoras automáticas para os cálculos mais comuns da Engenharia Ambiental: vazão e "
        "hidráulica, carga poluidora, eficiência de tratamento, balanço de massa, diluição de soluções e dosagem de "
        "produtos químicos. Você informa os dados, a planilha calcula na hora e mostra a fórmula usada, a "
        "substituição dos valores e uma explicação do resultado. Funciona como um pequeno aplicativo dentro da "
        "planilha, sem macros e sem internet.")
    sec("2   PARA QUEM FOI DESENVOLVIDO")
    par("Para estudantes de Engenharia Ambiental e de cursos próximos (Sanitária, Civil, Química, Gestão Ambiental) "
        "que precisam resolver listas de exercícios, conferir contas de provas e trabalhos, preparar relatórios de "
        "laboratório ou apoiar atividades de estágio. Você não precisa saber Excel: basta escolher opções nas listas "
        "e digitar números nos campos amarelos.")
    sec("3   COMO USAR EM 5 PASSOS")
    passos = [
        ("PASSO 1", "Escolha uma calculadora.", "No Painel Inicial, clique no card da calculadora que resolve o seu "
         "problema. Na dúvida, use a tabela Qual calculadora usar, logo abaixo."),
        ("PASSO 2", "Selecione um exemplo ou a opção \"Meu cálculo\".", "Cada ferramenta tem a lista Cenário. "
         "Exemplo 01 e Exemplo 02 carregam dados prontos. Meu cálculo usa os valores que você digitar."),
        ("PASSO 3", "Preencha os campos amarelos.", "Digite seus valores na coluna MEU CÁLCULO e escolha as unidades "
         "nas listas. A coluna Valor em uso mostra exatamente o que está entrando na conta."),
        ("PASSO 4", "Confira o resultado automático destacado em verde.", "O card verde mostra o resultado principal, "
         "outras unidades, uma interpretação e o status dos dados. Se faltar algo, aparece uma mensagem em vermelho "
         "dizendo o que corrigir."),
        ("PASSO 5", "Leia a explicação e verifique as unidades utilizadas.", "Em Entenda o cálculo você vê a fórmula, "
         "os valores usados e a conta montada. Confira se as unidades fazem sentido para o seu exercício."),
    ]
    reg_passos = r
    for p, t, d in passos:
        box(ws, f"B{r}", p, F(14, True, BRANCO), FL(VERDE_D), AL_C, BD(BRANCO))
        box(ws, f"C{r}:D{r}", None, F(10), FL(VERDE_L), AL_L, BD(BRANCO))
        ws[f"C{r}"].value = CellRichText(TextBlock(InlineFont(rFont=FONTE, sz=11, b=True, color=VERDE_T), t + "\n"),
                                         TextBlock(InlineFont(rFont=FONTE, sz=10, color=TXT), d))
        ws.row_dimensions[r].height = max(40, altura_texto(t + "\n" + d, 82, 10, False, 20, 10))
        r += 1
    sec("4   QUAIS CÁLCULOS ESTÃO DISPONÍVEIS")
    linhas, links_ = [], []
    for C in CALCS:
        nome = SHEETS[C["n"] + 1]
        linhas.append((f"{C['emoji']} {C['n']:02d} · {C['curto']}", f"Aba {nome}",
                       "; ".join(f"{T['title']}" for T in C["tools"]) if C["n"] not in (5, 6) else
                       ("Volume de estoque, concentração final, volume final ou concentração do estoque"
                        if C["n"] == 5 else "Massa de ativo; massa de produto comercial; volume de produto líquido")))
        links_.append(("B", nome))
    tabela(("Calculadora (clique para abrir)", "Onde fica", "O que ela resolve"), linhas, links_)
    sec("5   COMO ESCOLHER A CALCULADORA CERTA")
    linhas = [
        ("Descobrir a vazão de uma tubulação", "02 · Vazão e Hidráulica (Ferramenta 2)", "Tubo de 0,30 m com água a 1,5 m/s."),
        ("Descobrir a vazão a partir de volume e tempo", "02 · Vazão e Hidráulica (Ferramenta 1)", "Tanque de 360 m³ que enche em 2 horas."),
        ("Verificar a velocidade da água em um tubo", "02 · Vazão e Hidráulica (Ferramenta 3)", "0,05 m³/s passando em tubo de 0,20 m."),
        ("Calcular o tempo de detenção de um tanque", "02 · Vazão e Hidráulica (Ferramenta 4)", "Tanque de 500 m³ recebendo 50 m³/h."),
        ("Calcular a carga de DBO de um efluente", "03 · Carga Poluidora (Ferramenta 1)", "400 m³/dia com 180 mg/L de DBO."),
        ("Descobrir a concentração a partir da carga", "03 · Carga Poluidora (Ferramenta 2)", "24 kg/dia de DBO em 400 m³/dia."),
        ("Avaliar a eficiência de uma ETE", "04 · Eficiência de Tratamento", "DBO de 250 mg/L na entrada e 50 mg/L na saída."),
        ("Calcular a concentração após misturar dois efluentes", "05 · Balanço de Massa (Ferramenta 1)", "120 m³/dia a 200 mg/L + 80 m³/dia a 50 mg/L."),
        ("Calcular a concentração do lodo de um decantador", "05 · Balanço de Massa (Ferramenta 2)", "Entrada, clarificado e lodo de um decantador primário."),
        ("Saber se a massa está acumulando em um sistema", "05 · Balanço de Massa (Ferramenta 3)", "Tanque que recebe mais sólidos do que descarta."),
        ("Preparar uma solução diluída", "06 · Diluição de Soluções", "500 mL a 200 mg/L a partir de um padrão de 1.000 mg/L."),
        ("Calcular a quantidade diária de um produto químico", "07 · Dosagem Química", "4 mg/L de ativo em 500 m³/dia com produto a 50%."),
    ]
    destino = {"02": SHEETS[2], "03": SHEETS[3], "04": SHEETS[4], "05": SHEETS[5], "06": SHEETS[6], "07": SHEETS[7]}
    tabela(("Preciso calcular", "Qual calculadora utilizar", "Exemplo de aplicação"), linhas,
           [("C", destino[ln[1][:2]]) for ln in linhas])
    sec("6   COMO INFORMAR OS DADOS")
    for t in ["• Digite somente nos campos amarelos da coluna MEU CÁLCULO. As outras células estão protegidas para "
              "evitar que uma fórmula seja apagada sem querer.",
              "• Use vírgula para decimais (0,5), como no padrão brasileiro. Não digite a unidade junto do número: "
              "escreva 500, e não 500 m³.",
              "• Ao clicar em um campo amarelo aparece uma dica com o que deve ser digitado.",
              "• Se você digitar algo inválido (texto ou número negativo), a planilha avisa e não aceita.",
              "• Seus dados ficam guardados quando você troca para um exemplo. Ao voltar para Meu cálculo, eles "
              "continuam lá."]:
        par(t)
    sec("7   COMO ESCOLHER AS UNIDADES")
    par("Algumas ferramentas têm uma linha Unidade de ... com lista. Escolha a unidade em que o seu dado está e a "
        "calculadora converte tudo sozinha. Nas linhas sem lista, use exatamente a unidade indicada na coluna Unidade "
        "(por exemplo, m³/dia e mg/L). A tabela abaixo ajuda nas conversões feitas à mão.")
    tabela(("Grandeza", "Relação", "Observação"), [
        ("Vazão", "1 L/s = 3,6 m³/h = 86,4 m³/dia", "1 m³/s = 1.000 L/s"),
        ("Volume", "1 m³ = 1.000 L", "1 L = 1.000 mL"),
        ("Tempo", "1 h = 60 min = 3.600 s", "1 dia = 24 h = 86.400 s"),
        ("Comprimento", "1 m = 100 cm = 1.000 mm", "1 polegada = 2,54 cm"),
        ("Concentração", "1 mg/L = 1 g/m³", "Em água diluída, 1 mg/L é aproximadamente 1 ppm"),
        ("Carga", "kg/dia = m³/dia × mg/L × 0,001", "1 kg = 1.000 g"),
        ("Porcentagem", "50% = 0,50", "Divida por 100 antes de usar na fórmula"),
        ("Densidade", "kg/L", "Água ≈ 1 kg/L. Produtos químicos variam: consulte a ficha técnica"),
    ])
    sec("8   COMO UTILIZAR OS EXEMPLOS PRONTOS")
    par("Cada ferramenta tem dois exemplos cadastrados (colunas Exemplo 01 e Exemplo 02). Ao escolher um deles na "
        "lista Cenário, a calculadora usa esses dados nas mesmas fórmulas do Meu cálculo. Use os exemplos para "
        "aprender o caminho da conta e depois troque para Meu cálculo para testar seus valores. A aba 08: Exemplos "
        "Resolvidos mostra todos os exemplos explicados passo a passo, com um botão que leva para a calculadora.")
    sec("9   COMO INTERPRETAR OS RESULTADOS")
    par("O card verde traz o nome do resultado, o valor em destaque e a unidade; outros resultados e unidades; uma "
        "interpretação em linguagem simples; e o status dos dados. Com ✔ o cálculo está completo. Com ⚠, leia a "
        "mensagem: ela diz qual campo falta ou está incorreto. Nesse caso o resultado fica com um traço (—), porque "
        "a calculadora não inventa valores.")
    par("Resultados negativos aparecem quando fazem sentido. Uma eficiência negativa indica que o parâmetro aumentou; "
        "um acúmulo negativo indica que a massa no sistema está diminuindo. A interpretação explica cada caso.")
    sec("10   COMO CONFERIR AS FÓRMULAS")
    for t in ["• A área E · Entenda o cálculo mostra a fórmula, cada variável, os valores usados e a substituição "
              "numérica. Tudo se atualiza quando você muda um dado.",
              "• Clique em qualquer resultado e veja a fórmula na barra de fórmulas.",
              "• As conversões de unidade e as validações ficam na ÁREA TÉCNICA, à direita de cada calculadora "
              "(colunas K e L). Basta rolar a tela para a direita.",
              "• Os exemplos resolvidos podem ser refeitos à mão ou na calculadora do celular para conferir."]:
        par(t)
    sec("11   O SIGNIFICADO DAS CORES")
    tabela(("Cor", "Significado", "Onde aparece"), [
        ("Amarelo-claro", "Campo editável: digite ou escolha aqui", "Coluna MEU CÁLCULO e listas de escolha"),
        ("Verde", "Resultado automático e indicadores positivos", "Card de resultado e status ✔"),
        ("Azul-claro", "Valor em uso, que entrou na conta", "Coluna Valor em uso"),
        ("Azul", "Botões de navegação", "Topo das abas e botões dos cards"),
        ("Azul-marinho", "Títulos e cabeçalhos", "Faixas de título"),
        ("Cinza-claro", "Explicações, exemplos cadastrados e apoio", "Colunas Exemplo 01 e 02 e textos"),
        ("Vermelho suave", "Avisos, erros comuns e dados inválidos", "Status ⚠ e linha Erro comum"),
    ], cores=[(AMARELO, TXT), (VERDE_OK, VERDE_T), (AZUL_L, NAVY), (AZUL, BRANCO), (NAVY, BRANCO),
              (CINZA_L, CINZA_T), (VERM_L, VERM_T)])
    sec("12   CUIDADOS ANTES DE UTILIZAR UM RESULTADO")
    for t in ["• Confira se as unidades dos seus dados são as mesmas pedidas na calculadora.",
              "• Verifique as hipóteses de cada fórmula (por exemplo, mistura completa e sem reação no balanço de massa).",
              "• Os resultados são cálculos teóricos de apoio ao estudo. Projetos, laudos e decisões operacionais "
              "exigem normas técnicas, legislação aplicável e responsável técnico habilitado.",
              "• A calculadora não diz se um valor atende a uma norma ou padrão de lançamento. Compare com o padrão "
              "aplicável ao seu caso.",
              "• Arredonde o resultado final de acordo com o que o exercício ou o professor pedir.",
              "• Na dosagem química, considere ficha técnica, compatibilidade e segurança do produto."]:
        par(t, fill="FFF4E5")
    sec("13   GLOSSÁRIO DE SÍMBOLOS E UNIDADES")
    tabela(("Símbolo", "Significado", "Unidade usual"), [
        ("Q", "Vazão", "m³/s, L/s, m³/h, m³/dia"), ("V", "Volume", "m³, L, mL"), ("t", "Tempo", "s, min, h, dia"),
        ("D", "Diâmetro interno (na calculadora 07, D é a dosagem)", "m, mm  |  mg/L"),
        ("A", "Área da seção do tubo", "m²"), ("v", "Velocidade média do escoamento", "m/s"),
        ("TDH", "Tempo de detenção hidráulica", "h, min"), ("C", "Concentração", "mg/L"),
        ("DBO", "Demanda bioquímica de oxigênio", "mg/L"), ("DQO", "Demanda química de oxigênio", "mg/L"),
        ("SST", "Sólidos suspensos totais", "mg/L"), ("N, P", "Nitrogênio e fósforo", "mg/L"),
        ("Carga", "Massa de poluente por tempo", "kg/dia"), ("E", "Eficiência de remoção", "%"),
        ("Qe, Qa, Ql", "Vazões de entrada, clarificado e lodo", "m³/dia"),
        ("C₁, V₁", "Concentração e volume da solução estoque", "mg/L, mL"),
        ("C₂, V₂", "Concentração e volume da solução final", "mg/L, mL"),
        ("ρ", "Densidade do produto", "kg/L"), ("ETE / ETA", "Estação de tratamento de esgoto / de água", "-"),
    ])
    sec("14   PROTEÇÃO, CÓPIAS E COMPATIBILIDADE")
    for t in ["• Mantenha uma cópia original deste arquivo sem alterações. Para cada trabalho, salve uma cópia com "
              "outro nome (exemplo: Lista3_Hidraulica).",
              "• No Excel, as abas estão protegidas sem senha, só para evitar alterações acidentais. Para editar, use "
              "Revisão > Desproteger Planilha (não pede senha).",
              "• Google Planilhas: envie o arquivo .xlsx para o Google Drive, clique com o botão direito > Abrir com > "
              "Planilhas Google. Depois use Arquivo > Salvar como Planilhas Google. Os cálculos, listas, cores e "
              "gráficos funcionam. Se um botão de navegação não abrir a aba, use as abas na parte de baixo da tela.",
              "• No Google Planilhas, quem é dono do arquivo consegue editar qualquer célula. Para se proteger de "
              "apagar uma fórmula sem querer, use Dados > Proteger páginas e intervalos > Mostrar um aviso ao editar.",
              "• O arquivo não usa macros nem internet. LibreOffice Calc também abre e calcula; a aparência dos "
              "gráficos pode variar um pouco."]:
        par(t)
    espaco(ws, r, 12)
    r += 1
    botao(ws, f"B{r}:D{r}", "Tudo pronto?  Voltar ao Painel Inicial e escolher uma calculadora  ▶", PAINEL, cor=VERDE_D,
          size=11)
    ws.row_dimensions[r].height = 30
    ws.print_area = f"A1:E{r}"
    proteger(ws)
    return reg_passos


# ====================================================================== biblioteca
def biblioteca(wb, alturas, dyn, todos):
    ws = wb.create_sheet(BIBLIO)
    configurar_aba(ws, "BF9000")
    colunas(ws, WC)
    cx = Ctx(ws, alturas, dyn)
    cabecalho(cx, "📚  BIBLIOTECA DE EXEMPLOS RESOLVIDOS", "Estude o passo a passo e depois teste outros valores na "
                                                          "calculadora", "BF9000", 8)
    r = 5
    txt = ("Como estudar: leia a situação, tente resolver sozinho, confira a resolução e depois clique em Abrir na "
           "calculadora para testar outros valores. Todos os números desta aba são calculados por fórmulas a partir "
           "dos mesmos dados cadastrados nas calculadoras.")
    box(ws, f"B{r}:H{r}", txt, F(10, True, NAVY), FL(AZUL_L), AL_L)
    cx.altura(r, altura_texto(txt, LBH, 10, True))
    r += 2
    secao(ws, r, "ÍNDICE POR CATEGORIA (clique para ir)", NAVY, 11, 24)
    r += 1
    idx_cells = {}
    for C in CALCS:
        n_ex = 2 * len(C["tools"])
        c = box(ws, f"B{r}", f"{C['emoji']}  {C['n']:02d} · {C['curto']}", F(10, True, BRANCO), FL(ACC[C["n"]]), AL_L,
                BD(BRANCO))
        idx_cells[C["n"]] = c
        box(ws, f"C{r}:H{r}", f"{n_ex} exemplos: " + "; ".join(T["title"] for T in C["tools"]), F(9, color=CINZA_T),
            FL(CINZA_XL), AL_L, BD())
        cx.altura(r, 24)
        r += 1
    r += 1
    total = 0
    for C in CALCS:
        acc = ACC[C["n"]]
        calc_nome = SHEETS[C["n"] + 1]
        infos = todos[C["n"]]
        secao(ws, r, f"{C['emoji']}  CALCULADORA {C['n']:02d} · {C['nome'].upper()}", acc, 14, 32, f"B{r}:G{r}")
        c = box(ws, f"H{r}", "▲ Voltar ao índice", F(9, True, BRANCO, underline="single"), FL(acc), AL_R)
        link(c, BIBLIO, "A1")
        link(idx_cells[C["n"]], BIBLIO, f"B{r}")
        r += 1
        espaco(ws, r, 8)
        r += 1
        cont = 0
        for ti, T in enumerate(C["tools"]):
            info = infos[ti]
            for e in range(2):
                cont += 1
                total += 1
                mp = {k: f"{q(calc_nome)}!{info['ex'][k][e]}" for k in info["ex"]}
                mp["ok"] = "TRUE"
                mp["idx"] = str(e + 1)
                cx.prefix = f"[Ex {C['n']}.{cont}] "
                for k, tpl in T["calcs"]:
                    mp[k] = cx.helper(k, "=IFERROR(" + resolve(tpl, mp) + ',"")')
                titulo, sit = T["ex_info"][e]
                box(ws, f"B{r}:F{r}", f"EXEMPLO {C['n']}.{cont}   ·   {titulo}", F(11, True, NAVY),
                    FL(tint(acc, 0.78)), AL_L)
                b = box(ws, f"G{r}:H{r}", "▶ Abrir na calculadora", F(10, True, BRANCO), FL(NAVY), AL_C, BD(BRANCO))
                link(b, calc_nome, f"B{info['title_row']}")
                ws.row_dimensions[r].height = 26
                r += 1
                linha_texto(cx, r, "Situação prática", sit)
                r += 1
                linha_texto(cx, r, "Dados utilizados", "=" + cx.res(valores_tpl(T), mp), dyn=True, padrao=22)
                r += 1
                form = [c_ for rot, tp, c_ in T["entenda"] if tp == "formula"]
                ftxt = "\n".join(form)
                linha_texto(cx, r, "Fórmula", ftxt, fill=AZUL_XL, font=F(11, True, NAVY),
                            h=max(24, altura_texto(ftxt, LCH, 11, True)))
                r += 1
                subst = [c_ for rot, tp, c_ in T["entenda"] if tp == "dyn" and rot.startswith("Substitui")]
                stpl = '&CHAR(10)&'.join(subst)
                linha_texto(cx, r, "Substituição", "=" + cx.res(stpl, mp), dyn=True)
                r += 1
                linha_texto(cx, r, "Resolução", "=" + cx.res(T["resol"], mp), dyn=True, padrao=45)
                r += 1
                linha_texto(cx, r, "Resultado", "=" + cx.res(T["result"], mp), fill=VERDE_L, rfill=VERDE_OK,
                            font=F(11, True, VERDE_T), dyn=True, padrao=24)
                r += 1
                linha_texto(cx, r, "Explicação simples", "=" + cx.res(T["interp"], mp), dyn=True, padrao=40)
                r += 1
                como = (f"Na aba {calc_nome}" + (f", vá até a Ferramenta {ti + 1}" if len(C["tools"]) > 1 else "") +
                        f" e escolha Exemplo 0{e + 1} na lista Cenário. Depois troque para Meu cálculo e teste "
                        "outros valores.")
                linha_texto(cx, r, "Teste na calculadora", como, fill=AZUL_XL, rfill=AZUL_L)
                r += 1
                espaco(ws, r, 14)
                r += 1
        r += 1
    cx.prefix = ""
    ws.print_area = f"A1:I{r}"
    proteger(ws)
    return total


# ====================================================================== principal
def gerar(saida, alturas=None, reg_out=None):
    alturas = alturas or {}
    dyn = []
    wb = Workbook()
    painel(wb, alturas, dyn)
    linha_passos = guia(wb, alturas, dyn)
    link(wb[PAINEL]["F6"], GUIA, f"B{linha_passos - 1}")
    reg = {"calcs": {}}
    todos = {}
    for C in CALCS:
        creg = {}
        ws, infos = calculadora(wb, C, alturas, dyn, creg)
        todos[C["n"]] = infos
        reg["calcs"][C["n"]] = creg
    reg["n_exemplos"] = biblioteca(wb, alturas, dyn, todos)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)
    reg["dyn"] = dyn
    if reg_out:
        with open(reg_out, "w", encoding="utf-8") as fh:
            json.dump(reg, fh, ensure_ascii=False, indent=1)
    return reg


if __name__ == "__main__":
    alt = json.load(open(sys.argv[2], encoding="utf-8")) if len(sys.argv) > 2 and sys.argv[2] != "-" else {}
    gerar(sys.argv[1], alt, sys.argv[3] if len(sys.argv) > 3 else None)
    print("ok")
