# -*- coding: utf-8 -*-
"""Pipeline completo: gera, audita, ajusta alturas, gera de novo, audita e grava valores calculados no arquivo final.

Uso: python finalizar.py pasta_temp arquivo_final.xlsx
"""
import os
import shutil
import subprocess
import sys
import zipfile

from lxml import etree
from openpyxl import load_workbook

from gerar_kit import gerar
from lo_recalc import main as recalc

TMP, FINAL = sys.argv[1], sys.argv[2]
os.makedirs(TMP, exist_ok=True)
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def auditar(xlsx, reg, alt_out=None):
    cmd = [sys.executable, "testar_kit.py", xlsx, reg, os.path.join(TMP, "casos")] + ([alt_out] if alt_out else [])
    r = subprocess.run(cmd, capture_output=True, text=True)
    print("\n".join(l for l in r.stdout.splitlines() if not l.endswith(" ok")))
    if r.returncode != 0:
        print(r.stderr[-3000:])
        sys.exit("Auditoria falhou")


def gravar_valores(destino, calculado):
    """Copia os valores calculados (cache) para as células com fórmula do arquivo gerado pelo openpyxl."""
    wb = load_workbook(calculado, data_only=True)
    nomes = wb.sheetnames
    tmp = destino + ".tmp"
    with zipfile.ZipFile(destino) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                i = int(item.filename[len("xl/worksheets/sheet"):-4]) - 1
                ws = wb[nomes[i]]
                root = etree.fromstring(data)
                for c in root.iter(f"{{{NS}}}c"):
                    f = c.find(f"{{{NS}}}f")
                    if f is None:
                        continue
                    v = c.find(f"{{{NS}}}v")
                    if v is None:
                        v = etree.SubElement(c, f"{{{NS}}}v")
                    val = ws[c.get("r")].value
                    if isinstance(val, bool):
                        c.set("t", "b")
                        v.text = "1" if val else "0"
                    elif isinstance(val, (int, float)):
                        c.attrib.pop("t", None)
                        v.text = repr(float(val)) if isinstance(val, float) else str(val)
                    else:
                        c.set("t", "str")
                        v.text = "" if val is None else str(val)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    os.replace(tmp, destino)


p1, r1, alt = os.path.join(TMP, "p1.xlsx"), os.path.join(TMP, "r1.json"), os.path.join(TMP, "alturas.json")
gerar(p1, {}, r1)
print("== Auditoria (passada 1)")
auditar(p1, r1, alt)
import json
alturas = json.load(open(alt))
p2, r2 = os.path.join(TMP, "p2.xlsx"), os.path.join(TMP, "r2.json")
gerar(p2, alturas, r2)
print("== Auditoria (passada 2, arquivo final)")
auditar(p2, r2)
calc = os.path.join(TMP, "p2_calc.xlsx")
shutil.copy(p2, calc)
recalc([calc])
shutil.copy(p2, FINAL)
gravar_valores(FINAL, calc)
shutil.copy(r2, os.path.join(os.path.dirname(os.path.abspath(FINAL)), "registro_celulas.json"))
print("Arquivo final:", FINAL)
