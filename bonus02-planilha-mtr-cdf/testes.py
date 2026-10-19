# -*- coding: utf-8 -*-
"""
Testes automáticos da planilha (cenários 01 a 12 do briefing).
Cada cenário gera um arquivo com dados fictícios, recalcula no LibreOffice e confere os resultados.

Uso: python testes.py <pasta_de_trabalho>
"""
import os
import sys
import json
import datetime as dt
from concurrent.futures import ProcessPoolExecutor

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gerar_planilha import build, SHEET, LET, FIRST, CAP, PEN_FIRST  # noqa: E402

RECALC_DIR = os.environ.get(
    "RECALC_DIR",
    "/root/.claude/skills/synced/db866081-af29-49f9-8a50-2c9b1be0228c_e98eb667-8c8e-48d5-ac14-4f0d54e4e53c/xlsx/scripts")
sys.path.insert(0, RECALC_DIR)

D = dt.date
UA, UB = "Unidade Teste A", "Unidade Teste B"
T1 = "Transportadora Teste"
D1, D2 = "Destinadora Teste 1", "Destinadora Teste 2"
RA, RB = "Resíduo Teste A", "Resíduo Teste B"
SIS = "SINIR / MTR Nacional"


def base():
    return dict(
        UN=[dict(nome=UA, razao="Empresa Teste"), dict(nome=UB, razao="Empresa Teste")],
        PA=[dict(fantasia=T1, tipo="Transportador"), dict(fantasia=D1, tipo="Destinador"),
            dict(fantasia=D2, tipo="Destinador")],
        RS=[dict(nome=RA, cod="X-A"), dict(nome=RB, cod="X-B")],
        MTR=[], IT=[], REC=[], CDF=[], VIN=[])


def mtr(num, d=D(2026, 9, 1), un=UA, dest=D1, **kw):
    r = dict(num=num, sis=SIS, uf="SP", dtem=d, unid=un, transp=T1, dest=dest)
    r.update(kw)
    return r


def rec(item, q, u="kg", conf="Sim", evid="Consulta ao sistema", d=D(2026, 9, 5)):
    return dict(itemref=item, dt=d, qty=q, un=u, conf=conf, evid=evid)


SCENARIOS = {}


def scenario(fn):
    SCENARIOS[fn.__name__] = fn
    return fn


@scenario
def t01_vazia():
    return {}


@scenario
def t02_mtr_simples():
    b = base()
    b["MTR"] = [mtr("0001")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=500, un="kg")]
    return b


@scenario
def t03_varios_residuos():
    b = base()
    b["MTR"] = [mtr("0001")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=800, un="kg"),
               dict(mtrref="MTR-000001", res=RB, qty=1.2, un="t")]
    return b


@scenario
def t04_recebimento_conferido():
    b = base()
    b["MTR"] = [mtr("0001"), mtr("0002")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=500, un="kg"),
               dict(mtrref="MTR-000002", res=RA, qty=300, un="kg")]
    # MTR 1: conferido com evidência. MTR 2: marcado Sim mas sem evidência (não deve confirmar)
    b["REC"] = [rec("IT-000001", 500), rec("IT-000002", 300, evid=None)]
    return b


@scenario
def t05_um_cdf_dois_mtrs():
    b = base()
    b["MTR"] = [mtr("0001"), mtr("0002")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=500, un="kg"),
               dict(mtrref="MTR-000002", res=RA, qty=300, un="kg")]
    b["REC"] = [rec("IT-000001", 500), rec("IT-000002", 300)]
    b["CDF"] = [dict(num="CDF-77", sis=SIS, dest=D1, loc="Sim", conf="Sim")]
    b["VIN"] = [dict(cdfref="CDF-000001", itemref="IT-000001", qty=500, un="kg", conf="Sim"),
                dict(cdfref="CDF-000001", itemref="IT-000002", qty=300, un="kg", conf="Sim")]
    return b


@scenario
def t06_cobertura_parcial():
    b = base()
    b["MTR"] = [mtr("0001")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=1000, un="kg")]
    b["REC"] = [rec("IT-000001", 1000)]
    b["CDF"] = [dict(num="CDF-88", sis=SIS, dest=D1, loc="Sim", conf="Sim")]
    b["VIN"] = [dict(cdfref="CDF-000001", itemref="IT-000001", qty=600, un="kg", conf="Sim")]
    return b


@scenario
def t07_cdf_nao_conferido():
    b = base()
    b["MTR"] = [mtr("0001")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=500, un="kg")]
    b["REC"] = [rec("IT-000001", 500)]
    b["CDF"] = [dict(num="CDF-99", sis=SIS, dest=D1, loc="Sim", conf="Não")]
    b["VIN"] = [dict(cdfref="CDF-000001", itemref="IT-000001", qty=500, un="kg", conf="Sim")]
    return b


@scenario
def t08_duplicidade():
    b = base()
    b["MTR"] = [mtr("123"), mtr("123"), mtr("123", sis="Sistema estadual"), mtr("123", un=UB)]
    b["IT"] = [dict(mtrref=f"MTR-00000{k}", res=RA, qty=100, un="kg") for k in range(1, 5)]
    return b


@scenario
def t09_cancelado():
    b = base()
    b["MTR"] = [mtr("0001", sitsis="Cancelado"), mtr("0002")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=500, un="kg"),
               dict(mtrref="MTR-000002", res=RA, qty=300, un="kg")]
    return b


@scenario
def t10_diferenca_pesagem():
    b = base()
    b["MTR"] = [mtr("0001"), mtr("0002")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=1000, un="kg"),
               dict(mtrref="MTR-000002", res=RA, qty=1000, un="kg", anal="Sim")]
    b["REC"] = [rec("IT-000001", 980), rec("IT-000002", 980)]
    return b


@scenario
def t11_referencia_inexistente():
    b = base()
    b["MTR"] = [mtr("0001")]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=500, un="kg"),
               dict(mtrref="MTR-000099", res=RA, qty=700, un="kg")]
    b["REC"] = [rec("IT-000001", 500), rec("IT-000777", 50)]
    b["CDF"] = [dict(num="CDF-11", sis=SIS, dest=D1, loc="Sim", conf="Sim")]
    b["VIN"] = [dict(cdfref="CDF-000001", itemref="IT-000500", qty=500, un="kg", conf="Sim"),
                dict(cdfref="CDF-000099", itemref="IT-000001", qty=500, un="kg", conf="Sim")]
    return b


def _t12_data(painel=None, rel=None):
    b = base()
    b["MTR"] = [mtr("0001", d=D(2026, 7, 10), un=UA, dest=D1),
                mtr("0002", d=D(2026, 8, 12), un=UA, dest=D2),
                mtr("0003", d=D(2026, 8, 20), un=UB, dest=D1),
                mtr("0004", d=D(2026, 9, 3), un=UB, dest=D2)]
    b["IT"] = [dict(mtrref="MTR-000001", res=RA, qty=100, un="kg"),
               dict(mtrref="MTR-000001", res=RB, qty=200, un="kg"),
               dict(mtrref="MTR-000002", res=RA, qty=1, un="t"),
               dict(mtrref="MTR-000003", res=RB, qty=400, un="kg"),
               dict(mtrref="MTR-000004", res=RA, qty=50, un="kg"),
               dict(mtrref="MTR-000004", res=RB, qty=60, un="kg"),
               dict(mtrref="MTR-000004", res=RB, qty=70, un="kg")]
    b["REC"] = [rec("IT-000004", 400, d=D(2026, 8, 25))]
    b["CDF"] = [dict(num="C-1", sis=SIS, dest=D1, dtem=D(2026, 9, 1), loc="Sim", conf="Não")]
    b["VIN"] = [dict(cdfref="CDF-000001", itemref="IT-000004", qty=400, un="kg", conf="Sim")]
    b["PAINEL"] = painel or {}
    b["REL"] = rel or {}
    return b


@scenario
def t12_filtros_sem_filtro():
    return _t12_data()


@scenario
def t12_filtros_unidade_periodo():
    return _t12_data(painel=dict(f_un=UB, f_ini=dt.datetime(2026, 8, 1), f_fim=dt.datetime(2026, 8, 31)),
                     rel=dict(r_un=UA))


@scenario
def t12_filtros_residuo():
    return _t12_data(painel=dict(f_res=RB))


# ---------------------------------------------------------------- execução
def run_one(args):
    name, folder = args
    from recalc import recalc
    data = SCENARIOS[name]()
    path = os.path.join(folder, f"{name}.xlsx")
    build(path, data)
    res = recalc(path, 600)
    return name, path, res


class Book:
    def __init__(self, path):
        self.wb = load_workbook(path, data_only=True)

    def v(self, sk, key, n=1):
        ws = self.wb[SHEET[sk]]
        return ws[f"{LET[sk][key]}{FIRST + n - 1}"].value

    def cell(self, sheet, ref):
        return self.wb[sheet][ref].value

    def painel(self):
        p = self.wb[SHEET["PAI"]]
        cards = [p[c].value for c in ("B11", "E11", "H11", "K11", "N11", "Q11")]
        sec = [p[c].value for c in ("B17", "E17", "H17", "K17", "N17", "Q17")]
        donut = {p[f"V{r}"].value: p[f"W{r}"].value for r in range(7, 15)}
        return cards, sec, donut

    def pend_list(self):
        ws = self.wb[SHEET["PEN"]]
        out = []
        r = PEN_FIRST
        while ws[f"B{r}"].value not in (None, ""):
            out.append((ws[f"B{r}"].value, ws[f"G{r}"].value, ws[f"J{r}"].value))
            r += 1
        return out

    def report_rows(self, num):
        ws = self.wb[SHEET["REL"]]
        start = None
        for r in range(1, ws.max_row + 1):
            v = ws[f"B{r}"].value
            if isinstance(v, str) and v.startswith(f"RELATÓRIO {num:02d}"):
                start = r
                break
        rows = []
        r = start + 3
        while True:
            vals = [ws.cell(row=r, column=c).value for c in range(2, 11)]
            if vals[0] in (None, "") or vals[0] == "TOTAL":
                break
            rows.append(vals)
            r += 1
        return rows


def check(results, cond, msg):
    results.append(("OK" if cond else "FALHOU", msg))
    return cond


def verify(name, path, rc):
    b = Book(path)
    res = []
    check(res, rc.get("status") == "success" and rc.get("total_errors") == 0,
          f"Recalculo sem erros de fórmula ({rc.get('total_formulas')} fórmulas, {rc.get('total_errors')} erros)")
    cards, sec, donut = b.painel()

    if name == "t01_vazia":
        check(res, cards == [0] * 6, f"Cards do painel zerados: {cards}")
        check(res, sec == [0] * 6, f"Indicadores complementares zerados: {sec}")
        check(res, sum(donut.values()) == 0, "Gráfico de situação sem dados")
        check(res, b.pend_list() == [], "Central de pendências vazia")
        check(res, b.cell(SHEET["PAI"], "B70") == "Nenhuma pendência no momento.", "Mensagem de ausência de pendências")
        orig = path.replace(".xlsx", "_original.xlsx")
        build(orig, {})
        wbf = load_workbook(orig)
        names = set(wbf.sheetnames)
        links = []
        for ws in wbf.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.hyperlink is not None:
                        links.append((ws.title, c.coordinate, c.hyperlink.location))
        bad = [l for l in links if not l[2] or l[2].split("!")[0].strip("'") not in names]
        check(res, len(links) >= 20 and not bad, f"{len(links)} links internos, todos apontando para abas existentes")
        inicio = [l for l in links if l[0] == SHEET["INI"]]
        check(res, len(inicio) == 11, f"Aba INÍCIO com 6 botões + 5 atalhos de passos ({len(inicio)} links)")
        check(res, len(wbf.sheetnames) == 13, f"13 abas: {wbf.sheetnames}")

    elif name == "t02_mtr_simples":
        check(res, cards[0] == 1, f"1 MTR contabilizado (card 01 = {cards[0]})")
        check(res, sec[0] == 500, f"Quantidade movimentada = 500 kg ({sec[0]})")
        check(res, b.v("MTR", "status") == "Aguardando recebimento", f"Situação: {b.v('MTR', 'status')}")
        check(res, cards[1] == 1 and cards[2] == 0, f"Card aguardando recebimento = 1, confirmado = 0 ({cards[1]}, {cards[2]})")
        check(res, b.pend_list()[0][1] == "RECEBIMENTO NÃO CONFIRMADO", f"Pendência: {b.pend_list()}")

    elif name == "t03_varios_residuos":
        check(res, cards[0] == 1, f"Um único MTR (card 01 = {cards[0]})")
        check(res, b.v("MTR", "nit") == 2, f"Dois itens ({b.v('MTR', 'nit')})")
        check(res, b.v("MTR", "kgtot") == 2000, f"Total do MTR = 2.000 kg ({b.v('MTR', 'kgtot')})")
        check(res, b.v("IT", "kg", 2) == 1200, f"1,2 t convertido para 1.200 kg ({b.v('IT', 'kg', 2)})")
        check(res, sec[0] == 2000, f"Painel: 2.000 kg movimentados ({sec[0]})")

    elif name == "t04_recebimento_conferido":
        check(res, b.v("MTR", "status") == "Aguardando CDF", f"MTR 1 com recebimento conferido: {b.v('MTR', 'status')}")
        check(res, b.v("MTR", "kgcdf") == 0, "CDF não considerado localizado (0 kg vinculado)")
        check(res, cards[2] == 1, f"Card recebimento confirmado = 1 ({cards[2]})")
        check(res, cards[3] == 1, f"Card aguardando CDF = 1 ({cards[3]})")
        check(res, b.v("MTR", "status", 2) == "Informação a revisar",
              f"MTR 2 marcado Sim sem evidência não confirma: {b.v('MTR', 'status', 2)}")
        check(res, b.v("REC", "alerta", 2) == "Conferido sem evidência identificada", f"Alerta: {b.v('REC', 'alerta', 2)}")

    elif name == "t05_um_cdf_dois_mtrs":
        check(res, sec[3] == 1, f"Um CDF cadastrado ({sec[3]})")
        check(res, b.v("CDF", "nmtr") == 2, f"Dois MTRs relacionados ao CDF ({b.v('CDF', 'nmtr')})")
        check(res, b.v("CDF", "kg") == 800, f"CDF com 800 kg vinculados ({b.v('CDF', 'kg')})")
        check(res, (b.v("MTR", "kgcdf"), b.v("MTR", "kgcdf", 2)) == (500, 300), "MTR A = 500 kg, MTR B = 300 kg")
        check(res, sec[2] == 800, f"Painel: 800 kg vinculados a CDF, sem duplicar ({sec[2]})")
        check(res, cards[4] == 2 and cards[0] == 2, f"2 MTRs com documentação conferida ({cards})")

    elif name == "t06_cobertura_parcial":
        check(res, b.v("IT", "kgcdf") == 600, f"Cobertura registrada 600 kg ({b.v('IT', 'kgcdf')})")
        check(res, b.v("IT", "semcdf") == 400, f"Diferença a investigar 400 kg ({b.v('IT', 'semcdf')})")
        check(res, abs(b.v("IT", "cob") - 0.6) < 1e-9, f"Cobertura 60% ({b.v('IT', 'cob')})")
        check(res, b.v("MTR", "status") == "Vínculo de CDF parcial", f"Situação: {b.v('MTR', 'status')}")
        check(res, cards[4] == 0, "Nenhuma conclusão integral (card documentação conferida = 0)")

    elif name == "t07_cdf_nao_conferido":
        check(res, b.v("CDF", "sit") == "Localizado, aguardando conferência", f"CDF: {b.v('CDF', 'sit')}")
        check(res, b.v("MTR", "status") == "CDF localizado, a conferir", f"MTR: {b.v('MTR', 'status')}")
        check(res, sec[4] == 0 and cards[4] == 0, "Nenhuma conferência concluída reconhecida")
        check(res, b.pend_list()[0][1] == "CDF A CONFERIR", f"Pendência: {b.pend_list()}")

    elif name == "t08_duplicidade":
        d = [b.v("MTR", "dupal", k) for k in range(1, 5)]
        check(res, d[0] == "POSSÍVEL DUPLICIDADE" and d[1] == "POSSÍVEL DUPLICIDADE",
              "Mesmo número, sistema, UF e unidade: alerta nos dois registros")
        check(res, d[2] in (None, "") and d[3] in (None, ""),
              "Mesmo número em outro sistema ou outra unidade: sem alerta")
        check(res, b.v("MTR", "status") == "Informação a revisar", f"Situação: {b.v('MTR', 'status')}")

    elif name == "t09_cancelado":
        check(res, b.v("MTR", "status") == "Cancelado (informado)", f"Situação: {b.v('MTR', 'status')}")
        check(res, b.v("MTR", "num") == "0001", "Histórico preservado (registro continua na aba)")
        check(res, cards[0] == 2 and cards[1] == 1, f"Card 01 inclui o cancelado (2); card 02 só o ativo (1): {cards}")
        check(res, sec[0] == 300, f"Quantidade movimentada exclui o cancelado: 300 kg ({sec[0]})")
        check(res, all(p[0] != "MTR-000001" for p in b.pend_list()), "Cancelado fora da central de pendências")
        check(res, donut.get("Cancelado (informado)") == 1, "Gráfico mostra o cancelado em fatia própria")

    elif name == "t10_diferenca_pesagem":
        check(res, b.v("IT", "dif") == 20, f"Diferença de 20 kg ({b.v('IT', 'dif')})")
        check(res, b.v("IT", "sit") == "Quantidade a revisar", f"Item: {b.v('IT', 'sit')}")
        check(res, b.v("MTR", "pend") == "QUANTIDADE A REVISAR", f"Pendência: {b.v('MTR', 'pend')}")
        check(res, b.v("MTR", "status") == "Informação a revisar", "Sinaliza conferência, sem falar em infração")
        check(res, b.v("MTR", "status", 2) == "Aguardando CDF",
              f"Com DIFERENÇA ANALISADA? = Sim o MTR segue o fluxo: {b.v('MTR', 'status', 2)}")

    elif name == "t11_referencia_inexistente":
        check(res, b.v("IT", "sit", 2) == "MTR inexistente", f"Item com MTR inexistente: {b.v('IT', 'sit', 2)}")
        check(res, b.v("MTR", "kgtot") == 500, f"Total do MTR válido não muda: {b.v('MTR', 'kgtot')}")
        check(res, b.v("REC", "alerta", 2) == "Item inexistente (referência inválida)", f"{b.v('REC', 'alerta', 2)}")
        check(res, b.v("VIN", "ali") == "REVISAR VÍNCULO: item inexistente", f"{b.v('VIN', 'ali')}")
        check(res, b.v("VIN", "ali", 2) == "REVISAR VÍNCULO: CDF inexistente", f"{b.v('VIN', 'ali', 2)}")
        check(res, b.v("CDF", "sit") == "Necessita revisão de vínculos", f"CDF: {b.v('CDF', 'sit')}")
        check(res, b.v("MTR", "kgcdf") == 500 and b.v("MTR", "status") == "Informação a revisar",
              f"MTR não aparece como concluído: {b.v('MTR', 'status')}")
        check(res, cards[4] == 0, "Nenhuma situação falsa de conclusão")
        txt = b.cell(SHEET["PEN"], "H5")
        check(res, "itens sem MTR válido 1" in txt and "recebimentos 1" in txt and "vínculos 2" in txt,
              f"Central mostra referências inválidas: {txt[:110]}")

    elif name == "t12_filtros_sem_filtro":
        check(res, cards[0] == 4, f"4 MTRs (MTR com 3 itens contado uma vez): {cards[0]}")
        check(res, sum(donut.values()) == 4, f"Gráfico de situação soma 4 ({donut})")
        check(res, sec[0] == 1880, f"Total 1.880 kg ({sec[0]})")
        r1 = b.report_rows(1)
        by = {r[0]: r for r in r1}
        check(res, by.get("jul/26", [0] * 4)[2] == 1 and by.get("ago/26", [0] * 4)[2] == 2 and by.get("set/26", [0] * 4)[2] == 1,
              "Relatório 01: jul=1, ago=2, set=1 MTRs")
        check(res, by.get("ago/26", [0] * 4)[3] == 1400, f"Relatório 01: ago = 1.400 kg ({by.get('ago/26')})")
        r2 = {r[0]: r for r in b.report_rows(2)}
        check(res, r2[RA][3] == 1150 and r2[RB][3] == 730, f"Relatório 02 por resíduo: A=1.150 kg, B=730 kg")
        r3 = b.report_rows(3)
        check(res, len(r3) == 1 and r3[0][0] == "MTR-000003", f"Relatório 03: 1 recebimento ({r3})")
        r4 = b.report_rows(4)
        check(res, len(r4) == 1 and r4[0][6] == "Localizado, aguardando conferência", f"Relatório 04: {r4}")
        r5 = b.report_rows(5)
        check(res, len(r5) == 4, f"Relatório 05: 4 pendências ({[x[2] for x in r5]})")
        r6 = b.report_rows(6)
        check(res, len(r6) == 4 and r6[3][2] == 3 and r6[3][3] == 180,
              f"Relatório 06: 4 MTRs, MTR-000004 com 3 itens e 180 kg")
        pa = b.wb[SHEET["PAI"]]
        mon = {pa[f"W{r}"].value: pa[f"X{r}"].value for r in range(18, 30)}
        check(res, mon.get("ago/26") == 1400 and mon.get("set/26") == 180, f"Gráfico mensal: {mon}")
        top = [(pa[f"W{r}"].value, pa[f"X{r}"].value) for r in range(33, 35)]
        check(res, top == [(RA, 1150), (RB, 730)], f"Gráfico resíduos: {top}")
        dst = [(pa[f"W{r}"].value, pa[f"X{r}"].value) for r in range(44, 46)]
        check(res, sorted(dst) == [(D1, 2), (D2, 2)], f"Gráfico destinadores: {dst}")

    elif name == "t12_filtros_unidade_periodo":
        check(res, cards[0] == 1, f"Filtro Unidade B + agosto: 1 MTR ({cards[0]})")
        check(res, sec[0] == 400, f"400 kg no filtro ({sec[0]})")
        r1 = {r[0]: r for r in b.report_rows(1)}
        check(res, r1.get("jul/26", [0] * 4)[2] == 1 and r1.get("ago/26", [0] * 4)[2] == 1 and r1.get("set/26", [0] * 4)[2] == 0,
              "Relatório 01 com unidade A: jul=1, ago=1, set=0")
        r6 = b.report_rows(6)
        check(res, [x[0] for x in r6] == ["MTR-000001", "MTR-000002"], f"Relatório 06 filtrado: {[x[0] for x in r6]}")

    elif name == "t12_filtros_residuo":
        check(res, cards[0] == 3, f"Filtro resíduo B: 3 MTRs ({cards[0]})")
        check(res, sec[0] == 730, f"Só a quantidade do resíduo B: 730 kg ({sec[0]})")
    return res


def main(folder):
    os.makedirs(folder, exist_ok=True)
    jobs = [(n, folder) for n in SCENARIOS]
    out = {}
    with ProcessPoolExecutor(max_workers=4) as ex:
        for name, path, rc in ex.map(run_one, jobs):
            out[name] = verify(name, path, rc)
    lines, total, falhas = [], 0, 0
    for name in SCENARIOS:
        lines.append(f"\n### {name}")
        for st, msg in out[name]:
            total += 1
            falhas += st != "OK"
            lines.append(f"- [{st}] {msg}")
    head = f"# Resultado dos testes\n\n{total - falhas} de {total} verificações aprovadas.\n"
    txt = head + "\n".join(lines) + "\n"
    print(txt)
    return falhas


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1] if len(sys.argv) > 1 else "testes_out") else 0)
