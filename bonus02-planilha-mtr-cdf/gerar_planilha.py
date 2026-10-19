# -*- coding: utf-8 -*-
"""
Gerador da PLANILHA PRONTA DE CONTROLE DE MTR E CDF (Bônus 02).

Uso:
    python gerar_planilha.py            -> gera a versão limpa e a demonstrativa
    build(caminho, dados)               -> usado pelos testes e pela versão demonstrativa

Sem macros. Fórmulas compatíveis com Excel 2019 / Excel 365.
"""
import re
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.workbook.properties import CalcProperties
from openpyxl.chart import DoughnutChart, BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.pagebreak import Break

VERSAO = "1.0 | outubro de 2026"
FONT = "Calibri"
FIRST = 6  # primeira linha de dados das abas de registro

CAP = dict(MTR=500, IT=1000, REC=1000, CDF=500, VIN=1000, UN=100, PA=200, RS=200)

SHEET = dict(
    INI="INÍCIO", PAI="PAINEL", MTR="MTRs", IT="ITENS_MTR", REC="RECEBIMENTOS",
    CDF="CDFs", VIN="VÍNCULOS_CDF", PEN="PENDÊNCIAS", REL="RELATÓRIOS",
    UN="CAD_UNIDADES", PA="CAD_PARCEIROS", RS="CAD_RESÍDUOS", CFG="CONFIGURAÇÕES",
)

# ---------------------------------------------------------------- cores
NAVY, BLUE, LBLUE = "082C56", "1766A9", "EAF4FC"
GREEN, LGREEN = "259B52", "E8F7EE"
YELLOW, LRED, GREY, WHITE = "F6C453", "FBE8E7", "F4F7FA", "FFFFFF"
ORANGE, LORANGE, LYELLOW = "E8892B", "FDEBD3", "FEF4D8"
RED_TXT, DGREY, LINE = "B3261E", "5B6878", "D5DEE8"
BAND_IN, BAND_CALC = "DCEBF8", "E9EEF3"
TAB_CAD = "8A99AB"


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def font(size=11, bold=False, color="1F2933", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBORDER = Border()
UNLOCK = Protection(locked=False)
LOCK = Protection(locked=True)
A_WRAP = Alignment(wrap_text=True, vertical="center")
A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

FMT_DATE = "dd/mm/yyyy"
FMT_KG = "#,##0.00"
FMT_KGDIF = "#,##0.00;[Red]-#,##0.00"
FMT_INT = "0"
FMT_PCT = "0%"

# ---------------------------------------------------------------- textos fixos
SIT_INTERNAS = [
    ("Cadastro incompleto", "Faltam campos essenciais do MTR (número, sistema, data, unidade, transportador, destinatário) ou nenhum item de resíduo foi registrado."),
    ("Informação a revisar", "As regras internas encontraram algo a conferir: possível duplicidade, cadastro fora da lista, datas, item, recebimento ou vínculo com alerta, ou diferença de quantidade não analisada."),
    ("Aguardando recebimento", "Ainda não há recebimento informado e conferido para todos os itens do MTR."),
    ("Aguardando CDF", "Recebimento conferido, mas nenhuma quantidade foi vinculada a CDF."),
    ("Vínculo de CDF parcial", "Parte da quantidade recebida ainda não tem vínculo de CDF registrado."),
    ("CDF localizado, a conferir", "A quantidade tem vínculo de CDF, mas o certificado ou o vínculo ainda não foi marcado como conferido."),
    ("Documentação conferida", "Recebimento e CDF registrados, relacionados e conferidos internamente. Não significa conformidade legal garantida."),
    ("Cancelado (informado)", "O usuário informou a situação de cancelamento. O registro fica no histórico e sai dos indicadores operacionais."),
]
PEND_TIPOS = [
    ("MTR INCOMPLETO", "Média", "Completar os campos essenciais do MTR e incluir os itens de resíduo."),
    ("IDENTIFICAÇÃO A REVISAR", "Alta", "Conferir cadastros, datas, referências e possíveis duplicidades do registro."),
    ("QUANTIDADE A REVISAR", "Alta", "Comparar quantidades enviadas, recebidas e vinculadas com os documentos e registrar a análise no item."),
    ("RECEBIMENTO NÃO CONFIRMADO", "Média", "Consultar o sistema oficial e registrar ou conferir a evidência de recebimento."),
    ("CDF NÃO LOCALIZADO", "Média", "Verificar com o destinador a disponibilidade do CDF e registrar o vínculo."),
    ("VÍNCULO INCOMPLETO", "Média", "Localizar o CDF da quantidade restante ou registrar a análise da diferença."),
    ("CDF A CONFERIR", "Baixa", "Abrir o CDF, conferir os dados e marcar a conferência do certificado e do vínculo."),
]
UFS = ["AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS", "MG", "PA",
       "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC", "SP", "SE", "TO"]

DISCLAIMER = ("Esta ferramenta organiza registros internos e não substitui os documentos "
              "oficiais nem os procedimentos exigíveis.")

# ---------------------------------------------------------------- colunas das abas de registro
# (chave, cabeçalho, tipo, largura, formato, fórmula, validação)
# tipo: id | in (preencher) | calc (automático visível) | h (auxiliar oculta)
# Na fórmula: {chave} = célula da mesma linha | {chave:top} = intervalo $X$6:Xn |
# {chave:all} = coluna inteira da tabela | [ABA.chave] = coluna de outra aba | {i} = nº da linha na tabela


def nest(pairs, default):
    s = default
    for cond, val in reversed(pairs):
        s = f"IF({cond},{val},{s})"
    return s


def parse_id(ref):
    return f'=IF({{{ref}}}="","",TRIM(LEFT({{{ref}}},FIND(" |",{{{ref}}}&" |")-1)))'


def fac_formula(un):
    m = f"MATCH({{{un}}},cfg_units,0)"
    return (f'=IF({{{un}}}="","",IFERROR(IF(ISNUMBER(INDEX(cfg_fac,{m})),'
            f'IF(INDEX(cfg_fac,{m})>0,INDEX(cfg_fac,{m}),""),""),""))')


def kg_formula(act, qty, fac):
    return (f'=IF({{{act}}}=0,"",IF(NOT(ISNUMBER({{{qty}}})),"",'
            f'IF(ISNUMBER({{{fac}}}),{{{qty}}}*{{{fac}}},"UNIDADE A CONFERIR")))')


DV_DATE = ("date", None)
DV_QTY = ("qty", None)

MTR_COLS = [
    ("id", "ID interno do MTR", "id", 13, None, None, None),
    ("num", "Número oficial do MTR", "in", 18, "@", None, None),
    ("sis", "Sistema de origem", "in", 22, None, None, ("list", "lst_sistemas")),
    ("uf", "UF do sistema", "in", 8, None, None, ("list", "lst_uf")),
    ("dtem", "Data de emissão", "in", 12, FMT_DATE, None, DV_DATE),
    ("dtsai", "Data de saída", "in", 12, FMT_DATE, None, DV_DATE),
    ("unid", "Unidade geradora", "in", 26, None, None, ("list", "lst_unidades")),
    ("transp", "Transportador", "in", 28, None, None, ("list", "lst_transp")),
    ("dest", "Unidade destinatária", "in", 28, None, None, ("list", "lst_dest")),
    ("ufdest", "UF de destino", "in", 8, None, None, ("list", "lst_uf")),
    ("sitsis", "Situação informada no sistema", "in", 18, None, None, ("list", "lst_sit_mtr")),
    ("dtcons", "Data da última consulta", "in", 12, FMT_DATE, None, DV_DATE),
    ("link", "Link ou caminho do arquivo MTR", "in", 30, None, None, None),
    ("resp", "Responsável pelo lançamento", "in", 18, None, None, None),
    ("obs", "Observações", "in", 30, None, None, None),
    ("kgtot", "Quantidade total vinculada ao MTR (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([IT.kg],[IT.mtrid],{id},[IT.qok],1))', None),
    ("kgrec", "Quantidade recebida registrada (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([REC.kg],[REC.mtr],{id},[REC.qok],1))', None),
    ("kgcdf", "Quantidade vinculada a CDF (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([VIN.kg],[VIN.mtr],{id},[VIN.qok],1))', None),
    ("status", "Situação interna de acompanhamento", "calc", 26, None, "=" + nest([
        ("{act}=0", '""'),
        ("{canc}=1", '"Cancelado (informado)"'),
        ("{incomp}=1", '"Cadastro incompleto"'),
        ("OR({dup}=1,{ident}=1,{nitstruct}>0,{nitqty}>0,{nrecissue}>0,{nvinissue}>0)", '"Informação a revisar"'),
        ("{recconf}=0", '"Aguardando recebimento"'),
        ("N({kgcdf})=0", '"Aguardando CDF"'),
        ("{nitcov}<N({nit})", '"Vínculo de CDF parcial"'),
        ("{cdfdone}=0", '"CDF localizado, a conferir"'),
    ], '"Documentação conferida"'), None),
    ("pend", "Pendência principal", "calc", 26, None, "=" + nest([
        ("{act}=0", '""'),
        ('{status}="Cadastro incompleto"', '"MTR INCOMPLETO"'),
        ('{status}="Informação a revisar"',
         'IF(AND({nitqty}>0,{dup}=0,{ident}=0,{nitstruct}=0,{nrecissue}=0,{nvinissue}=0),'
         '"QUANTIDADE A REVISAR","IDENTIFICAÇÃO A REVISAR")'),
        ('{status}="Aguardando recebimento"', '"RECEBIMENTO NÃO CONFIRMADO"'),
        ('{status}="Aguardando CDF"', '"CDF NÃO LOCALIZADO"'),
        ('{status}="Vínculo de CDF parcial"', '"VÍNCULO INCOMPLETO"'),
        ('{status}="CDF localizado, a conferir"', '"CDF A CONFERIR"'),
    ], '""'), None),
    ("dias", "Dias em acompanhamento", "calc", 15, FMT_INT,
     '=IF({pend}="","",IF({ref}="","Não calculável",IF({ref}>TODAY(),"Data futura",TODAY()-{ref})))', None),
    ("kgrecconf", "Quantidade recebida conferida (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([REC.kg],[REC.mtr],{id},[REC.confok],1))', None),
    ("nit", "Nº de itens", "calc", 11, FMT_INT,
     '=IF({act}=0,"",COUNTIFS([IT.mtrid],{id},[IT.act],1))', None),
    ("dupal", "Alerta de duplicidade", "calc", 16, None, '=IF({dup}=1,"POSSÍVEL DUPLICIDADE","")', None),
    ("todas", "Todas as pendências identificadas", "calc", 42, None,
     '=IF({act}=0,"",IF({todasraw}="","Nenhuma",LEFT({todasraw},LEN({todasraw})-2)))', None),
    # ---- auxiliares
    ("act", "aux registro ativo", "h", 6, None, "=IF(COUNTA({num}:{obs})>0,1,0)", None),
    ("canc", "aux cancelado", "h", 6, None,
     '=IF(AND({act}=1,{sitsis}<>""),IF(TRIM({sitsis})=TRIM(cfg_cancel),1,0),0)', None),
    ("dupkey", "aux chave duplicidade", "h", 6, None,
     '=IF(AND({act}=1,{num}<>""),{sis}&"|"&{uf}&"|"&{unid}&"|"&TRIM({num}),"")', None),
    ("dup", "aux duplicado", "h", 6, None,
     '=IF({dupkey}="",0,IF(SUMPRODUCT(--({dupkey:all}={dupkey}))>1,1,0))', None),
    ("ident", "aux identificação a revisar", "h", 6, None,
     '=IF({act}=0,0,IF(OR(AND({unid}<>"",ISNA(MATCH({unid},[UN.nome],0))),'
     'AND({transp}<>"",ISNA(MATCH({transp},cmp_transp,0))),'
     'AND({dest}<>"",ISNA(MATCH({dest},cmp_dest,0))),'
     'AND({dtem}<>"",NOT(ISNUMBER({dtem}))),AND({dtsai}<>"",NOT(ISNUMBER({dtsai}))),'
     'AND(ISNUMBER({dtem}),ISNUMBER({dtsai}),{dtsai}<{dtem})),1,0))', None),
    ("incomp", "aux incompleto", "h", 6, None,
     '=IF({act}=0,0,IF(OR({num}="",{sis}="",{dtem}="",{unid}="",{transp}="",{dest}="",N({nit})=0),1,0))', None),
    ("nitstruct", "aux itens com dado a revisar", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([IT.mtrid],{id},[IT.struct],1))", None),
    ("nitrecok", "aux itens recebidos conferidos", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([IT.mtrid],{id},[IT.recok],1))", None),
    ("nitqty", "aux itens com quantidade a revisar", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([IT.mtrid],{id},[IT.qtyiss],1))", None),
    ("nitcdfok", "aux itens com CDF ok", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([IT.mtrid],{id},[IT.cdfok],1))", None),
    ("nitcov", "aux itens com cobertura", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([IT.mtrid],{id},[IT.covok],1))", None),
    ("nrecissue", "aux recebimentos com alerta", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([REC.mtr],{id},[REC.issue],1))", None),
    ("nvinissue", "aux vínculos com alerta", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([VIN.mtr],{id},[VIN.issue],1))", None),
    ("nrec", "aux nº recebimentos", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([REC.mtr],{id},[REC.act],1))", None),
    ("recconf", "aux recebimento conferido", "h", 6, None,
     "=IF(AND({act}=1,N({nit})>0,{nitrecok}=N({nit})),1,0)", None),
    ("cdfdone", "aux cdf completo", "h", 6, None,
     "=IF(AND({act}=1,N({nit})>0,{nitcdfok}=N({nit})),1,0)", None),
    ("maxrec", "aux última data recebimento", "h", 6, FMT_DATE,
     "=IF({act}=0,0,_xlfn.MAXIFS([REC.dt],[REC.mtr],{id}))", None),
    ("ref", "aux data de referência", "h", 6, FMT_DATE,
     '=IF({pend}="","",IF({pend}="RECEBIMENTO NÃO CONFIRMADO",'
     'IF(ISNUMBER({dtsai}),{dtsai},IF(ISNUMBER({dtem}),{dtem},"")),'
     'IF(OR({pend}="CDF NÃO LOCALIZADO",{pend}="VÍNCULO INCOMPLETO",{pend}="CDF A CONFERIR"),'
     'IF({maxrec}>0,{maxrec},IF(ISNUMBER({dtem}),{dtem},"")),IF(ISNUMBER({dtem}),{dtem},""))))', None),
    ("overdue", "aux acima do prazo interno", "h", 6, None,
     '=IF(NOT(ISNUMBER({dias})),0,IF({pend}="RECEBIMENTO NÃO CONFIRMADO",'
     'IF(AND(ISNUMBER(cfg_rec),{dias}>cfg_rec),1,0),'
     'IF(OR({pend}="CDF NÃO LOCALIZADO",{pend}="VÍNCULO INCOMPLETO",{pend}="CDF A CONFERIR"),'
     'IF(AND(ISNUMBER(cfg_cdf),{dias}>cfg_cdf),1,0),0)))', None),
    ("prio", "aux prioridade", "h", 6, None,
     '=IF({pend}="","",IF({overdue}=1,"Alta",IFERROR(INDEX(cfg_pp,MATCH({pend},cfg_pt,0))&"","Média")))', None),
    ("prov", "aux providência", "h", 6, None,
     '=IF({pend}="","",IFERROR(INDEX(cfg_pv,MATCH({pend},cfg_pt,0))&"",""))', None),
    ("pkey", "aux ordem pendência", "h", 6, None,
     '=IF({pend}="","",IF({prio}="Alta",1,IF({prio}="Média",2,3))*1000000000'
     '+(999999-IF(ISNUMBER({dias}),MIN({dias},999999),0))*1000+{i})', None),
    ("haspend", "aux tem pendência", "h", 6, None, '=IF({pend}<>"",1,0)', None),
    ("todasraw", "aux texto pendências", "h", 6, None,
     '=IF({act}=0,"",IF({canc}=1,"Cancelamento informado; ","")'
     '&IF({incomp}=1,"Campos essenciais ou itens faltando; ","")'
     '&IF({dup}=1,"Possível duplicidade do número; ","")'
     '&IF({ident}=1,"Cadastro ou datas a revisar; ","")'
     '&IF({nitstruct}>0,"Item de resíduo com dado a revisar; ","")'
     '&IF({nitqty}>0,"Diferença de quantidade a analisar; ","")'
     '&IF({nrecissue}>0,"Recebimento com alerta; ","")'
     '&IF({nvinissue}>0,"Vínculo de CDF com alerta; ","")'
     '&IF(AND({canc}=0,{recconf}=0),"Recebimento não confirmado; ","")'
     '&IF(AND({canc}=0,{recconf}=1,{cdfdone}=0),"CDF pendente ou a conferir; ",""))', None),
    ("passb", "aux filtro painel sem resíduo", "h", 6, None,
     '=IF({act}=0,0,IF(AND(OR(f_ini="",AND(ISNUMBER({dtem}),{dtem}>=f_ini)),'
     'OR(f_fim="",AND(ISNUMBER({dtem}),{dtem}<=f_fim)),'
     'OR(f_un="",f_un="(Todas)",{unid}=f_un),OR(f_tr="",f_tr="(Todos)",{transp}=f_tr),'
     'OR(f_de="",f_de="(Todos)",{dest}=f_de),OR(f_sit="",f_sit="(Todas)",{status}=f_sit)),1,0))', None),
    ("pass", "aux filtro painel", "h", 6, None,
     '=IF({passb}=0,0,IF(OR(f_res="",f_res="(Todos)"),1,IF(COUNTIFS([IT.mtrid],{id},[IT.res],f_res)>0,1,0)))', None),
    ("mes", "aux mês emissão", "h", 6, FMT_DATE,
     '=IF(ISNUMBER({dtem}),DATE(YEAR({dtem}),MONTH({dtem}),1),"")', None),
    ("rpass", "aux filtro relatórios", "h", 6, None,
     '=IF({act}=0,0,IF(AND(OR(r_ini="",AND(ISNUMBER({dtem}),{dtem}>=r_ini)),'
     'OR(r_fim="",AND(ISNUMBER({dtem}),{dtem}<=r_fim)),'
     'OR(r_un="",r_un="(Todas)",{unid}=r_un),OR(r_de="",r_de="(Todos)",{dest}=r_de)),1,0))', None),
    ("rseq", "aux ordem relatório 06", "h", 6, None, '=IF({rpass}=1,COUNTIF({rpass:top},1),"")', None),
    ("r5key", "aux ordem relatório 05", "h", 6, None,
     '=IF(AND({pkey}<>"",OR(r_un="",r_un="(Todas)",{unid}=r_un),OR(r_de="",r_de="(Todos)",{dest}=r_de)),{pkey},"")', None),
    ("label", "aux rótulo seleção", "h", 6, None,
     '=IF({act}=1,{id}&" | nº "&IF({num}="","(sem número)",{num})&" | "&IF({unid}="","(sem unidade)",{unid}),"")', None),
    ("lseq", "aux ordem seleção", "h", 6, None, '=IF({act}=1,COUNTIF({act:top},1),"")', None),
    ("partic", "aux participante", "h", 6, None,
     '=IF({pend}="","",IF({pend}="RECEBIMENTO NÃO CONFIRMADO","Transp.: "&{transp}&" | Dest.: "&{dest},'
     'IF(OR({pend}="CDF NÃO LOCALIZADO",{pend}="VÍNCULO INCOMPLETO",{pend}="CDF A CONFERIR"),'
     '"Dest.: "&{dest},"Gerador: "&{unid})))', None),
]

IT_STATUS = "=" + nest([
    ("{act}=0", '""'),
    ('{mtrid}=""', '"Sem MTR associado"'),
    ("{mok}=0", '"MTR inexistente"'),
    ('{res}=""', '"Resíduo não identificado"'),
    ("{resok}=0", '"Resíduo fora do cadastro"'),
    ("{qtyok}=0", '"Quantidade inválida"'),
    ('{un}=""', '"Unidade não informada"'),
    ("NOT(ISNUMBER({kg}))", '"Unidade a conferir"'),
    ('AND({dup}=1,{anal}<>"Sim")', '"Possível duplicidade"'),
    ("{canc}=1", '"MTR cancelado"'),
    ("{nreciss}>0", '"Recebimento com alerta"'),
    ("{nrec}=0", '"Aguardando recebimento"'),
    ("{nrecok}<{nrec}", '"Recebimento a conferir"'),
    ("{qtyiss}=1", '"Quantidade a revisar"'),
    ("{nviniss}>0", '"Vínculo a revisar"'),
    ("{nvin}=0", '"Aguardando CDF"'),
    ("{covok}=0", '"Vínculo de CDF parcial"'),
    ("{nvinok}<{nvin}", '"CDF a conferir"'),
], '"Documentação conferida"')

IT_COLS = [
    ("id", "ID interno do item", "id", 12, None, None, None),
    ("mtrref", "ID interno do MTR", "in", 34, None, None, ("list", "lst_mtr")),
    ("res", "Resíduo cadastrado", "in", 28, None, None, ("list", "lst_res")),
    ("cod", "Código informado para o resíduo", "calc", 16, None,
     '=IF(OR({act}=0,{resok}=0),"",INDEX([RS.cod],MATCH({res},[RS.nome],0))&"")', None),
    ("desc", "Descrição do resíduo", "calc", 26, None,
     '=IF(OR({act}=0,{resok}=0),"",INDEX([RS.desc],MATCH({res},[RS.nome],0))&"")', None),
    ("cls", "Classificação informada", "calc", 18, None,
     '=IF(OR({act}=0,{resok}=0),"",INDEX([RS.cls],MATCH({res},[RS.nome],0))&"")', None),
    ("qty", "Quantidade enviada", "in", 12, "#,##0.000", None, DV_QTY),
    ("un", "Unidade de medida", "in", 9, None, None, ("list", "lst_un")),
    ("kg", "Quantidade convertida (kg)", "calc", 14, FMT_KG, None, None),
    ("obs", "Observações", "in", 28, None, None, None),
    ("anal", "Diferença analisada?", "in", 11, None, None, ("list", "lst_simnao")),
    ("kgrec", "Quantidade recebida vinculada (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([REC.kg],[REC.itemid],{id},[REC.qok],1))', None),
    ("kgcdf", "Quantidade com CDF vinculada (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([VIN.kg],[VIN.itemid],{id},[VIN.qok],1))', None),
    ("dif", "Diferença enviado − recebido (kg)", "calc", 14, FMT_KGDIF,
     '=IF(OR({act}=0,{nrec}=0,NOT(ISNUMBER({kg}))),"",{kg}-N({kgrec}))', None),
    ("cob", "Cobertura documental por CDF", "calc", 11, FMT_PCT,
     '=IF(OR({act}=0,{base}<=0),"",N({kgcdf})/{base})', None),
    ("sit", "Situação do item", "calc", 24, None, IT_STATUS, None),
    ("semcdf", "Recebido ainda sem vínculo de CDF (kg)", "calc", 14, FMT_KG,
     '=IF(OR({act}=0,{nrec}=0),"",MAX(0,N({kgrec})-N({kgcdf})))', None),
    # ---- auxiliares
    ("act", "aux registro ativo", "h", 6, None,
     "=IF(COUNTA({mtrref}:{res},{qty}:{un},{obs}:{anal})>0,1,0)", None),
    ("mtrid", "aux id mtr", "h", 6, None, parse_id("mtrref"), None),
    ("mrow", "aux linha mtr", "h", 6, None,
     '=IF({mtrid}="","",IFERROR(MATCH({mtrid},[MTR.id],0),""))', None),
    ("mok", "aux mtr válido", "h", 6, None,
     '=IF({mrow}="",0,IF(INDEX([MTR.act],{mrow})=1,1,0))', None),
    ("resok", "aux resíduo válido", "h", 6, None,
     '=IF({res}="",0,IF(ISNUMBER(MATCH({res},[RS.nome],0)),1,0))', None),
    ("qtyok", "aux quantidade positiva", "h", 6, None,
     "=IF(AND(ISNUMBER({qty}),{qty}>0),1,0)", None),
    ("fac", "aux fator kg", "h", 6, None, fac_formula("un"), None),
    ("qok", "aux quantidade válida", "h", 6, None, "=IF(AND({qtyok}=1,ISNUMBER({kg})),1,0)", None),
    ("canc", "aux mtr cancelado", "h", 6, None, "=IF({mok}=1,INDEX([MTR.canc],{mrow}),0)", None),
    ("nrec", "aux nº recebimentos", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([REC.itemid],{id},[REC.act],1))", None),
    ("nrecok", "aux recebimentos conferidos", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([REC.itemid],{id},[REC.confok],1))", None),
    ("nreciss", "aux recebimentos com alerta", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([REC.itemid],{id},[REC.issue],1))", None),
    ("nvin", "aux nº vínculos", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([VIN.itemid],{id},[VIN.act],1))", None),
    ("nvinok", "aux vínculos ok", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([VIN.itemid],{id},[VIN.ok],1))", None),
    ("nviniss", "aux vínculos com alerta", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([VIN.itemid],{id},[VIN.issue],1))", None),
    ("dkey", "aux chave duplicidade", "h", 6, None,
     '=IF(AND({act}=1,{mtrid}<>"",{res}<>""),{mtrid}&"|"&{res}&"|"&{qty}&"|"&{un},"")', None),
    ("dup", "aux duplicado", "h", 6, None,
     '=IF({dkey}="",0,IF(SUMPRODUCT(--({dkey:all}={dkey}))>1,1,0))', None),
    ("struct", "aux dado a revisar", "h", 6, None,
     '=IF({act}=0,0,IF(OR({mok}=0,{resok}=0,{qok}=0,AND({dup}=1,{anal}<>"Sim")),1,0))', None),
    ("recok", "aux recebimento conferido", "h", 6, None,
     "=IF(AND({act}=1,{nrec}>0,{nrecok}={nrec}),1,0)", None),
    ("base", "aux referência cobertura", "h", 6, None,
     "=IF({act}=0,0,IF({nrec}>0,N({kgrec}),IF(ISNUMBER({kg}),{kg},0)))", None),
    ("over", "aux vínculo acima", "h", 6, None,
     '=IF({act}=0,0,IF(AND({anal}<>"Sim",N({kgcdf})>{base}+cfg_tol),1,0))', None),
    ("qtyiss", "aux quantidade a revisar", "h", 6, None,
     '=IF({act}=0,0,IF({anal}="Sim",0,IF(OR(AND({nrec}>0,ISNUMBER({kg}),'
     'ABS({kg}-N({kgrec}))>cfg_tol),{over}=1),1,0)))', None),
    ("covok", "aux cobertura ok", "h", 6, None,
     '=IF(AND({act}=1,N({kgcdf})>0,OR(N({kgcdf})>={base}-cfg_tol,{anal}="Sim")),1,0)', None),
    ("cdfok", "aux cdf ok", "h", 6, None,
     "=IF(AND({covok}=1,{nvin}>0,{nvinok}={nvin},{over}=0),1,0)", None),
    ("pass", "aux filtro painel", "h", 6, None,
     '=IF({mok}=0,0,IF(AND(INDEX([MTR.passb],{mrow})=1,{canc}=0,'
     'OR(f_res="",f_res="(Todos)",{res}=f_res)),1,0))', None),
    ("mes", "aux mês emissão", "h", 6, FMT_DATE, '=IF({mok}=0,"",INDEX([MTR.mes],{mrow}))', None),
    ("rpass", "aux filtro relatórios", "h", 6, None,
     "=IF({mok}=0,0,IF(AND(INDEX([MTR.rpass],{mrow})=1,{canc}=0),1,0))", None),
    ("label", "aux rótulo seleção", "h", 6, None,
     '=IF({act}=1,{id}&" | MTR "&IF({mok}=1,INDEX([MTR.num],{mrow})&"","?")&" | "'
     '&IF({res}="","(sem resíduo)",{res}),"")', None),
    ("lseq", "aux ordem seleção", "h", 6, None, '=IF({act}=1,COUNTIF({act:top},1),"")', None),
    ("mdest", "aux destinatário do MTR", "h", 6, None,
     '=IF({mok}=1,INDEX([MTR.dest],{mrow})&"","")', None),
]
# fórmula do kg depende de chaves já definidas
IT_COLS = [c if c[0] != "kg" else c[:5] + (kg_formula("act", "qty", "fac"),) + c[6:] for c in IT_COLS]

REC_ALERT = "=" + nest([
    ("{act}=0", '""'),
    ('{itemref}=""', '"Item de resíduo não informado"'),
    ("{iok}=0", '"Item inexistente (referência inválida)"'),
    ("NOT(ISNUMBER({qty}))", '"Quantidade não informada"'),
    ("{qty}<=0", '"Quantidade não positiva"'),
    ('{un}=""', '"Unidade não informada"'),
    ("NOT(ISNUMBER({kg}))", '"Unidade a conferir"'),
    ("NOT(ISNUMBER({dt}))", '"Data de recebimento não informada"'),
    ("AND(ISNUMBER({memis}),N({memis})>0,{dt}<N({memis}))", '"Data anterior à emissão do MTR"'),
    ('AND({conf}="Sim",{evid}="",{link}="")', '"Conferido sem evidência identificada"'),
    ('{conf}<>"Sim"', '"Aguardando conferência"'),
    ("INDEX([IT.qtyiss],{irow})=1", '"Conferido (ver diferença no item)"'),
], '"Conferido"')

REC_COLS = [
    ("id", "ID interno do recebimento", "id", 13, None, None, None),
    ("mtr", "ID do MTR", "calc", 13, None, '=IF({iok}=1,INDEX([IT.mtrid],{irow})&"","")', None),
    ("itemref", "ID do item de resíduo", "in", 36, None, None, ("list", "lst_itens")),
    ("dest", "Unidade destinatária", "calc", 26, None,
     '=IF(N({mrow})>0,INDEX([MTR.dest],{mrow})&"","")', None),
    ("dt", "Data de recebimento informada", "in", 13, FMT_DATE, None, DV_DATE),
    ("qty", "Quantidade recebida", "in", 12, "#,##0.000", None, DV_QTY),
    ("un", "Unidade de medida", "in", 9, None, None, ("list", "lst_un")),
    ("kg", "Quantidade convertida (kg)", "calc", 14, FMT_KG, kg_formula("act", "qty", "fac"), None),
    ("sitsis", "Situação informada no sistema", "in", 18, None, None, ("list", "lst_sit_mtr")),
    ("conf", "Recebimento conferido?", "in", 12, None, None, ("list", "lst_simnao")),
    ("dtconf", "Data da conferência", "in", 12, FMT_DATE, None, DV_DATE),
    ("evid", "Documento ou evidência consultada", "in", 26, None, None, None),
    ("link", "Link ou caminho do registro", "in", 28, None, None, None),
    ("resp", "Responsável pela conferência", "in", 18, None, None, None),
    ("obs", "Observações", "in", 26, None, None, None),
    ("alerta", "Alerta do recebimento", "calc", 30, None, REC_ALERT, None),
    # ---- auxiliares
    ("act", "aux registro ativo", "h", 6, None,
     "=IF(COUNTA({itemref},{dt}:{un},{sitsis}:{obs})>0,1,0)", None),
    ("itemid", "aux id item", "h", 6, None, parse_id("itemref"), None),
    ("irow", "aux linha item", "h", 6, None,
     '=IF({itemid}="","",IFERROR(MATCH({itemid},[IT.id],0),""))', None),
    ("iok", "aux item válido", "h", 6, None,
     '=IF({irow}="",0,IF(INDEX([IT.act],{irow})=1,1,0))', None),
    ("mrow", "aux linha mtr", "h", 6, None, '=IF({iok}=1,INDEX([IT.mrow],{irow}),"")', None),
    ("memis", "aux emissão do mtr", "h", 6, FMT_DATE,
     '=IF(N({mrow})>0,INDEX([MTR.dtem],{mrow}),"")', None),
    ("fac", "aux fator kg", "h", 6, None, fac_formula("un"), None),
    ("qok", "aux quantidade válida", "h", 6, None,
     "=IF(AND(ISNUMBER({qty}),ISNUMBER({kg})),IF({kg}>0,1,0),0)", None),
    ("issue", "aux alerta", "h", 6, None,
     '=IF({act}=0,0,IF(OR({iok}=0,{qok}=0,NOT(ISNUMBER({dt})),AND({conf}="Sim",{evid}="",{link}=""),'
     'AND(ISNUMBER({dt}),ISNUMBER({memis}),N({memis})>0,{dt}<N({memis}))),1,0))', None),
    ("confok", "aux conferido válido", "h", 6, None,
     '=IF(AND({act}=1,{iok}=1,{qok}=1,{conf}="Sim",OR({evid}<>"",{link}<>"")),1,0)', None),
    ("rpass", "aux filtro relatórios", "h", 6, None,
     '=IF({act}=0,0,IF(AND(OR(r_ini="",AND(ISNUMBER({dt}),{dt}>=r_ini)),'
     'OR(r_fim="",AND(ISNUMBER({dt}),{dt}<=r_fim)),'
     'OR(r_un="",r_un="(Todas)",IF(N({mrow})>0,INDEX([MTR.unid],{mrow}),"")=r_un),'
     'OR(r_de="",r_de="(Todos)",{dest}=r_de)),1,0))', None),
    ("rseq", "aux ordem relatório 03", "h", 6, None, '=IF({rpass}=1,COUNTIF({rpass:top},1),"")', None),
]

CDF_SIT = "=" + nest([
    ("{act}=0", '""'),
    ('OR({num}="",{dest}="")', '"Informações incompletas"'),
    ("{dup}=1", '"Possível duplicidade"'),
    ("{destok}=0", '"Destinador fora do cadastro"'),
    ('AND({conf}="Sim",N({nit})=0)', '"Necessita revisão de vínculos"'),
    ("{nviniss}>0", '"Necessita revisão de vínculos"'),
    ('{conf}="Sim"', '"Conferido"'),
    ('{loc}="Sim"', '"Localizado, aguardando conferência"'),
], '"Cadastrado"')

CDF_COLS = [
    ("id", "ID interno do CDF", "id", 13, None, None, None),
    ("num", "Número oficial do CDF", "in", 18, "@", None, None),
    ("sis", "Sistema de origem", "in", 22, None, None, ("list", "lst_sistemas")),
    ("dtem", "Data de emissão", "in", 12, FMT_DATE, None, DV_DATE),
    ("periodo", "Data ou período de destinação informado", "in", 20, None, None, None),
    ("dest", "Unidade destinadora", "in", 28, None, None, ("list", "lst_dest")),
    ("modal", "Modalidade de destinação informada", "in", 26, None, None, ("list", "lst_modal")),
    ("sitsis", "Situação informada no sistema", "in", 16, None, None, ("list", "lst_sit_cdf")),
    ("loc", "Certificado localizado?", "in", 12, None, None, ("list", "lst_simnao")),
    ("conf", "Certificado conferido?", "in", 12, None, None, ("list", "lst_simnao")),
    ("dtconf", "Data da conferência", "in", 12, FMT_DATE, None, DV_DATE),
    ("link", "Link ou caminho do PDF", "in", 28, None, None, None),
    ("resp", "Responsável pela conferência", "in", 18, None, None, None),
    ("obs", "Observações", "in", 26, None, None, None),
    ("kg", "Quantidade total vinculada ao CDF (kg)", "calc", 14, FMT_KG,
     '=IF({act}=0,"",SUMIFS([VIN.kg],[VIN.cdfid],{id},[VIN.qok],1))', None),
    ("nit", "Nº de itens de resíduo vinculados", "calc", 10, FMT_INT,
     '=IF({act}=0,"",COUNTIFS([VIN.cdfid],{id},[VIN.act],1))', None),
    ("sit", "Situação de controle do CDF", "calc", 28, None, CDF_SIT, None),
    ("nmtr", "Nº de MTRs relacionados", "calc", 10, FMT_INT,
     '=IF({act}=0,"",COUNTIFS([VIN.cdfid],{id},[VIN.pairfirst],1))', None),
    # ---- auxiliares
    ("act", "aux registro ativo", "h", 6, None, "=IF(COUNTA({num}:{obs})>0,1,0)", None),
    ("dupkey", "aux chave duplicidade", "h", 6, None,
     '=IF(AND({act}=1,{num}<>""),{sis}&"|"&TRIM({num}),"")', None),
    ("dup", "aux duplicado", "h", 6, None,
     '=IF({dupkey}="",0,IF(SUMPRODUCT(--({dupkey:all}={dupkey}))>1,1,0))', None),
    ("destok", "aux destinador válido", "h", 6, None,
     '=IF({dest}="",0,IF(ISNUMBER(MATCH({dest},cmp_dest,0)),1,0))', None),
    ("confflag", "aux conferido válido", "h", 6, None,
     '=IF(AND({act}=1,{num}<>"",{dest}<>"",{conf}="Sim",{dup}=0),1,0)', None),
    ("nviniss", "aux vínculos com alerta", "h", 6, None,
     "=IF({act}=0,0,COUNTIFS([VIN.cdfid],{id},[VIN.issue],1))", None),
    ("label", "aux rótulo seleção", "h", 6, None,
     '=IF({act}=1,{id}&" | CDF "&IF({num}="","(sem número)",{num})&" | "'
     '&IF({dest}="","(sem destinador)",{dest}),"")', None),
    ("lseq", "aux ordem seleção", "h", 6, None, '=IF({act}=1,COUNTIF({act:top},1),"")', None),
    ("rpass", "aux filtro relatórios", "h", 6, None,
     '=IF({act}=0,0,IF(AND(OR(r_ini="",AND(ISNUMBER({dtem}),{dtem}>=r_ini)),'
     'OR(r_fim="",AND(ISNUMBER({dtem}),{dtem}<=r_fim)),'
     'OR(r_de="",r_de="(Todos)",{dest}=r_de)),1,0))', None),
    ("rkey", "aux ordem relatório 04", "h", 6, None,
     '=IF({rpass}=1,IF({dest}="",0,COUNTIF({dest:all},"<"&{dest}))*1000+{i},"")', None),
]

VIN_COLS = [
    ("id", "ID interno do vínculo", "id", 13, None, None, None),
    ("cdfref", "ID interno do CDF", "in", 34, None, None, ("list", "lst_cdf")),
    ("itemref", "ID interno do item de resíduo", "in", 36, None, None, ("list", "lst_itens")),
    ("mtr", "ID do MTR correspondente", "calc", 13, None,
     '=IF({iok}=1,INDEX([IT.mtrid],{irow})&"","")', None),
    ("res", "Resíduo correspondente", "calc", 24, None,
     '=IF({iok}=1,INDEX([IT.res],{irow})&"","")', None),
    ("qty", "Quantidade vinculada ao CDF", "in", 12, "#,##0.000", None, DV_QTY),
    ("un", "Unidade de medida", "in", 9, None, None, ("list", "lst_un")),
    ("kg", "Quantidade convertida (kg)", "calc", 14, FMT_KG, kg_formula("act", "qty", "fac"), None),
    ("dest", "Unidade destinadora vinculada", "calc", 26, None,
     '=IF({cok}=1,INDEX([CDF.dest],{crow})&"","")', None),
    ("conf", "Conferência do vínculo", "in", 12, None, None, ("list", "lst_simnao")),
    ("obs", "Observações", "in", 26, None, None, None),
    ("alq", "Alerta de quantidade", "calc", 30, None, "=" + nest([
        ("{act}=0", '""'),
        ("NOT(ISNUMBER({qty}))", '"Quantidade não informada"'),
        ("{qty}<=0", '"Quantidade não positiva"'),
        ('{un}=""', '"Unidade não informada"'),
        ("NOT(ISNUMBER({kg}))", '"Unidade a conferir"'),
        ("{iover}=1", '"REVISAR VÍNCULO: soma vinculada maior que a referência do item"'),
    ], '"OK"'), None),
    ("ali", "Alerta de identificação", "calc", 30, None, "=" + nest([
        ("{act}=0", '""'),
        ('{cdfref}=""', '"CDF não informado"'),
        ("{cok}=0", '"REVISAR VÍNCULO: CDF inexistente"'),
        ('{itemref}=""', '"Item não informado"'),
        ("{iok}=0", '"REVISAR VÍNCULO: item inexistente"'),
        ('AND({mism}=1,{conf}<>"Sim")', '"REVISAR VÍNCULO: destinador diferente do MTR"'),
        ("{mism}=1", '"Destinador diferente do MTR (analisado)"'),
    ], '"OK"'), None),
    # ---- auxiliares
    ("act", "aux registro ativo", "h", 6, None,
     "=IF(COUNTA({cdfref}:{itemref},{qty}:{un},{conf}:{obs})>0,1,0)", None),
    ("cdfid", "aux id cdf", "h", 6, None, parse_id("cdfref"), None),
    ("crow", "aux linha cdf", "h", 6, None,
     '=IF({cdfid}="","",IFERROR(MATCH({cdfid},[CDF.id],0),""))', None),
    ("cok", "aux cdf válido", "h", 6, None,
     '=IF({crow}="",0,IF(INDEX([CDF.act],{crow})=1,1,0))', None),
    ("itemid", "aux id item", "h", 6, None, parse_id("itemref"), None),
    ("irow", "aux linha item", "h", 6, None,
     '=IF({itemid}="","",IFERROR(MATCH({itemid},[IT.id],0),""))', None),
    ("iok", "aux item válido", "h", 6, None,
     '=IF({irow}="",0,IF(INDEX([IT.act],{irow})=1,1,0))', None),
    ("fac", "aux fator kg", "h", 6, None, fac_formula("un"), None),
    ("qok", "aux quantidade válida", "h", 6, None,
     "=IF(AND(ISNUMBER({qty}),ISNUMBER({kg})),IF({kg}>0,1,0),0)", None),
    ("mdest", "aux destinatário do MTR", "h", 6, None,
     '=IF({iok}=1,INDEX([IT.mdest],{irow})&"","")', None),
    ("mism", "aux destinador diferente", "h", 6, None,
     '=IF(AND({cok}=1,{iok}=1,{dest}<>"",{mdest}<>""),IF({dest}<>{mdest},1,0),0)', None),
    ("issue", "aux alerta", "h", 6, None,
     '=IF({act}=0,0,IF(OR({cok}=0,{iok}=0,{qok}=0,AND({mism}=1,{conf}<>"Sim")),1,0))', None),
    ("cconf", "aux cdf conferido", "h", 6, None, "=IF({cok}=1,INDEX([CDF.confflag],{crow}),0)", None),
    ("ok", "aux vínculo ok", "h", 6, None,
     '=IF(AND({act}=1,{issue}=0,{cconf}=1,{conf}="Sim"),1,0)', None),
    ("pairfirst", "aux primeiro par cdf mtr", "h", 6, None,
     '=IF(AND({act}=1,{cdfid}<>"",{mtr}<>""),IF(COUNTIFS({cdfid:top},{cdfid},{mtr:top},{mtr})=1,1,0),0)', None),
    ("iover", "aux item acima", "h", 6, None, "=IF({iok}=1,INDEX([IT.over],{irow}),0)", None),
]

UN_COLS = [
    ("id", "ID interno da unidade", "id", 12, None, None, None),
    ("razao", "Razão social", "in", 30, None, None, None),
    ("nome", "Nome da unidade", "in", 28, None, None, None),
    ("cnpj", "CNPJ ou identificação aplicável", "in", 20, "@", None, None),
    ("uf", "UF", "in", 7, None, None, ("list", "lst_uf")),
    ("mun", "Município", "in", 18, None, None, None),
    ("papel", "Papel da unidade", "in", 20, None, None, ("list", "lst_papel")),
    ("sistema", "Sistema ambiental relacionado", "in", 22, None, None, ("list", "lst_sistemas")),
    ("resp", "Responsável interno", "in", 18, None, None, None),
    ("sitcad", "Situação do cadastro", "in", 12, None, None, ("list", "lst_sitcad")),
    ("obs", "Observações", "in", 26, None, None, None),
    ("alerta", "Alerta do cadastro", "calc", 30, None,
     '=IF({nome}="",IF(COUNTA({razao}:{obs})>0,"Informe o nome da unidade",""),'
     'IF(SUMPRODUCT(--({nome:all}={nome}))>1,"Nome repetido: diferencie as unidades","OK"))', None),
    ("has", "aux tem nome", "h", 6, None, '=IF({nome}<>"",1,0)', None),
    ("seq", "aux ordem seleção", "h", 6, None, '=IF({has}=1,COUNTIF({has:top},1),"")', None),
]

PA_COLS = [
    ("id", "ID interno do parceiro", "id", 12, None, None, None),
    ("razao", "Razão social", "in", 30, None, None, None),
    ("fantasia", "Nome fantasia", "in", 24, None, None, None),
    ("cnpj", "CNPJ ou identificação aplicável", "in", 20, "@", None, None),
    ("tipo", "Tipo de atuação", "in", 15, None, None, ("list", "lst_tipo")),
    ("estab", "Unidade/estabelecimento", "in", 20, None, None, None),
    ("mun", "Município", "in", 16, None, None, None),
    ("uf", "UF", "in", 7, None, None, ("list", "lst_uf")),
    ("contato", "Contato", "in", 16, None, None, None),
    ("email", "E-mail", "in", 22, None, None, None),
    ("tel", "Telefone", "in", 15, "@", None, None),
    ("sitcad", "Situação informada do cadastro", "in", 16, None, None, None),
    ("docaut", "Documento de autorização informado", "in", 22, None, None, None),
    ("valid", "Data de validade informada", "in", 12, FMT_DATE, None, DV_DATE),
    ("obs", "Observações", "in", 24, None, None, None),
    ("nomesel", "Nome para seleção", "calc", 30, None,
     '=IF(AND({razao}="",{fantasia}=""),"",IF({fantasia}<>"",{fantasia},{razao})'
     '&IF({estab}<>""," | "&{estab},""))', None),
    ("alerta", "Alerta do cadastro", "calc", 30, None,
     '=IF({nomesel}="",IF(COUNTA({razao}:{obs})>0,"Informe razão social ou nome fantasia",""),'
     'IF({tipo}="","Informe o tipo de atuação",'
     'IF(SUMPRODUCT(--({nomesel:all}={nomesel}))>1,"Nome repetido: diferencie pelo estabelecimento",'
     'IF(AND(ISNUMBER({valid}),{valid}<TODAY()),"Validade informada vencida: conferir","OK"))))', None),
    ("istr", "aux transportador", "h", 6, None,
     '=IF(AND({nomesel}<>"",OR({tipo}="Transportador",{tipo}="Ambos")),1,0)', None),
    ("trseq", "aux ordem transportador", "h", 6, None, '=IF({istr}=1,COUNTIF({istr:top},1),"")', None),
    ("isde", "aux destinador", "h", 6, None,
     '=IF(AND({nomesel}<>"",OR({tipo}="Destinador",{tipo}="Ambos")),1,0)', None),
    ("deseq", "aux ordem destinador", "h", 6, None, '=IF({isde}=1,COUNTIF({isde:top},1),"")', None),
    ("dcount", "aux MTRs no painel", "h", 6, None,
     '=IF({isde}=1,COUNTIFS([MTR.dest],{nomesel},[MTR.pass],1,[MTR.canc],0),"")', None),
    ("dkey", "aux ordem gráfico", "h", 6, None,
     '=IF(N({dcount})>0,{dcount}+(1000-{i})/10000,"")', None),
]

RS_COLS = [
    ("id", "ID interno do resíduo", "id", 12, None, None, None),
    ("nome", "Nome interno", "in", 30, None, None, None),
    ("desc", "Descrição", "in", 32, None, None, None),
    ("proc", "Processo de origem", "in", 24, None, None, None),
    ("cod", "Código informado", "in", 16, "@", None, None),
    ("ref", "Referência utilizada para o código", "in", 26, None, None, None),
    ("cls", "Classificação informada", "in", 20, None, None, None),
    ("unpad", "Unidade padrão", "in", 9, None, None, ("list", "lst_un")),
    ("obs", "Observações", "in", 26, None, None, None),
    ("alerta", "Alerta do cadastro", "calc", 24, None,
     '=IF({nome}="",IF(COUNTA({desc}:{obs})>0,"Informe o nome interno",""),'
     'IF(SUMPRODUCT(--({nome:all}={nome}))>1,"Nome repetido","OK"))', None),
    ("has", "aux tem nome", "h", 6, None, '=IF({nome}<>"",1,0)', None),
    ("seq", "aux ordem seleção", "h", 6, None, '=IF({has}=1,COUNTIF({has:top},1),"")', None),
    ("kgd", "aux kg no painel", "h", 6, None,
     '=IF({has}=0,"",SUMIFS([IT.kg],[IT.res],{nome},[IT.pass],1,[IT.qok],1))', None),
    ("kkey", "aux ordem gráfico", "h", 6, None,
     '=IF(N({kgd})>0,{kgd}+(1000-{i})/1000000,"")', None),
]

COLS = dict(MTR=MTR_COLS, IT=IT_COLS, REC=REC_COLS, CDF=CDF_COLS, VIN=VIN_COLS,
            UN=UN_COLS, PA=PA_COLS, RS=RS_COLS)
LET = {sk: {c[0]: CL(i + 1) for i, c in enumerate(cols)} for sk, cols in COLS.items()}
IDPREFIX = dict(MTR=("MTR-", 6), IT=("IT-", 6), REC=("REC-", 6), CDF=("CDF-", 6), VIN=("VIN-", 6),
                UN=("UN-", 3), PA=("PA-", 3), RS=("RS-", 3))


def rng(sk, key):
    col = LET[sk][key]
    last = FIRST + CAP[sk] - 1
    return f"'{SHEET[sk]}'!${col}${FIRST}:${col}${last}"


_xref = re.compile(r"\[(\w+)\.(\w+)\]")
_local = re.compile(r"\{(\w+)(?::(top|all))?\}")


def render(tpl, sk=None, r=None):
    out = _xref.sub(lambda m: rng(m.group(1), m.group(2)), tpl)
    if sk is None:
        return out
    last = FIRST + CAP[sk] - 1

    def loc(m):
        key, mode = m.group(1), m.group(2)
        if key == "i":
            return str(r - FIRST + 1)
        if key == "r":
            return str(r)
        col = LET[sk][key]
        if mode == "top":
            return f"${col}${FIRST}:{col}{r}"
        if mode == "all":
            return f"${col}${FIRST}:${col}${last}"
        return f"{col}{r}"
    return _local.sub(loc, out)


# ---------------------------------------------------------------- utilidades de layout
def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def put(ws, ref, value=None, f=None, fl=None, al=None, bd=None, fmt=None, prot=None, merge=True):
    """Escreve em uma célula ou intervalo (mesclado) aplicando estilo em todas as células."""
    if ":" in ref:
        if merge:
            ws.merge_cells(ref)
        cells = [c for row in ws[ref] for c in row]
    else:
        cells = [ws[ref]]
    first = cells[0]
    if value is not None:
        first.value = value
    for c in cells:
        if f is not None:
            c.font = f
        if fl is not None:
            c.fill = fl
        if al is not None:
            c.alignment = al
        if bd is not None:
            c.border = bd
        if fmt is not None:
            c.number_format = fmt
        if prot is not None:
            c.protection = prot
    return first


def link(cell, sheet, target="A1"):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{target}")


def band(ws, ref, text, size=20, sub=None, subref=None, color=NAVY):
    put(ws, ref, text, f=font(size, True, WHITE), fl=fill(color),
        al=Alignment(horizontal="left", vertical="center", indent=1))
    if sub:
        put(ws, subref, sub, f=font(11, False, "D6E4F3"), fl=fill(color),
            al=Alignment(horizontal="left", vertical="center", indent=1))


def set_protect(ws):
    p = ws.protection
    p.sheet = True
    p.autoFilter = False
    p.sort = False
    p.formatColumns = False
    p.formatRows = False
    p.selectLockedCells = False
    p.selectUnlockedCells = False


def dv_add(ws, kind, ref, extra=None):
    if kind == "list":
        dv = DataValidation(type="list", formula1=f"={extra}", allow_blank=True)
        dv.error = "Valor fora da lista. Confira o cadastro ou ajuste a lista em CONFIGURAÇÕES."
        dv.errorTitle = "Valor fora da lista"
        dv.errorStyle = "warning"
    elif kind == "date":
        dv = DataValidation(type="date", operator="between", formula1="36526", formula2="73050",
                            allow_blank=True)
        dv.error = "Informe uma data válida (dd/mm/aaaa)."
        dv.errorTitle = "Data"
        dv.errorStyle = "warning"
    elif kind == "qty":
        dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
        dv.error = "A quantidade deve ser um número maior que zero."
        dv.errorTitle = "Quantidade"
        dv.errorStyle = "warning"
    elif kind == "inline":
        dv = DataValidation(type="list", formula1=extra, allow_blank=True)
        dv.errorStyle = "warning"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(ref)
    return dv


# regras de cor por texto (ordem = prioridade)
CF_TEXT = [
    ("grey", ["Cancelado", "cancelado"]),
    ("red", ["revisar", "REVISAR", "incompleto", "INCOMPLETO", "inexistente", "inválid", "não informad",
             "fora do cadastro", "duplicidade", "DUPLICIDADE", "UNIDADE A CONFERIR", "Unidade a conferir",
             "sem evidência", "anterior à emissão", "não positiva", "não identificado", "Sem MTR",
             "Necessita", "incompletas", "Nome repetido", "Informe", "vencida"]),
    ("yellow", ["a conferir", "A CONFERIR", "Aguardando conferência", "Localizado", "ver diferença",
                "analisado", "Data futura", "Não calculável"]),
    ("orange", ["Aguardando", "parcial", "Cadastrado", "NÃO CONFIRMADO", "NÃO LOCALIZADO", "VÍNCULO INCOMPLETO"]),
    ("green", ["conferida", "Conferido", "OK", "Nenhuma"]),
]
CF_STYLE = {
    "grey": (GREY, DGREY), "red": (LRED, RED_TXT), "yellow": (LYELLOW, "8A6100"),
    "orange": (LORANGE, "9A4A00"), "green": (LGREEN, "1E7B43"),
}


def cf_text(ws, ref, anchor):
    for cat, words in CF_TEXT:
        bg, fg = CF_STYLE[cat]
        cond = "OR(" + ",".join(f'ISNUMBER(SEARCH("{w}",{anchor}))' for w in words) + ")"
        if cat == "green":
            cond = f'OR(EXACT({anchor},"OK"),' + ",".join(
                f'ISNUMBER(SEARCH("{w}",{anchor}))' for w in words if w != "OK") + ")"
        ws.conditional_formatting.add(ref, FormulaRule(
            formula=[cond], fill=fill(bg), font=Font(name=FONT, color=fg, bold=True), stopIfTrue=True))


# ---------------------------------------------------------------- abas de registro
SHEET_INFO = {
    "MTR": ("CADASTRO DE MTRs", "Cada linha é um MTR. Preencha as colunas azuis com os dados do documento oficial. "
            "As colunas cinzas são calculadas pela planilha.",
            "Depois de cadastrar o MTR, registre os resíduos dele na aba ITENS_MTR."),
    "IT": ("ITENS DE RESÍDUO DOS MTRs", "Cada linha é um resíduo de um MTR. Um MTR pode ter vários itens.",
           "Escolha o MTR e o resíduo pelas listas. Toneladas (t) são convertidas para kg (1 t = 1.000 kg). "
           "Marque DIFERENÇA ANALISADA? = Sim só depois de conferir uma diferença de quantidade."),
    "REC": ("CONTROLE DE RECEBIMENTOS", "Cada linha é um registro de recebimento de um item de resíduo.",
            "O recebimento só conta como confirmado com RECEBIMENTO CONFERIDO? = Sim e um documento ou link "
            "de evidência informado. Data preenchida sozinha não confirma nada."),
    "CDF": ("CADASTRO DE CDFs", "Cada linha é um Certificado de Destinação Final realmente disponível.",
            "Um CDF pode cobrir vários MTRs. A ligação com os resíduos é feita na aba VÍNCULOS_CDF."),
    "VIN": ("VÍNCULOS ENTRE CDF E ITENS DE MTR", "Cada linha liga um CDF a um item de resíduo, com a quantidade "
            "que o certificado documenta para esse item.",
            "Exemplo: CDF X com 500 kg do item do MTR A e 300 kg do item do MTR B = 2 linhas, 800 kg no CDF X."),
    "UN": ("CADASTRO DE UNIDADES", "Estabelecimentos da sua organização (matriz, filiais, obras).",
           "Use um nome diferente para cada unidade. Esse nome aparece nas listas das outras abas."),
    "PA": ("CADASTRO DE PARCEIROS", "Transportadores e destinadores.",
           "O nome para seleção é montado com o nome fantasia (ou razão social) e o estabelecimento. "
           "Estar cadastrado não significa estar autorizado: confira a documentação nas fontes oficiais."),
    "RS": ("CADASTRO DE RESÍDUOS", "Identificação interna padronizada dos resíduos.",
           "Preencha código e classificação conforme a referência técnica que você consultou. "
           "A planilha não atribui código nem periculosidade."),
}
FREEZE = dict(MTR="C6", IT="B6", REC="B6", CDF="C6", VIN="B6", UN="C6", PA="B6", RS="C6")
TABLE_NAME = dict(MTR="tbMTR", IT="tbItens", REC="tbRecebimentos", CDF="tbCDF", VIN="tbVinculos",
                  UN="tbUnidades", PA="tbParceiros", RS="tbResiduos")
CF_COLS = dict(MTR=["status", "pend", "dias", "dupal", "todas"], IT=["sit", "cod"], REC=["alerta"],
               CDF=["sit"], VIN=["alq", "ali"], UN=["alerta"], PA=["alerta"], RS=["alerta"])

S_ID = dict(f=font(10, True, NAVY), fl=fill(GREY))
S_IN = dict(f=font(10), fl=fill(LBLUE))
S_CALC = dict(f=font(10, False, "2B3A4A"), fl=fill(GREY))
S_H = dict(f=font(9, False, "8899AA"))


def build_data_sheet(wb, sk, rows=None):
    ws = wb.create_sheet(SHEET[sk])
    cols = COLS[sk]
    cap = CAP[sk]
    last = FIRST + cap - 1
    vis = [c for c in cols if c[2] != "h"]
    nvis = len(vis)
    lastvis = CL(nvis)
    title, sub, tip = SHEET_INFO[sk]
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    # cabeçalho
    ws.row_dimensions[1].height = 34
    put(ws, f"A1:{lastvis}1", title, f=font(18, True, WHITE), fl=fill(NAVY),
        al=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[2].height = 30
    put(ws, f"A2:{lastvis}2", sub + "  " + tip, f=font(10, False, "1F2933"), fl=fill(LBLUE),
        al=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    ws.row_dimensions[3].height = 20
    for j, (txt, tgt, cell) in enumerate([("◀ INÍCIO", SHEET["INI"], "A3"), ("PAINEL", SHEET["PAI"], "B3"),
                                          ("PENDÊNCIAS", SHEET["PEN"], "C3")]):
        c = put(ws, cell, txt, f=Font(name=FONT, size=9, bold=True, color=BLUE, underline="single"),
                al=Alignment(horizontal="left", vertical="center"))
        link(c, tgt)
    put(ws, "D3", "Azul-claro = preencher   |   Cinza = automático   |   Linhas prontas: "
        f"{cap:,}".replace(",", ".") + " registros", f=font(9, False, DGREY, True),
        al=Alignment(horizontal="left", vertical="center"))
    ws.row_dimensions[5].height = 46

    # rótulos de grupo (linha 4) e cabeçalhos (linha 5)
    for i, c in enumerate(cols):
        key, head, kind, width = c[:4]
        col = CL(i + 1)
        ws.column_dimensions[col].width = width
        h = ws.cell(row=5, column=i + 1, value=head)
        tag = ws.cell(row=4, column=i + 1)
        if kind == "h":
            ws.column_dimensions[col].hidden = True
            h.font = font(8, False, "8899AA")
            continue
        h.alignment = A_CENTER
        h.border = Border(left=Side(style="thin", color=WHITE), right=Side(style="thin", color=WHITE))
        if kind == "calc":
            h.fill, h.font = fill(BLUE), font(10, True, WHITE)
            tag.value, tag.font, tag.fill = "AUTOMÁTICO", font(8, True, DGREY), fill(GREY)
        else:
            h.fill, h.font = fill(NAVY), font(10, True, WHITE)
            tag.value = "PREENCHER" if kind == "in" else "ID FIXO"
            tag.font, tag.fill = font(8, True, BLUE if kind == "in" else NAVY), fill(LBLUE if kind == "in" else GREY)
        tag.alignment = A_CENTER

    # linhas de dados
    pre, ndig = IDPREFIX[sk]
    data = rows or []
    for n in range(cap):
        r = FIRST + n
        rec = data[n] if n < len(data) else {}
        for i, c in enumerate(cols):
            key, head, kind, width, fmt, formula, dv = c
            cell = ws.cell(row=r, column=i + 1)
            if kind == "id":
                cell.value = f"{pre}{n + 1:0{ndig}d}"
                cell.font, cell.fill = S_ID["f"], S_ID["fl"]
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif kind == "in":
                v = rec.get(key)
                if v is not None:
                    cell.value = v
                cell.font, cell.fill = S_IN["f"], S_IN["fl"]
                cell.protection = UNLOCK
                cell.alignment = Alignment(vertical="center")
            elif kind == "calc":
                cell.value = render(formula, sk, r)
                cell.font, cell.fill = S_CALC["f"], S_CALC["fl"]
                cell.alignment = Alignment(vertical="center")
            else:
                cell.value = render(formula, sk, r)
                cell.font = S_H["f"]
            if fmt:
                cell.number_format = fmt

    # validações
    for i, c in enumerate(cols):
        dv = c[6]
        if dv:
            col = CL(i + 1)
            dv_add(ws, dv[0], f"{col}{FIRST}:{col}{last}", dv[1])

    # tabela estruturada
    tab = Table(displayName=TABLE_NAME[sk], ref=f"A5:{lastvis}{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)

    # formatação condicional: cores de situação, depois faixas alternadas
    for key in CF_COLS[sk]:
        col = LET[sk][key]
        cf_text(ws, f"{col}{FIRST}:{col}{last}", f"{col}{FIRST}")
    if sk == "MTR":
        b, ac, d, ov = LET[sk]["num"], LET[sk]["dup"], LET[sk]["dias"], LET[sk]["overdue"]
        ws.conditional_formatting.add(f"{b}{FIRST}:{b}{last}", FormulaRule(
            formula=[f"${ac}{FIRST}=1"], fill=fill(LRED), font=Font(name=FONT, color=RED_TXT, bold=True)))
        ws.conditional_formatting.add(f"{d}{FIRST}:{d}{last}", FormulaRule(
            formula=[f"${ov}{FIRST}=1"], fill=fill(LRED), font=Font(name=FONT, color=RED_TXT, bold=True),
            stopIfTrue=True))
    if sk == "IT":
        q = LET[sk]["qty"]
        ws.conditional_formatting.add(f"{q}{FIRST}:{q}{last}", FormulaRule(
            formula=[f'AND(${LET[sk]["act"]}{FIRST}=1,${LET[sk]["qtyok"]}{FIRST}=0)'], fill=fill(LRED)))
        kgc = LET[sk]["kg"]
        cf_text(ws, f"{kgc}{FIRST}:{kgc}{last}", f"{kgc}{FIRST}")
    # faixas alternadas
    ins = [CL(i + 1) for i, c in enumerate(cols) if c[2] == "in"]
    calcs = [CL(i + 1) for i, c in enumerate(cols) if c[2] == "calc"]
    in_ref = " ".join(f"{c}{FIRST}:{c}{last}" for c in ins)
    calc_ref = " ".join(f"{c}{FIRST}:{c}{last}" for c in calcs)
    ws.conditional_formatting.add(in_ref, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(BAND_IN)))
    if calcs:
        ws.conditional_formatting.add(calc_ref, FormulaRule(formula=["MOD(ROW(),2)=0"], fill=fill(BAND_CALC)))

    ws.freeze_panes = FREEZE[sk]
    ws.sheet_properties.tabColor = TAB_CAD if sk in ("UN", "PA", "RS") else BLUE
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"
    set_protect(ws)
    return ws


# ---------------------------------------------------------------- CONFIGURAÇÕES
CFG_LISTS = [
    # (coluna, título, valores, nome, linhas)
    ("B", "Sistemas de origem", ["SINIR / MTR Nacional", "Sistema estadual", "Outro sistema aplicável"],
     "lst_sistemas", 10),
    ("C", "Situações informadas no sistema (MTR e recebimento)",
     ["Salvo", "Emitido", "Recebido", "Armazenamento temporário", "Cancelado", "Outra (ver observações)"],
     "lst_sit_mtr", 12),
    ("D", "Situações informadas no sistema (CDF)", ["Emitido", "Disponível para consulta", "Cancelado",
                                                    "Outra (ver observações)"], "lst_sit_cdf", 10),
    ("E", "Modalidades de destinação (exemplos editáveis)",
     ["Reciclagem", "Reutilização", "Compostagem", "Rerrefino", "Coprocessamento", "Recuperação energética",
      "Tratamento térmico", "Tratamento físico-químico", "Aterro sanitário", "Aterro industrial",
      "Outra (conforme CDF)"], "lst_modal", 20),
    ("F", "UF", UFS, "lst_uf", 27),
    ("G", "Unidade de medida", ["kg", "t"], "lst_un", 10),
    ("I", "Sim / Não", ["Sim", "Não"], "lst_simnao", 2),
    ("J", "Tipo de atuação (parceiros)", ["Transportador", "Destinador", "Ambos", "Armazenador temporário",
                                          "Outro"], "lst_tipo", 8),
    ("K", "Papel da unidade", ["Geradora", "Armazenador temporário", "Destinadora", "Outro"], "lst_papel", 8),
    ("L", "Situação do cadastro", ["Ativo", "Inativo"], "lst_sitcad", 5),
]
LIST_TOP = 17
# listas automáticas ocultas
HIDDEN_LISTS = [
    # (coluna, título, fonte(compact), cap, seq)
    ("AA", "Unidades", ("UN", "nome", "seq")),
    ("AB", "Transportadores", ("PA", "nomesel", "trseq")),
    ("AC", "Destinadores", ("PA", "nomesel", "deseq")),
    ("AD", "Resíduos", ("RS", "nome", "seq")),
    ("AE", "MTRs", ("MTR", "label", "lseq")),
    ("AF", "Itens", ("IT", "label", "lseq")),
    ("AG", "CDFs", ("CDF", "label", "lseq")),
]
FILTER_LISTS = [("AH", "(Todas)", "AA", "UN", "seq", "lst_f_unid"),
                ("AI", "(Todos)", "AB", "PA", "trseq", "lst_f_transp"),
                ("AJ", "(Todos)", "AC", "PA", "deseq", "lst_f_dest"),
                ("AK", "(Todos)", "AD", "RS", "seq", "lst_f_res")]
HL_FIRST = 6
PEND_ROW = 52
SIT_ROW = 63


def build_config(wb):
    ws = wb.create_sheet(SHEET["CFG"])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_CAD
    setw(ws, {"A": 2, "B": 44, "C": 30, "D": 30, "E": 30, "F": 7, "G": 14, "H": 11, "I": 9, "J": 22,
              "K": 22, "L": 16})
    ws.row_dimensions[1].height = 34
    band(ws, "A1:L1", "CONFIGURAÇÕES", 18)
    put(ws, "A2:L2", "Parâmetros internos e listas usadas nas abas. Altere apenas as células azul-claro. "
        "Nas listas, preencha de cima para baixo, sem deixar linhas vazias no meio.",
        f=font(10), fl=fill(LBLUE), al=A_LEFT)
    c = put(ws, "A3", "◀ INÍCIO", f=Font(name=FONT, size=9, bold=True, color=BLUE, underline="single"))
    link(c, SHEET["INI"])

    put(ws, "B5:C5", "PARÂMETROS GERAIS", f=font(12, True, WHITE), fl=fill(BLUE), al=A_LEFT)
    params = [
        ("Nome da organização usuária", None, "cfg_org", None),
        ("Unidade padrão de massa", "kg", "cfg_massa", None),
        ("Dias para alerta interno de recebimento", 15, "cfg_rec", FMT_INT),
        ("Dias para alerta interno de CDF", 15, "cfg_cdf", FMT_INT),
        ("Tolerância para alerta de diferença de quantidade (kg)", 0, "cfg_tol", FMT_KG),
        ("Termo que indica MTR cancelado (Situação informada no sistema)", "Cancelado", "cfg_cancel", None),
        ("Versão da planilha", VERSAO, "cfg_versao", None),
    ]
    for k, (lab, val, name, fmt) in enumerate(params):
        r = 6 + k
        put(ws, f"B{r}", lab, f=font(10, True, NAVY), bd=BORDER, al=A_LEFT)
        editable = name not in ("cfg_massa", "cfg_versao")
        cc = put(ws, f"C{r}", val, f=font(10, True), fl=fill(LBLUE if editable else GREY), bd=BORDER,
                 al=Alignment(horizontal="center", vertical="center"), fmt=fmt,
                 prot=UNLOCK if editable else LOCK)
        wb.defined_names[name] = DefinedName(name, attr_text=f"'{SHEET['CFG']}'!$C${r}")
    put(ws, "D8:L8", "Parâmetro interno de organização. Não é prazo legal.", f=font(9, False, DGREY, True))
    put(ws, "D9:L9", "Parâmetro interno de organização. Não é prazo legal.", f=font(9, False, DGREY, True))
    put(ws, "D10:L10", "Diferenças maiores que este valor geram alerta de revisão. Use 0 para alertar qualquer diferença.",
        f=font(9, False, DGREY, True))
    put(ws, "D11:L11", "A planilha considera cancelado o MTR cuja situação informada seja igual a este termo.",
        f=font(9, False, DGREY, True))
    ws.row_dimensions[13].height = 30
    put(ws, "B13:L13", "Prazos de alerta são parâmetros internos de acompanhamento, não prazos legais.",
        f=font(13, True, "7A4E00"), fl=fill(YELLOW), al=Alignment(horizontal="center", vertical="center"))

    put(ws, "B15:L15", "LISTAS EDITÁVEIS", f=font(12, True, WHITE), fl=fill(BLUE), al=A_LEFT)
    ws.row_dimensions[16].height = 44
    for col, title, vals, name, n in CFG_LISTS:
        put(ws, f"{col}16", title, f=font(9, True, WHITE), fl=fill(NAVY), al=A_CENTER)
        for k in range(n):
            r = LIST_TOP + k
            put(ws, f"{col}{r}", vals[k] if k < len(vals) else None, f=font(10), fl=fill(LBLUE), bd=BORDER,
                prot=UNLOCK)
        top = f"'{SHEET['CFG']}'!${col}${LIST_TOP}"
        full = f"'{SHEET['CFG']}'!${col}${LIST_TOP}:${col}${LIST_TOP + n - 1}"
        wb.defined_names[name] = DefinedName(name, attr_text=f"OFFSET({top},0,0,MAX(1,COUNTA({full})),1)")
    # fatores
    put(ws, "H16", "Fator para kg", f=font(9, True, WHITE), fl=fill(NAVY), al=A_CENTER)
    for k in range(10):
        r = LIST_TOP + k
        put(ws, f"H{r}", [1, 1000][k] if k < 2 else None, f=font(10), fl=fill(LBLUE), bd=BORDER,
            prot=UNLOCK, fmt="#,##0.####")
    wb.defined_names["cfg_units"] = DefinedName("cfg_units", attr_text=f"'{SHEET['CFG']}'!$G${LIST_TOP}:$G${LIST_TOP + 9}")
    wb.defined_names["cfg_fac"] = DefinedName("cfg_fac", attr_text=f"'{SHEET['CFG']}'!$H${LIST_TOP}:$H${LIST_TOP + 9}")
    put(ws, "G28:L31", "Unidades sem fator (por exemplo L ou m³) não são somadas e aparecem como UNIDADE A CONFERIR. "
        "Só informe um fator de volume para massa quando houver densidade fundamentada e documentada.",
        f=font(9, False, DGREY, True), al=A_WRAP)
    put(ws, "B30:D33", "As modalidades de destinação são exemplos editáveis. Use o termo que aparece no CDF. "
        "A lista de situações também pode ser ajustada aos termos do sistema que você utiliza.",
        f=font(9, False, DGREY, True), al=A_WRAP)

    # tipos de pendência
    put(ws, f"B{PEND_ROW - 2}:L{PEND_ROW - 2}", "TIPOS DE PENDÊNCIA (prioridade e próxima providência editáveis)",
        f=font(12, True, WHITE), fl=fill(BLUE), al=A_LEFT)
    for j, h in enumerate(["Tipo de pendência", "Prioridade interna base", "Próxima providência sugerida"]):
        put(ws, f"{'BCD'[j]}{PEND_ROW - 1}" if j < 2 else f"D{PEND_ROW - 1}:L{PEND_ROW - 1}", h,
            f=font(10, True, WHITE), fl=fill(NAVY), al=A_CENTER)
    for k, (t, p, prov) in enumerate(PEND_TIPOS):
        r = PEND_ROW + k
        put(ws, f"B{r}", t, f=font(10, True, NAVY), fl=fill(GREY), bd=BORDER, al=A_LEFT)
        put(ws, f"C{r}", p, f=font(10), fl=fill(LBLUE), bd=BORDER, al=A_CENTER, prot=UNLOCK)
        put(ws, f"D{r}:L{r}", prov, f=font(10), fl=fill(LBLUE), bd=BORDER, al=A_LEFT, prot=UNLOCK)
    dv_add(ws, "inline", f"C{PEND_ROW}:C{PEND_ROW + 6}", '"Alta,Média,Baixa"')
    for name, col in (("cfg_pt", "B"), ("cfg_pp", "C"), ("cfg_pv", "D")):
        wb.defined_names[name] = DefinedName(
            name, attr_text=f"'{SHEET['CFG']}'!${col}${PEND_ROW}:${col}${PEND_ROW + 6}")
    put(ws, f"B{PEND_ROW + 7}:L{PEND_ROW + 7}", "Quando a pendência passa do número de dias configurado acima, "
        "a prioridade sobe para Alta. Prioridade é um critério de organização interna, não uma classificação legal.",
        f=font(9, False, DGREY, True), al=A_WRAP)

    # situações internas (referência)
    put(ws, f"B{SIT_ROW - 1}:L{SIT_ROW - 1}", "SITUAÇÕES INTERNAS DE ACOMPANHAMENTO (calculadas, para consulta)",
        f=font(12, True, WHITE), fl=fill(BLUE), al=A_LEFT)
    for k, (s, d) in enumerate(SIT_INTERNAS):
        r = SIT_ROW + k
        ws.row_dimensions[r].height = 30
        put(ws, f"B{r}", s, f=font(10, True, NAVY), fl=fill(GREY), bd=BORDER, al=A_LEFT)
        put(ws, f"C{r}:L{r}", d, f=font(10), bd=BORDER, al=A_LEFT)
    cf_text(ws, f"B{SIT_ROW}:B{SIT_ROW + 7}", f"B{SIT_ROW}")

    # paleta
    pr = SIT_ROW + 10
    put(ws, f"B{pr}:L{pr}", "PALETA DE CORES DA PLANILHA", f=font(12, True, WHITE), fl=fill(BLUE), al=A_LEFT)
    pal = [("Azul-marinho", NAVY, "Títulos e cabeçalhos"), ("Azul médio", BLUE, "Tabelas, gráficos e navegação"),
           ("Azul-claro", LBLUE, "Campos para preencher"), ("Verde", GREEN, "Confirmações"),
           ("Verde-claro", LGREEN, "Informação conferida"), ("Amarelo", YELLOW, "Atenção"),
           ("Vermelho suave", LRED, "Inconsistências"), ("Cinza-claro", GREY, "Campos automáticos")]
    for k, (n, hexc, use) in enumerate(pal):
        r = pr + 1 + k
        put(ws, f"B{r}", f"{n}  #{hexc}", f=font(10, True, NAVY), bd=BORDER, al=A_LEFT)
        put(ws, f"C{r}", None, fl=fill(hexc), bd=BORDER)
        put(ws, f"D{r}:E{r}", use, f=font(10), bd=BORDER, al=A_LEFT)

    # listas automáticas ocultas
    put(ws, "AA4", "LISTAS AUTOMÁTICAS (NÃO EDITAR)", f=font(9, True, NAVY))
    for col, title, (sk, key, seq) in HIDDEN_LISTS:
        n = CAP[sk]
        ws[f"{col}5"] = title
        for k in range(n):
            ws[f"{col}{HL_FIRST + k}"] = (f'=IFERROR(INDEX({rng(sk, key)},MATCH({k + 1},{rng(sk, seq)},0))&"","")')
        nm = {"AA": "lst_unidades", "AB": "lst_transp", "AC": "lst_dest", "AD": "lst_res", "AE": "lst_mtr",
              "AF": "lst_itens", "AG": "lst_cdf"}[col]
        top = f"'{SHEET['CFG']}'!${col}${HL_FIRST}"
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"OFFSET({top},0,0,MAX(1,MAX({rng(sk, seq)})),1)")
        cmp = {"AB": "cmp_transp", "AC": "cmp_dest"}.get(col)
        if cmp:
            wb.defined_names[cmp] = DefinedName(
                cmp, attr_text=f"'{SHEET['CFG']}'!${col}${HL_FIRST}:${col}${HL_FIRST + n - 1}")
    for col, first, src, sk, seq, nm in FILTER_LISTS:
        n = CAP[sk]
        ws[f"{col}5"] = "Filtro " + nm[6:]
        ws[f"{col}{HL_FIRST}"] = first
        for k in range(n):
            ws[f"{col}{HL_FIRST + 1 + k}"] = f"={src}{HL_FIRST + k}"
        top = f"'{SHEET['CFG']}'!${col}${HL_FIRST}"
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"OFFSET({top},0,0,MAX(1,MAX({rng(sk, seq)}))+1,1)")
    ws["AL5"] = "Filtro situação"
    sits = ["(Todas)"] + [s for s, _ in SIT_INTERNAS]
    for k, s in enumerate(sits):
        ws[f"AL{HL_FIRST + k}"] = s
    wb.defined_names["lst_f_sit"] = DefinedName(
        "lst_f_sit", attr_text=f"'{SHEET['CFG']}'!$AL${HL_FIRST}:$AL${HL_FIRST + len(sits) - 1}")
    for col in ["AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL"]:
        ws.column_dimensions[col].hidden = True
        ws.column_dimensions[col].width = 30
    ws.print_area = "A1:L82"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    set_protect(ws)
    return ws


# ---------------------------------------------------------------- PAINEL
def value_labels():
    dl = DataLabelList()
    dl.showVal = True
    dl.showSerName = False
    dl.showCatName = False
    dl.showLegendKey = False
    dl.showPercent = False
    dl.showLeaderLines = False
    dl.numFmt = "#,##0;;;"
    return dl


STATUS_COLORS = ["E06C5F", "C0392B", ORANGE, YELLOW, "D9A21B", BLUE, GREEN, "AAB4C0"]
MONTHS = 'CHOOSE(MONTH({d}),"jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez")&"/"&RIGHT(YEAR({d}),2)'


def build_painel(wb, filters=None):
    ws = wb.create_sheet(SHEET["PAI"])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN
    ws.sheet_view.zoomScale = 90
    setw(ws, {"A": 2, "T": 2, "U": 3})
    for i in range(2, 20):
        ws.column_dimensions[CL(i)].width = 9.6
    for col in "VWXYZ":
        ws.column_dimensions[col].width = 14
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 38
    ws.row_dimensions[3].height = 22
    put(ws, "B2:P2", "♻  PAINEL DE CONTROLE — MTR E CDF", f=font(22, True, WHITE), fl=fill(NAVY),
        al=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "B3:P3", "Acompanhe movimentações, documentos e pendências em um só lugar.",
        f=font(11, False, "D6E4F3"), fl=fill(NAVY), al=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "Q2:S2", "BÔNUS 02", f=font(11, True, WHITE), fl=fill(GREEN), al=A_CENTER)
    put(ws, "Q3:S3", '=IF(cfg_org="","Atualizado em "&TEXT(DAY(TODAY()),"00")&"/"&TEXT(MONTH(TODAY()),"00")'
        '&"/"&YEAR(TODAY()),cfg_org)', f=font(9, False, WHITE), fl=fill(NAVY), al=A_CENTER)
    for col, txt, tgt in (("B4", "◀ INÍCIO", SHEET["INI"]), ("D4", "PENDÊNCIAS", SHEET["PEN"]),
                          ("F4", "RELATÓRIOS", SHEET["REL"]), ("H4", "MTRs", SHEET["MTR"])):
        c = put(ws, col, txt, f=Font(name=FONT, size=9, bold=True, color=BLUE, underline="single"))
        link(c, tgt)

    # filtros
    f = filters or {}
    flt = [("B", "C", "Período inicial", "f_ini", f.get("f_ini"), "date"),
           ("D", "E", "Período final", "f_fim", f.get("f_fim"), "date"),
           ("F", "H", "Unidade geradora", "f_un", f.get("f_un", "(Todas)"), "lst_f_unid"),
           ("I", "K", "Transportador", "f_tr", f.get("f_tr", "(Todos)"), "lst_f_transp"),
           ("L", "N", "Destinador", "f_de", f.get("f_de", "(Todos)"), "lst_f_dest"),
           ("O", "P", "Tipo de resíduo", "f_res", f.get("f_res", "(Todos)"), "lst_f_res"),
           ("Q", "S", "Situação documental", "f_sit", f.get("f_sit", "(Todas)"), "lst_f_sit")]
    ws.row_dimensions[6].height = 18
    ws.row_dimensions[7].height = 24
    for a, b, lab, name, val, dvk in flt:
        put(ws, f"{a}6:{b}6", "▼ " + lab, f=font(9, True, NAVY), al=Alignment(horizontal="left", vertical="bottom"))
        put(ws, f"{a}7:{b}7", val, f=font(10, True, NAVY), fl=fill(LBLUE),
            bd=Border(bottom=Side(style="medium", color=BLUE), top=thin, left=thin, right=thin),
            al=A_CENTER, fmt=FMT_DATE if dvk == "date" else None, prot=UNLOCK)
        wb.defined_names[name] = DefinedName(name, attr_text=f"'{SHEET['PAI']}'!${a}$7")
        if dvk == "date":
            dv_add(ws, "date", f"{a}7")
        else:
            dv_add(ws, "list", f"{a}7", dvk)
    put(ws, "B8:S8", "Os filtros de período usam a DATA DE EMISSÃO DO MTR. Deixe as datas em branco e as listas em "
        "(Todas)/(Todos) para ver tudo. MTRs com cancelamento informado ficam fora dos cards 02 a 06 e das quantidades.",
        f=font(8.5, False, DGREY, True), al=Alignment(horizontal="left", vertical="center", wrap_text=True))
    ws.row_dimensions[8].height = 24

    # cards
    M = lambda k: rng("MTR", k)
    cards = [
        ("📄  TOTAL DE MTRs", NAVY, f"=COUNTIFS({M('pass')},1)",
         f'="Inclui "&COUNTIFS({M("pass")},1,{M("canc")},1)&" com cancelamento informado"'),
        ("🚚⏱  AGUARDANDO RECEBIMENTO", ORANGE, f"=COUNTIFS({M('pass')},1,{M('canc')},0,{M('recconf')},0)",
         '="Sem recebimento informado e conferido"'),
        ("✅  RECEBIMENTO CONFIRMADO", GREEN, f"=COUNTIFS({M('pass')},1,{M('canc')},0,{M('recconf')},1)",
         '="Todos os itens com recebimento conferido"'),
        ("🔎  AGUARDANDO CDF", "D98E04",
         f"=COUNTIFS({M('pass')},1,{M('canc')},0,{M('recconf')},1,{M('cdfdone')},0)",
         '="Recebidos, sem CDF suficiente ou conferido"'),
        ("🗂  DOCUMENTAÇÃO CONFERIDA", GREEN, f'=COUNTIFS({M("pass")},1,{M("status")},"Documentação conferida")',
         '="Conferência interna concluída"'),
        ("⚠  REGISTROS A REVISAR", "C0392B",
         f'=COUNTIFS({M("pass")},1,{M("status")},"Informação a revisar")'
         f'+COUNTIFS({M("pass")},1,{M("status")},"Cadastro incompleto")',
         '="Inconsistências pelas regras internas"'),
    ]
    ws.row_dimensions[10].height = 30
    ws.row_dimensions[11].height = 30
    ws.row_dimensions[12].height = 22
    ws.row_dimensions[13].height = 26
    for k, (title, color, formula, cap) in enumerate(cards):
        a, b = CL(2 + 3 * k), CL(4 + 3 * k)
        side = Side(style="thin", color=color)
        put(ws, f"{a}10:{b}10", title, f=font(9.5, True, WHITE), fl=fill(color), al=A_CENTER)
        put(ws, f"{a}11:{b}12", formula, f=font(30, True, color), fl=fill(WHITE), al=A_CENTER, fmt="#,##0",
            bd=Border(left=side, right=side))
        put(ws, f"{a}13:{b}13", cap, f=font(8.5, False, DGREY, True), fl=fill(WHITE), al=A_CENTER,
            bd=Border(left=side, right=side, bottom=Side(style="thick", color=color)))

    # indicadores complementares
    I = lambda k: rng("IT", k)
    C = lambda k: rng("CDF", k)
    put(ws, "B15:S15", "INDICADORES COMPLEMENTARES", f=font(11, True, NAVY),
        bd=Border(bottom=Side(style="medium", color=NAVY)))
    sec = [
        ("Quantidade movimentada (kg)", f"=SUMIFS({I('kg')},{I('pass')},1,{I('qok')},1)", FMT_KG,
         f'="Itens fora da soma por unidade ou quantidade a conferir: "&COUNTIFS({I("pass")},1,{I("qok")},0)'),
        ("Com recebimento informado (kg)", f"=SUMIFS({I('kgrec')},{I('pass')},1)", FMT_KG,
         '="Soma dos recebimentos registrados"'),
        ("Vinculada a CDF (kg)", f"=SUMIFS({I('kgcdf')},{I('pass')},1)", FMT_KG, '="Soma dos vínculos registrados"'),
        ("CDFs cadastrados", f"=COUNTIF({C('act')},1)", "#,##0", '="Total geral, sem filtros"'),
        ("CDFs conferidos", f'=COUNTIF({C("sit")},"Conferido")', "#,##0", '="Total geral, sem filtros"'),
        ("Pendências acompanhadas", f"=COUNTIFS({M('pass')},1,{M('haspend')},1)", "#,##0",
         '="MTRs com pendência principal"'),
    ]
    ws.row_dimensions[17].height = 28
    for k, (lab, formula, fmt, note) in enumerate(sec):
        a, b = CL(2 + 3 * k), CL(4 + 3 * k)
        put(ws, f"{a}16:{b}16", lab, f=font(9, True, DGREY), fl=fill(GREY), al=A_CENTER)
        put(ws, f"{a}17:{b}17", formula, f=font(16, True, NAVY), fl=fill(GREY), al=A_CENTER, fmt=fmt)
        put(ws, f"{a}18:{b}18", note, f=font(7.5, False, DGREY, True), fl=fill(GREY), al=A_CENTER)
        ws.row_dimensions[18].height = 24

    put(ws, "B20:S20", "GRÁFICOS", f=font(11, True, NAVY), bd=Border(bottom=Side(style="medium", color=NAVY)))

    # dados dos gráficos
    put(ws, "V5:Y5", "DADOS DOS GRÁFICOS (automático)", f=font(9, True, WHITE), fl=fill(BLUE), al=A_CENTER)
    put(ws, "V6", "Situação", f=font(9, True, NAVY))
    put(ws, "W6", "MTRs", f=font(9, True, NAVY))
    for k, (s, _) in enumerate(SIT_INTERNAS):
        r = 7 + k
        ws[f"V{r}"] = s
        ws[f"W{r}"] = f'=COUNTIFS({M("pass")},1,{M("status")},V{r})'
    put(ws, "V17", "Mês", f=font(9, True, NAVY))
    put(ws, "W17", "Mês/ano", f=font(9, True, NAVY))
    put(ws, "X17", "kg", f=font(9, True, NAVY))
    ws["Y17"] = "Base"
    ws["Y18"] = (f"=IF(ISNUMBER(f_fim),f_fim,IF(MAX({M('dtem')})>0,MAX({M('dtem')}),TODAY()))")
    ws["Y18"].number_format = FMT_DATE
    for k in range(12):
        r = 18 + k
        if k == 0:
            ws[f"V{r}"] = "=IF(ISNUMBER(f_ini),DATE(YEAR(f_ini),MONTH(f_ini),1),EDATE(DATE(YEAR(Y18),MONTH(Y18),1),-11))"
        else:
            ws[f"V{r}"] = f"=EDATE(V{r - 1},1)"
        ws[f"V{r}"].number_format = FMT_DATE
        ws[f"W{r}"] = "=" + MONTHS.format(d=f"V{r}")
        ws[f"X{r}"] = f"=SUMIFS({I('kg')},{I('pass')},1,{I('qok')},1,{I('mes')},V{r})"
    R = lambda k: rng("RS", k)
    P = lambda k: rng("PA", k)
    put(ws, "V32", "chave", f=font(9, True, NAVY))
    put(ws, "W32", "Resíduo", f=font(9, True, NAVY))
    put(ws, "X32", "kg", f=font(9, True, NAVY))
    for k in range(8):
        r = 33 + k
        ws[f"V{r}"] = f'=IFERROR(LARGE({R("kkey")},{k + 1}),"")'
        ws[f"W{r}"] = f'=IF(V{r}="","",INDEX({R("nome")},MATCH(V{r},{R("kkey")},0)))'
        ws[f"X{r}"] = f'=IF(V{r}="",0,INDEX({R("kgd")},MATCH(V{r},{R("kkey")},0)))'
    put(ws, "V43", "chave", f=font(9, True, NAVY))
    put(ws, "W43", "Destinador", f=font(9, True, NAVY))
    put(ws, "X43", "MTRs", f=font(9, True, NAVY))
    for k in range(8):
        r = 44 + k
        ws[f"V{r}"] = f'=IFERROR(LARGE({P("dkey")},{k + 1}),"")'
        ws[f"W{r}"] = f'=IF(V{r}="","",INDEX({P("nomesel")},MATCH(V{r},{P("dkey")},0)))'
        ws[f"X{r}"] = f'=IF(V{r}="",0,INDEX({P("dcount")},MATCH(V{r},{P("dkey")},0)))'
    for row in ws.iter_rows(min_row=6, max_row=51, min_col=22, max_col=25):
        for c in row:
            if c.font is None or not c.font.b:
                c.font = font(9, False, DGREY)

    # gráficos
    def style_axes(ch):
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.majorGridlines = None
        ch.legend = None

    d1 = DoughnutChart(holeSize=58)
    d1.title = "Situação dos MTRs"
    d1.add_data(Reference(ws, min_col=23, min_row=6, max_row=14), titles_from_data=True)
    d1.set_categories(Reference(ws, min_col=22, min_row=7, max_row=14))
    s = d1.series[0]
    for k, colr in enumerate(STATUS_COLORS):
        pt = DataPoint(idx=k)
        pt.graphicalProperties.solidFill = colr
        pt.graphicalProperties.line.solidFill = WHITE
        s.dPt.append(pt)
    d1.dataLabels = value_labels()
    d1.legend.position = "r"
    d1.width, d1.height = 16.4, 9.2
    ws.add_chart(d1, "B21")

    c2 = BarChart()
    c2.type = "col"
    c2.title = "Movimentação por mês (kg, data de emissão)"
    c2.add_data(Reference(ws, min_col=24, min_row=17, max_row=29), titles_from_data=True)
    c2.set_categories(Reference(ws, min_col=23, min_row=18, max_row=29))
    c2.series[0].graphicalProperties.solidFill = BLUE
    c2.series[0].graphicalProperties.line.solidFill = BLUE
    c2.gapWidth = 60
    style_axes(c2)
    c2.y_axis.numFmt = "#,##0"
    c2.width, c2.height = 16.4, 9.2
    ws.add_chart(c2, "K21")

    c3 = BarChart()
    c3.type = "bar"
    c3.title = "Resíduos mais movimentados (kg)"
    c3.add_data(Reference(ws, min_col=24, min_row=32, max_row=40), titles_from_data=True)
    c3.set_categories(Reference(ws, min_col=23, min_row=33, max_row=40))
    c3.series[0].graphicalProperties.solidFill = GREEN
    c3.series[0].graphicalProperties.line.solidFill = GREEN
    c3.gapWidth = 50
    style_axes(c3)
    c3.x_axis.scaling.orientation = "maxMin"
    c3.y_axis.numFmt = "#,##0"
    c3.dataLabels = value_labels()
    c3.width, c3.height = 16.4, 9.2
    ws.add_chart(c3, "B40")

    c4 = BarChart()
    c4.type = "bar"
    c4.title = "Principais destinadores (nº de MTRs)"
    c4.add_data(Reference(ws, min_col=24, min_row=43, max_row=51), titles_from_data=True)
    c4.set_categories(Reference(ws, min_col=23, min_row=44, max_row=51))
    c4.series[0].graphicalProperties.solidFill = NAVY
    c4.series[0].graphicalProperties.line.solidFill = NAVY
    c4.gapWidth = 50
    style_axes(c4)
    c4.x_axis.scaling.orientation = "maxMin"
    c4.dataLabels = value_labels()
    c4.width, c4.height = 16.4, 9.2
    ws.add_chart(c4, "K40")

    # quadro de atenção
    put(ws, "B59:S59", "⚠  PENDÊNCIAS QUE PRECISAM DE ACOMPANHAMENTO", f=font(12, True, WHITE), fl=fill("C0392B"),
        al=A_LEFT)
    ws.row_dimensions[59].height = 26
    put(ws, "B60:S60", "As 8 primeiras da Central de Pendências, por prioridade e tempo. Este quadro mostra todos os "
        "registros, sem os filtros do topo.", f=font(8.5, False, DGREY, True), al=A_LEFT)
    heads = [("B", "C", "ID do MTR", "B"), ("D", "E", "Nº oficial", "C"), ("F", "I", "Destinador", "F"),
             ("J", "M", "Tipo de pendência", "G"), ("N", "O", "Tempo (dias)", "I"), ("P", "S", "Próxima ação", "K")]
    ws.row_dimensions[61].height = 22
    for a, b, h, _ in heads:
        put(ws, f"{a}61:{b}61", h, f=font(9.5, True, WHITE), fl=fill(NAVY), al=A_CENTER)
    for k in range(8):
        r = 62 + k
        src = PEN_FIRST + k
        ws.row_dimensions[r].height = 30
        for a, b, h, scol in heads:
            put(ws, f"{a}{r}:{b}{r}", f"='{SHEET['PEN']}'!{scol}{src}&\"\"" if scol != "I" else
                f"='{SHEET['PEN']}'!{scol}{src}",
                f=font(9, scol in "BG", NAVY if scol == "B" else "1F2933"),
                fl=fill(GREY if k % 2 else WHITE), bd=Border(bottom=thin),
                al=A_CENTER if scol in "BCI" else A_LEFT, fmt=FMT_INT if scol == "I" else None)
    cf_text(ws, "J62:J69", "J62")
    put(ws, "B70:S70", '=IF(COUNT(' + rng("MTR", "pkey") + ')=0,"Nenhuma pendência no momento.",'
        '"Total de pendências na central: "&COUNT(' + rng("MTR", "pkey") + '))',
        f=font(9, True, NAVY), al=A_LEFT)
    c = put(ws, "P71:S71", "Abrir a central completa ▶", f=Font(name=FONT, size=9, bold=True, color=BLUE,
                                                                 underline="single"), al=A_CENTER)
    link(c, SHEET["PEN"])
    ws.row_dimensions[73].height = 30
    put(ws, "B73:S73", DISCLAIMER, f=font(9, True, NAVY, True), fl=fill(LBLUE), al=A_CENTER)

    ws.print_area = "A1:S73"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    set_protect(ws)
    return ws


# ---------------------------------------------------------------- PENDÊNCIAS
PEN_FIRST = 11
PEN_CAP = 500


def build_pendencias(wb):
    ws = wb.create_sheet(SHEET["PEN"])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ORANGE
    ws.sheet_view.zoomScale = 90
    setw(ws, {"A": 2, "B": 13, "C": 17, "D": 22, "E": 24, "F": 24, "G": 26, "H": 13, "I": 12, "J": 11,
              "K": 46, "L": 46, "M": 6})
    ws.row_dimensions[1].height = 36
    put(ws, "B1:L1", "CENTRAL DE PENDÊNCIAS DOCUMENTAIS", f=font(20, True, WHITE), fl=fill(NAVY),
        al=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "B2:L2", "Identifique rapidamente os registros que precisam de atenção.", f=font(11, False, "D6E4F3"),
        fl=fill(NAVY), al=Alignment(horizontal="left", vertical="center", indent=1))
    for col, txt, tgt in (("B3", "◀ INÍCIO", SHEET["INI"]), ("C3", "PAINEL", SHEET["PAI"]),
                          ("D3", "MTRs", SHEET["MTR"])):
        c = put(ws, col, txt, f=Font(name=FONT, size=9, bold=True, color=BLUE, underline="single"))
        link(c, tgt)
    M = lambda k: rng("MTR", k)
    cards = [
        ("B", "C", "Aguardando recebimento", ORANGE, f'=COUNTIF({M("pend")},"RECEBIMENTO NÃO CONFIRMADO")'),
        ("D", "D", "Aguardando CDF", "D98E04",
         f'=COUNTIF({M("pend")},"CDF NÃO LOCALIZADO")+COUNTIF({M("pend")},"VÍNCULO INCOMPLETO")'
         f'+COUNTIF({M("pend")},"CDF A CONFERIR")'),
        ("E", "E", "Informações incompletas", "C0392B", f'=COUNTIF({M("pend")},"MTR INCOMPLETO")'),
        ("F", "F", "Divergências a investigar", "C0392B",
         f'=COUNTIF({M("pend")},"QUANTIDADE A REVISAR")+COUNTIF({M("pend")},"IDENTIFICAÇÃO A REVISAR")'),
        ("G", "G", "Acima do período interno", NAVY, f"=COUNTIF({M('overdue')},1)"),
    ]
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 36
    for a, b, lab, color, formula in cards:
        ref5 = f"{a}5:{b}5" if a != b else f"{a}5"
        ref6 = f"{a}6:{b}6" if a != b else f"{a}6"
        put(ws, ref5, lab, f=font(9.5, True, WHITE), fl=fill(color), al=A_CENTER)
        put(ws, ref6, formula, f=font(22, True, color), fl=fill(WHITE), al=A_CENTER, fmt="#,##0",
            bd=Border(bottom=Side(style="thick", color=color), left=thin, right=thin))
    R = lambda k: rng("REC", k)
    V = lambda k: rng("VIN", k)
    I = lambda k: rng("IT", k)
    put(ws, "H5:L6", f'="Registros sem referência válida (não entram na lista abaixo): itens sem MTR válido "'
        f'&COUNTIFS({I("act")},1,{I("mok")},0)&" | recebimentos "&COUNTIFS({R("act")},1,{R("iok")},0)'
        f'&" | vínculos "&(COUNTIFS({V("act")},1,{V("cok")},0)+COUNTIFS({V("act")},1,{V("cok")},1,{V("iok")},0))'
        f'&". Confira as colunas de alerta nas abas ITENS_MTR, RECEBIMENTOS e VÍNCULOS_CDF."',
        f=font(9, False, DGREY, True), fl=fill(GREY), al=A_WRAP)
    put(ws, "B8:L8", "Prioridade é um critério operacional interno. Pendências e diferenças internas não constituem, "
        "por si só, infrações. Use a lista para decidir o que conferir primeiro.",
        f=font(9, False, DGREY, True), al=A_LEFT)
    heads = ["ID do MTR", "Número oficial do MTR", "Unidade geradora", "Transportador", "Destinador",
             "Tipo de pendência", "Data de referência", "Dias em acompanhamento", "Prioridade interna",
             "Próxima providência sugerida", "Observações (todas as pendências do MTR)"]
    hr = PEN_FIRST - 1
    ws.row_dimensions[hr].height = 40
    for j, h in enumerate(heads):
        put(ws, f"{CL(2 + j)}{hr}", h, f=font(10, True, WHITE), fl=fill(NAVY), al=A_CENTER)
    ws[f"M{hr}"] = "aux linha"
    keys = ["id", "num", "unid", "transp", "dest", "pend", "ref", "dias", "prio", "prov", "todas"]
    fmts = {"ref": FMT_DATE, "dias": FMT_INT}
    last = PEN_FIRST + PEN_CAP - 1
    for k in range(PEN_CAP):
        r = PEN_FIRST + k
        ws[f"M{r}"] = f'=IFERROR(MATCH(SMALL({M("pkey")},{k + 1}),{M("pkey")},0),"")'
        ws.row_dimensions[r].height = 30
        for j, key in enumerate(keys):
            c = ws.cell(row=r, column=2 + j)
            c.value = f'=IF($M{r}="","",INDEX({M(key)},$M{r})&"")' if key not in ("ref", "dias") else \
                f'=IF($M{r}="","",INDEX({M(key)},$M{r}))'
            c.font = font(9.5, j == 0, NAVY if j == 0 else "1F2933")
            c.alignment = Alignment(vertical="center", wrap_text=j >= 2, horizontal="center" if j in (0, 6, 7, 8) else None)
            if key in fmts:
                c.number_format = fmts[key]
    ws.column_dimensions["M"].hidden = True
    tab = Table(displayName="tbPendencias", ref=f"B{hr}:L{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)
    rr = f"B{PEN_FIRST}:L{last}"
    ws.conditional_formatting.add(f"J{PEN_FIRST}:J{last}", FormulaRule(
        formula=[f'J{PEN_FIRST}="Alta"'], fill=fill(LRED), font=Font(name=FONT, color=RED_TXT, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(f"J{PEN_FIRST}:J{last}", FormulaRule(
        formula=[f'J{PEN_FIRST}="Média"'], fill=fill(LORANGE), font=Font(name=FONT, color="9A4A00", bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(f"J{PEN_FIRST}:J{last}", FormulaRule(
        formula=[f'J{PEN_FIRST}="Baixa"'], fill=fill(LYELLOW), font=Font(name=FONT, color="8A6100", bold=True), stopIfTrue=True))
    cf_text(ws, f"G{PEN_FIRST}:G{last}", f"G{PEN_FIRST}")
    ws.conditional_formatting.add(rr, FormulaRule(formula=[f'AND($B{PEN_FIRST}<>"",MOD(ROW(),2)=0)'], fill=fill(GREY)))
    ws.freeze_panes = f"C{PEN_FIRST}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    set_protect(ws)
    return ws


# ---------------------------------------------------------------- RELATÓRIOS
def build_relatorios(wb, filters=None):
    ws = wb.create_sheet(SHEET["REL"])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_view.zoomScale = 90
    setw(ws, {"A": 2, "B": 16, "C": 20, "D": 24, "E": 22, "F": 20, "G": 22, "H": 22, "I": 18, "J": 24,
              "K": 3, "L": 10})
    ws.row_dimensions[1].height = 36
    put(ws, "B1:J1", "RELATÓRIOS AUTOMÁTICOS", f=font(20, True, WHITE), fl=fill(NAVY),
        al=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "B2:J2", "Consultas e consolidações feitas com os dados cadastrados. Para imprimir ou gerar PDF, use "
        "Arquivo > Imprimir ou Arquivo > Exportar.", f=font(10, False, "D6E4F3"), fl=fill(NAVY),
        al=Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True))
    for col, txt, tgt in (("B3", "◀ INÍCIO", SHEET["INI"]), ("C3", "PAINEL", SHEET["PAI"])):
        c = put(ws, col, txt, f=Font(name=FONT, size=9, bold=True, color=BLUE, underline="single"))
        link(c, tgt)
    f = filters or {}
    put(ws, "B5:J5", "FILTROS DOS RELATÓRIOS", f=font(11, True, NAVY), bd=Border(bottom=Side(style="medium", color=NAVY)))
    flt = [("B", "B", "Período inicial", "r_ini", f.get("r_ini"), "date"),
           ("C", "C", "Período final", "r_fim", f.get("r_fim"), "date"),
           ("D", "E", "Unidade geradora", "r_un", f.get("r_un", "(Todas)"), "lst_f_unid"),
           ("F", "G", "Destinador", "r_de", f.get("r_de", "(Todos)"), "lst_f_dest")]
    ws.row_dimensions[7].height = 24
    for a, b, lab, name, val, dvk in flt:
        r6 = f"{a}6:{b}6" if a != b else f"{a}6"
        r7 = f"{a}7:{b}7" if a != b else f"{a}7"
        put(ws, r6, "▼ " + lab, f=font(9, True, NAVY))
        put(ws, r7, val, f=font(10, True, NAVY), fl=fill(LBLUE), al=A_CENTER,
            bd=Border(bottom=Side(style="medium", color=BLUE), top=thin, left=thin, right=thin),
            fmt=FMT_DATE if dvk == "date" else None, prot=UNLOCK)
        wb.defined_names[name] = DefinedName(name, attr_text=f"'{SHEET['REL']}'!${a}$7")
        if dvk == "date":
            dv_add(ws, "date", f"{a}7")
        else:
            dv_add(ws, "list", f"{a}7", dvk)
    put(ws, "H6:J7", "Cada relatório informa qual data o período considera. Deixe em branco ou em (Todas)/(Todos) "
        "para ver tudo.", f=font(8.5, False, DGREY, True), al=A_WRAP)
    ws["L7"] = ('=IF(AND(r_ini="",r_fim=""),"Todo o período",IF(ISNUMBER(r_ini),TEXT(DAY(r_ini),"00")&"/"'
                '&TEXT(MONTH(r_ini),"00")&"/"&YEAR(r_ini),"início")&" a "&IF(ISNUMBER(r_fim),TEXT(DAY(r_fim),"00")'
                '&"/"&TEXT(MONTH(r_fim),"00")&"/"&YEAR(r_fim),"hoje"))')
    ws["L6"] = '=IF(OR(r_un="",r_un="(Todas)"),"Todas",r_un)'
    M = lambda k: rng("MTR", k)
    I = lambda k: rng("IT", k)
    Rr = lambda k: rng("REC", k)
    Cc = lambda k: rng("CDF", k)
    breaks = []

    def header(row, num, title, note):
        put(ws, f"B{row}:J{row}", f"RELATÓRIO {num:02d} | {title}", f=font(13, True, WHITE), fl=fill(BLUE), al=A_LEFT)
        ws.row_dimensions[row].height = 26
        put(ws, f"B{row + 1}:G{row + 1}", note, f=font(9, False, DGREY, True), al=A_WRAP)
        ws.row_dimensions[row + 1].height = 28

    def heads(row, names):
        ws.row_dimensions[row].height = 34
        for j, h in enumerate(names):
            put(ws, f"{CL(2 + j)}{row}", h, f=font(9.5, True, WHITE), fl=fill(NAVY), al=A_CENTER)

    def table(name, row1, row2, ncols):
        t = Table(displayName=name, ref=f"B{row1}:{CL(1 + ncols)}{row2}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(t)
        ws.conditional_formatting.add(f"B{row1 + 1}:{CL(1 + ncols)}{row2}", FormulaRule(
            formula=[f"MOD(ROW(),2)=0"], fill=fill(GREY)))

    def cellf(r, c, v, fmt=None, bold=False):
        x = ws.cell(row=r, column=c, value=v)
        x.font = font(9.5, bold, NAVY if bold else "1F2933")
        x.alignment = Alignment(vertical="center")
        if fmt:
            x.number_format = fmt
        return x

    # R01
    s = 10
    header(s, 1, "MTRs POR PERÍODO", "Data considerada: data de emissão do MTR. Mostra 12 meses a partir do período "
           "inicial (ou os 12 meses até o período final ou o MTR mais recente). Filtros: período, unidade e destinador.")
    heads(s + 2, ["Mês/ano", "Unidade", "Nº de MTRs", "Quantidade movimentada (kg)", "Documentação conferida",
                  "Em acompanhamento", "A revisar ou incompletos", "Cancelados"])
    ws[f"L{s + 1}"] = (f"=IF(ISNUMBER(r_fim),r_fim,IF(MAX({M('dtem')})>0,MAX({M('dtem')}),TODAY()))")
    ws[f"L{s + 1}"].number_format = FMT_DATE
    acomp = ["Aguardando recebimento", "Aguardando CDF", "Vínculo de CDF parcial", "CDF localizado, a conferir"]
    for k in range(12):
        r = s + 3 + k
        if k == 0:
            ws[f"L{r}"] = (f"=IF(ISNUMBER(r_ini),DATE(YEAR(r_ini),MONTH(r_ini),1),"
                           f"EDATE(DATE(YEAR(L{s + 1}),MONTH(L{s + 1}),1),-11))")
        else:
            ws[f"L{r}"] = f"=EDATE(L{r - 1},1)"
        ws[f"L{r}"].number_format = FMT_DATE
        base = f"{M('rpass')},1,{M('mes')},$L{r}"
        cellf(r, 2, "=" + MONTHS.format(d=f"L{r}"), bold=True)
        cellf(r, 3, "=$L$6")
        cellf(r, 4, f"=COUNTIFS({base})", "#,##0")
        cellf(r, 5, f"=SUMIFS({I('kg')},{I('rpass')},1,{I('qok')},1,{I('mes')},$L{r})", FMT_KG)
        cellf(r, 6, f'=COUNTIFS({base},{M("status")},"Documentação conferida")', "#,##0")
        cellf(r, 7, "=" + "+".join(f'COUNTIFS({base},{M("status")},"{x}")' for x in acomp), "#,##0")
        cellf(r, 8, f'=COUNTIFS({base},{M("status")},"Informação a revisar")+COUNTIFS({base},{M("status")},"Cadastro incompleto")', "#,##0")
        cellf(r, 9, f"=COUNTIFS({base},{M('canc')},1)", "#,##0")
    table("tbRel01", s + 2, s + 14, 8)
    tr = s + 15
    cellf(tr, 2, "TOTAL", bold=True)
    for c in range(4, 10):
        x = cellf(tr, c, f"=SUM({CL(c)}{s + 3}:{CL(c)}{s + 14})", FMT_KG if c == 5 else "#,##0", bold=True)
    for c in range(2, 10):
        ws.cell(row=tr, column=c).fill = fill(LBLUE)
        ws.cell(row=tr, column=c).border = Border(top=Side(style="medium", color=NAVY))
    put(ws, f"B{tr + 1}:J{tr + 1}", "MTRs sem data de emissão não aparecem nesta tabela (eles ficam como Cadastro "
        "incompleto na Central de Pendências).", f=font(8.5, False, DGREY, True), al=A_LEFT)

    # R02
    s = tr + 4
    breaks.append(s - 1)
    NR2 = 50
    header(s, 2, "MOVIMENTAÇÃO POR RESÍDUO", "Data considerada: data de emissão do MTR. Lista até 50 resíduos do "
           "cadastro. MTRs com cancelamento informado não entram na soma. Filtros: período, unidade e destinador.")
    heads(s + 2, ["Resíduo", "Código cadastrado", "Nº de itens", "Quantidade movimentada (kg)", "Unidade geradora",
                  "Período"])
    for k in range(NR2):
        r = s + 3 + k
        cellf(r, 2, f"='{SHEET['CFG']}'!AD{HL_FIRST + k}&\"\"", bold=True)
        cellf(r, 3, f'=IF(B{r}="","",INDEX({rng("RS", "cod")},MATCH(B{r},{rng("RS", "nome")},0))&"")')
        cellf(r, 4, f'=IF(B{r}="","",COUNTIFS({I("res")},B{r},{I("rpass")},1))', "#,##0")
        cellf(r, 5, f'=IF(B{r}="","",SUMIFS({I("kg")},{I("res")},B{r},{I("rpass")},1,{I("qok")},1))', FMT_KG)
        cellf(r, 6, f'=IF(B{r}="","",$L$6)')
        cellf(r, 7, f'=IF(B{r}="","",$L$7)')
    table("tbRel02", s + 2, s + 2 + NR2, 6)
    tr = s + 3 + NR2
    cellf(tr, 2, "TOTAL", bold=True)
    cellf(tr, 4, f"=SUM(D{s + 3}:D{tr - 1})", "#,##0", True)
    cellf(tr, 5, f"=SUM(E{s + 3}:E{tr - 1})", FMT_KG, True)
    for c in range(2, 8):
        ws.cell(row=tr, column=c).fill = fill(LBLUE)
        ws.cell(row=tr, column=c).border = Border(top=Side(style="medium", color=NAVY))

    # R03
    s = tr + 3
    breaks.append(s - 1)
    NR3 = 200
    header(s, 3, "RECEBIMENTOS", "Data considerada: data de recebimento informada. Filtros: período, unidade "
           "geradora do MTR e destinador. Mostra até 200 registros; para mais, use os filtros da aba RECEBIMENTOS.")
    put(ws, f"H{s + 1}:J{s + 1}", f'="Registros encontrados: "&COUNTIF({Rr("rpass")},1)', f=font(9.5, True, NAVY),
        al=A_LEFT)
    heads(s + 2, ["ID do MTR", "Número oficial do MTR", "Destinatário", "Data de chegada informada",
                  "Quantidade registrada (kg)", "Situação de conferência", "ID do recebimento"])
    for k in range(NR3):
        r = s + 3 + k
        ws[f"L{r}"] = f'=IFERROR(MATCH({k + 1},{Rr("rseq")},0),"")'
        g = lambda key, fmt=None, b=False, txt=True: (
            f'=IF($L{r}="","",INDEX({Rr(key)},$L{r}){chr(38) + chr(34) * 2 if txt else ""})')
        cellf(r, 2, g("mtr"), bold=True)
        cellf(r, 3, f'=IF($L{r}="","",IFERROR(INDEX({M("num")},INDEX({Rr("mrow")},$L{r}))&"",""))')
        cellf(r, 4, g("dest"))
        cellf(r, 5, g("dt", txt=False), FMT_DATE)
        cellf(r, 6, g("kg", txt=False), FMT_KG)
        cellf(r, 7, g("alerta"))
        cellf(r, 8, g("id"))
    table("tbRel03", s + 2, s + 2 + NR3, 7)
    cf_text(ws, f"G{s + 3}:G{s + 2 + NR3}", f"G{s + 3}")

    # R04
    s = s + 3 + NR3 + 2
    breaks.append(s - 1)
    NR4 = 150
    header(s, 4, "CDFs POR DESTINADOR", "Data considerada: data de emissão do CDF. Ordenado por destinador. "
           "Filtros: período e destinador. Mostra até 150 certificados.")
    put(ws, f"H{s + 1}:J{s + 1}", f'="CDFs encontrados: "&COUNTIF({Cc("rpass")},1)', f=font(9.5, True, NAVY), al=A_LEFT)
    heads(s + 2, ["Destinador", "Número do CDF", "ID interno do CDF", "Data de emissão", "Quantidade vinculada (kg)",
                  "Nº de MTRs relacionados", "Situação interna de conferência"])
    for k in range(NR4):
        r = s + 3 + k
        ws[f"L{r}"] = f'=IFERROR(MATCH(SMALL({Cc("rkey")},{k + 1}),{Cc("rkey")},0),"")'
        g = lambda key, txt=True: f'=IF($L{r}="","",INDEX({Cc(key)},$L{r}){chr(38) + chr(34) * 2 if txt else ""})'
        cellf(r, 2, g("dest"), bold=True)
        cellf(r, 3, g("num"))
        cellf(r, 4, g("id"))
        cellf(r, 5, g("dtem", False), FMT_DATE)
        cellf(r, 6, g("kg", False), FMT_KG)
        cellf(r, 7, g("nmtr", False), "#,##0")
        cellf(r, 8, g("sit"))
    table("tbRel04", s + 2, s + 2 + NR4, 7)
    cf_text(ws, f"H{s + 3}:H{s + 2 + NR4}", f"H{s + 3}")

    # R05
    s = s + 3 + NR4 + 2
    breaks.append(s - 1)
    NR5 = 150
    header(s, 5, "PENDÊNCIAS DOCUMENTAIS", "Situação atual, sem filtro de período. Filtros: unidade geradora e "
           "destinador. Ordenado por prioridade interna e tempo. Mostra até 150 pendências.")
    put(ws, f"H{s + 1}:J{s + 1}", f'="Pendências encontradas: "&COUNT({M("r5key")})', f=font(9.5, True, NAVY), al=A_LEFT)
    heads(s + 2, ["ID do MTR", "Número oficial do MTR", "Pendência principal", "Participante relacionado",
                  "Tempo em acompanhamento (dias)", "Prioridade interna", "Próxima providência"])
    for k in range(NR5):
        r = s + 3 + k
        ws[f"L{r}"] = f'=IFERROR(MATCH(SMALL({M("r5key")},{k + 1}),{M("r5key")},0),"")'
        g = lambda key, txt=True: f'=IF($L{r}="","",INDEX({M(key)},$L{r}){chr(38) + chr(34) * 2 if txt else ""})'
        cellf(r, 2, g("id"), bold=True)
        cellf(r, 3, g("num"))
        cellf(r, 4, g("pend"))
        cellf(r, 5, g("partic"))
        cellf(r, 6, g("dias", False), FMT_INT)
        cellf(r, 7, g("prio"))
        x = cellf(r, 8, g("prov"))
        x.alignment = Alignment(vertical="center", wrap_text=True)
    table("tbRel05", s + 2, s + 2 + NR5, 7)
    cf_text(ws, f"D{s + 3}:D{s + 2 + NR5}", f"D{s + 3}")

    # R06
    s = s + 3 + NR5 + 2
    breaks.append(s - 1)
    NR6 = CAP["MTR"]
    header(s, 6, "CONCILIAÇÃO DOCUMENTAL", "Data considerada: data de emissão do MTR. Filtros: período, unidade e "
           "destinador. Diferenças são indicações para conferência, não conclusões sobre a regularidade.")
    put(ws, f"H{s + 1}:J{s + 1}", f'="MTRs listados: "&COUNTIF({M("rpass")},1)', f=font(9.5, True, NAVY), al=A_LEFT)
    heads(s + 2, ["ID do MTR", "Número oficial do MTR", "Itens", "Quantidade enviada (kg)",
                  "Quantidade recebida registrada (kg)", "Quantidade vinculada a CDF (kg)",
                  "Diferença enviado − recebido (kg)", "Diferença recebido − CDF (kg)", "Situação de conferência"])
    for k in range(NR6):
        r = s + 3 + k
        ws[f"L{r}"] = f'=IFERROR(MATCH({k + 1},{M("rseq")},0),"")'
        g = lambda key, txt=True: f'=IF($L{r}="","",INDEX({M(key)},$L{r}){chr(38) + chr(34) * 2 if txt else ""})'
        cellf(r, 2, g("id"), bold=True)
        cellf(r, 3, g("num"))
        cellf(r, 4, g("nit", False), "#,##0")
        cellf(r, 5, g("kgtot", False), FMT_KG)
        cellf(r, 6, g("kgrec", False), FMT_KG)
        cellf(r, 7, g("kgcdf", False), FMT_KG)
        cellf(r, 8, f'=IF($L{r}="","",IF(INDEX({M("nrec")},$L{r})=0,"",N(E{r})-N(F{r})))', FMT_KGDIF)
        cellf(r, 9, f'=IF($L{r}="","",IF(INDEX({M("nrec")},$L{r})=0,"",N(F{r})-N(G{r})))', FMT_KGDIF)
        cellf(r, 10, g("status"))
    table("tbRel06", s + 2, s + 2 + NR6, 9)
    cf_text(ws, f"J{s + 3}:J{s + 2 + NR6}", f"J{s + 3}")
    lastrow = s + 2 + NR6
    put(ws, f"B{lastrow + 2}:J{lastrow + 2}", DISCLAIMER, f=font(9, True, NAVY, True), fl=fill(LBLUE), al=A_CENTER)

    ws.column_dimensions["L"].hidden = True
    for b in breaks:
        ws.row_breaks.append(Break(id=b))
    ws.print_area = f"A1:J{lastrow + 2}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "A8"
    set_protect(ws)
    return ws


# ---------------------------------------------------------------- INÍCIO
def build_inicio(wb):
    ws = wb.active
    ws.title = SHEET["INI"]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    ws.column_dimensions["A"].width = 3
    for i in range(2, 17):
        ws.column_dimensions[CL(i)].width = 9.3
    ws.column_dimensions["Q"].width = 3
    for r, h in {1: 10, 2: 44, 3: 26, 4: 16}.items():
        ws.row_dimensions[r].height = h
    put(ws, "B2:M2", "PLANILHA PRONTA DE CONTROLE DE MTR E CDF", f=font(24, True, WHITE), fl=fill(NAVY),
        al=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "B3:M3", "Organize movimentações, acompanhe documentos e identifique pendências.",
        f=font(13, False, "D6E4F3"), fl=fill(NAVY), al=Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "B4:M4", None, fl=fill(NAVY))
    put(ws, "N2:P2", "BÔNUS 02", f=font(13, True, WHITE), fl=fill(GREEN), al=A_CENTER)
    put(ws, "N3:P4", "♻", f=font(26, True, WHITE), fl=fill(NAVY), al=A_CENTER)

    # fluxo visual
    ws.row_dimensions[6].height = 40
    steps = [("B", "D", "📄  MTR"), ("F", "H", "🚚  TRANSPORTE"), ("J", "L", "🏭  DESTINATÁRIO"), ("N", "P", "📜  CDF")]
    for a, b, t in steps:
        put(ws, f"{a}6:{b}6", t, f=font(12, True, NAVY), fl=fill(LBLUE), al=A_CENTER,
            bd=Border(bottom=Side(style="medium", color=BLUE)))
    for col in "EIM":
        put(ws, f"{col}6", "➜", f=font(18, True, GREEN), al=A_CENTER)
    put(ws, "B7:P7", "Recebimento e destinação são etapas diferentes: a planilha acompanha cada uma separadamente.",
        f=font(9, False, DGREY, True), al=Alignment(horizontal="center", vertical="center"))

    # 5 passos
    put(ws, "B9:P9", "COMECE EM 5 PASSOS", f=font(14, True, NAVY), bd=Border(bottom=Side(style="medium", color=NAVY)))
    passos = [("01 — CADASTRE AS UNIDADES", "Informe os estabelecimentos que serão utilizados.", "UN"),
              ("02 — CADASTRE OS PARCEIROS", "Adicione transportadores e destinadores.", "PA"),
              ("03 — REGISTRE OS MTRs", "Inclua os documentos e seus resíduos.", "MTR"),
              ("04 — ACOMPANHE O RECEBIMENTO", "Atualize a situação com base nos registros disponíveis.", "REC"),
              ("05 — VINCULE OS CDFs", "Registre os certificados e confira o painel.", "CDF")]
    for r, h in {10: 30, 11: 36, 12: 40, 13: 18}.items():
        ws.row_dimensions[r].height = h
    for k, (t, d, sk) in enumerate(passos):
        a, b = CL(2 + 3 * k), CL(4 + 3 * k)
        num, title = t.split(" — ")
        put(ws, f"{a}10:{b}10", num, f=font(20, True, WHITE), fl=fill(BLUE if k % 2 == 0 else NAVY), al=A_CENTER)
        put(ws, f"{a}11:{b}11", title, f=font(10.5, True, NAVY), fl=fill(GREY), al=A_CENTER)
        put(ws, f"{a}12:{b}12", d, f=font(9.5, False, "33414F"), fl=fill(GREY), al=A_CENTER)
        c = put(ws, f"{a}13:{b}13", f"abrir {SHEET[sk]} ▶", f=Font(name=FONT, size=8.5, bold=True, color=BLUE,
                                                                     underline="single"), fl=fill(GREY), al=A_CENTER)
        link(c, SHEET[sk], "B6")

    # navegação
    put(ws, "B15:P15", "NAVEGAÇÃO RÁPIDA", f=font(14, True, NAVY), bd=Border(bottom=Side(style="medium", color=NAVY)))
    botoes = [("ABRIR PAINEL", "PAI", "A1", GREEN), ("CADASTRAR MTR", "MTR", "B6", NAVY),
              ("CONSULTAR RECEBIMENTOS", "REC", "C6", BLUE), ("REGISTRAR CDF", "CDF", "B6", NAVY),
              ("VER PENDÊNCIAS", "PEN", "A1", ORANGE), ("ABRIR RELATÓRIOS", "REL", "A1", BLUE)]
    for k, (t, sk, tgt, colr) in enumerate(botoes):
        row = 16 + (k // 3) * 2
        a, b = CL(2 + (k % 3) * 5), CL(5 + (k % 3) * 5)
        ws.row_dimensions[row].height = 34
        c = put(ws, f"{a}{row}:{b}{row}", t + "  ▶", f=font(12, True, WHITE), fl=fill(colr), al=A_CENTER)
        link(c, SHEET[sk], tgt)

    # legenda
    put(ws, "B21:P21", "LEGENDA DE CORES", f=font(14, True, NAVY), bd=Border(bottom=Side(style="medium", color=NAVY)))
    leg = [("AZUL-CLARO", LBLUE, NAVY, "informações para preencher."),
           ("CINZA-CLARO", GREY, DGREY, "informações calculadas automaticamente."),
           ("VERDE", LGREEN, "1E7B43", "informação conferida."),
           ("AMARELO", LYELLOW, "8A6100", "atenção ou revisão necessária."),
           ("VERMELHO-CLARO", LRED, RED_TXT, "inconsistência relevante ou registro que exige providência.")]
    ws.row_dimensions[22].height = 24
    ws.row_dimensions[23].height = 36
    for k, (n, bg, fg, d) in enumerate(leg):
        a, b = CL(2 + 3 * k), CL(4 + 3 * k)
        put(ws, f"{a}22:{b}22", n, f=font(10, True, fg), fl=fill(bg), al=A_CENTER, bd=BORDER)
        put(ws, f"{a}23:{b}23", d, f=font(9, False, "33414F"), al=A_CENTER)
    put(ws, "B24:P24", "Na Central de Pendências, LARANJA indica pendência em acompanhamento e CINZA indica informação "
        "não disponível ou cancelamento informado.", f=font(9, False, DGREY, True), al=A_LEFT)

    # como funciona
    put(ws, "B26:P26", "COMO A PLANILHA FUNCIONA", f=font(14, True, NAVY),
        bd=Border(bottom=Side(style="medium", color=NAVY)))
    regras = [
        "Sequência de preenchimento: CADASTROS → MTRs → ITENS_MTR → RECEBIMENTOS → CDFs → VÍNCULOS_CDF → PAINEL.",
        "Cada registro já tem um ID interno fixo (MTR-000001, IT-000001, CDF-000001). Ele não muda ao ordenar ou "
        "filtrar. Nas outras abas, escolha o MTR, o item ou o CDF pela lista suspensa.",
        "Você registra o que leu nos documentos e sistemas oficiais. A planilha não emite MTR nem CDF, não se conecta "
        "ao SINIR e não altera registros oficiais.",
        "O recebimento só conta como confirmado com RECEBIMENTO CONFERIDO? = Sim e um documento ou link de evidência. "
        "Recebimento confirmado não significa destinação comprovada: isso depende do CDF.",
        "Um CDF pode cobrir vários MTRs. Na aba VÍNCULOS_CDF, cada linha liga o CDF a um item, com a quantidade "
        "documentada. A planilha mostra quanto já tem vínculo e quanto ainda falta.",
        "Diferenças de quantidade geram alerta de revisão, nunca conclusão automática de irregularidade. Depois de "
        "analisar, marque DIFERENÇA ANALISADA? = Sim no item.",
        "MTR com cancelamento informado continua no histórico, mas sai dos cards 02 a 06 e das quantidades em kg. "
        "O termo de cancelamento é definido em CONFIGURAÇÕES.",
        "Prazos de alerta (15 dias como sugestão) são parâmetros internos, ajustáveis em CONFIGURAÇÕES. Não são "
        "prazos legais.",
        "Capacidade pronta: 500 MTRs, 1.000 itens, 1.000 recebimentos, 500 CDFs e 1.000 vínculos. Para volumes "
        "maiores, recomendamos usar um arquivo por ano.",
        "Compatibilidade: Excel 365 ou Excel 2019 em diante (usa a função MÁXIMOSES). Sem macros e sem internet. "
        "O funcionamento no Google Sheets não foi testado.",
        "As colunas automáticas ficam protegidas sem senha para evitar que fórmulas sejam apagadas. Filtros funcionam "
        "normalmente. Se precisar ordenar ou mudar a estrutura, use Revisão > Desproteger Planilha.",
    ]
    for k, t in enumerate(regras):
        r = 27 + k
        ws.row_dimensions[r].height = 30
        put(ws, f"B{r}", k + 1, f=font(11, True, WHITE), fl=fill(GREEN if k % 2 == 0 else BLUE), al=A_CENTER)
        put(ws, f"C{r}:P{r}", t, f=font(10, False, "1F2933"), fl=fill(GREY if k % 2 == 0 else WHITE), al=A_LEFT)
    r = 27 + len(regras) + 1
    ws.row_dimensions[r].height = 36
    put(ws, f"B{r}:P{r}", DISCLAIMER, f=font(10.5, True, NAVY, True), fl=fill(LBLUE), al=A_CENTER,
        bd=Border(top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY)))
    put(ws, f"B{r + 1}:P{r + 1}", f"Versão {VERSAO}", f=font(8.5, False, DGREY), al=A_CENTER)
    ws.print_area = f"A1:Q{r + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    set_protect(ws)
    return ws


# ---------------------------------------------------------------- montagem
def build(path, data=None):
    data = data or {}
    wb = Workbook()
    build_inicio(wb)
    build_painel(wb, data.get("PAINEL"))
    for sk in ("MTR", "IT", "REC", "CDF", "VIN"):
        build_data_sheet(wb, sk, data.get(sk))
    build_pendencias(wb)
    build_relatorios(wb, data.get("REL"))
    for sk in ("UN", "PA", "RS"):
        build_data_sheet(wb, sk, data.get(sk))
    build_config(wb)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.active = 0
    wb.save(path)
    return path


# ---------------------------------------------------------------- dados demonstrativos (fictícios)
def demo_data():
    D = dt.date
    UN = [
        dict(razao="Empresa Fictícia Exemplo S.A.", nome="Unidade Fictícia Alfa (Matriz)", cnpj="00.000.000/0001-00 (fictício)",
             uf="SP", mun="Cidade Exemplo", papel="Geradora", sistema="SINIR / MTR Nacional", resp="Analista Exemplo",
             sitcad="Ativo", obs="EXEMPLO FICTÍCIO"),
        dict(razao="Empresa Fictícia Exemplo S.A.", nome="Unidade Fictícia Beta (Filial)", cnpj="00.000.000/0002-00 (fictício)",
             uf="MG", mun="Cidade Modelo", papel="Geradora", sistema="SINIR / MTR Nacional", resp="Técnico Exemplo",
             sitcad="Ativo", obs="EXEMPLO FICTÍCIO"),
    ]
    PA = [
        dict(razao="Transportes Exemplo Ltda (fictícia)", fantasia="Transporte Exemplo", cnpj="11.111.111/0001-00 (fictício)",
             tipo="Transportador", estab="Base SP", mun="Cidade Exemplo", uf="SP", contato="Contato Fictício",
             email="contato@exemplo.invalid", tel="(00) 0000-0000", sitcad="Ativo (informado)",
             docaut="Licença fictícia nº 000", valid=D(2027, 6, 30), obs="EXEMPLO FICTÍCIO"),
        dict(razao="Logística Modelo Ltda (fictícia)", fantasia="Logística Modelo", tipo="Transportador", estab="Base MG",
             mun="Cidade Modelo", uf="MG", obs="EXEMPLO FICTÍCIO"),
        dict(razao="Recicladora Exemplo Ltda (fictícia)", fantasia="Recicladora Exemplo", tipo="Destinador",
             estab="Planta 1", mun="Cidade Exemplo", uf="SP", docaut="Licença fictícia nº 111", valid=D(2026, 9, 30),
             obs="EXEMPLO FICTÍCIO. Validade vencida de propósito para mostrar o alerta."),
        dict(razao="Central de Tratamento Exemplo (fictícia)", fantasia="Central Tratamento Exemplo", tipo="Destinador",
             estab="Unidade Sul", mun="Cidade Modelo", uf="MG", obs="EXEMPLO FICTÍCIO"),
        dict(razao="Ambiental Demo Ltda (fictícia)", fantasia="Ambiental Demo", tipo="Ambos", estab="Matriz",
             mun="Cidade Exemplo", uf="SP", obs="EXEMPLO FICTÍCIO"),
    ]
    T1, T2 = "Transporte Exemplo | Base SP", "Logística Modelo | Base MG"
    D1, D2, D3 = "Recicladora Exemplo | Planta 1", "Central Tratamento Exemplo | Unidade Sul", "Ambiental Demo | Matriz"
    RS = [
        dict(nome="Resíduo Exemplo A (fictício)", desc="Sucata metálica de exemplo", proc="Manutenção",
             cod="CÓD-FICT-A", ref="Referência fictícia", cls="Classificação fictícia 1", unpad="kg", obs="EXEMPLO FICTÍCIO"),
        dict(nome="Resíduo Exemplo B (fictício)", desc="Papel e papelão de exemplo", proc="Administrativo",
             cod="CÓD-FICT-B", ref="Referência fictícia", cls="Classificação fictícia 2", unpad="kg", obs="EXEMPLO FICTÍCIO"),
        dict(nome="Resíduo Exemplo C (fictício)", desc="Lodo de exemplo", proc="Estação de tratamento",
             cod="CÓD-FICT-C", ref="Referência fictícia", cls="Classificação fictícia 3", unpad="t", obs="EXEMPLO FICTÍCIO"),
        dict(nome="Resíduo Exemplo D (fictício)", desc="Óleo usado de exemplo", proc="Oficina",
             cod="CÓD-FICT-D", ref="Referência fictícia", cls="Classificação fictícia 4", unpad="kg", obs="EXEMPLO FICTÍCIO"),
    ]
    A, B = UN[0]["nome"], UN[1]["nome"]
    RA, RB, RC, RD = (r["nome"] for r in RS)
    S = "SINIR / MTR Nacional"
    mt = [  # (num, data emissão, data saída, unidade, transp, dest, situação, obs)
        ("FICT-2026-0001", D(2026, 5, 12), D(2026, 5, 13), A, T1, D1, "Recebido", "Fluxo completo conferido"),
        ("FICT-2026-0002", D(2026, 6, 3), D(2026, 6, 4), A, T1, D1, "Recebido", "Mesmo CDF do MTR FICT-2026-0001"),
        ("FICT-2026-0003", D(2026, 7, 8), D(2026, 7, 9), B, T2, D2, "Recebido", "Dois resíduos: 800 kg + 1,2 t"),
        ("FICT-2026-0004", D(2026, 8, 5), D(2026, 8, 6), B, T2, D2, "Recebido", "Diferença de pesagem 1.000 x 980 kg"),
        ("FICT-2026-0005", D(2026, 8, 20), D(2026, 8, 21), A, T1, D3, "Recebido", "CDF cobre só 600 de 1.000 kg"),
        ("FICT-2026-0006", D(2026, 10, 14), D(2026, 10, 15), A, T1, D1, "Emitido", "Aguardando recebimento"),
        ("FICT-2026-0007", D(2026, 8, 25), D(2026, 8, 26), B, T2, D2, "Emitido", "Sem recebimento há mais tempo"),
        ("FICT-2026-0008", D(2026, 9, 10), D(2026, 9, 11), A, D3, D3, "Recebido", "Recebido, CDF ainda não localizado"),
        ("FICT-2026-0009", D(2026, 7, 15), None, A, T1, D1, "Cancelado", "Cancelamento informado no sistema"),
        ("FICT-2026-0010", D(2026, 10, 1), None, B, None, D2, "Emitido", "Transportador não preenchido"),
        ("FICT-2026-0011", D(2026, 9, 28), D(2026, 9, 29), B, T2, D1, "Recebido", "Recebimento lançado, falta conferir"),
    ]
    MTR, labels = [], []
    for k, (num, de, ds, un, tr, de_, sit, obs) in enumerate(mt):
        MTR.append(dict(num=num, sis=S, uf="SP" if un == A else "MG", dtem=de, dtsai=ds, unid=un, transp=tr, dest=de_,
                        ufdest="SP", sitsis=sit, dtcons=D(2026, 10, 16), link=f"C:/Exemplo/MTR/{num}.pdf",
                        resp="Analista Exemplo", obs="EXEMPLO FICTÍCIO. " + obs))
        labels.append(f"MTR-{k + 1:06d} | nº {num} | {un}")
    it = [  # (mtr idx, resíduo, qtd, un)
        (0, RA, 500, "kg"), (1, RA, 300, "kg"), (2, RB, 800, "kg"), (2, RC, 1.2, "t"), (3, RB, 1000, "kg"),
        (4, RD, 1000, "kg"), (5, RA, 450, "kg"), (6, RC, 2.5, "t"), (7, RB, 650, "kg"), (8, RA, 200, "kg"),
        (9, RC, 3, "t"), (10, RD, 380, "kg"),
    ]
    IT, ilabels = [], []
    for k, (m, r, q, u) in enumerate(it):
        IT.append(dict(mtrref=labels[m], res=r, qty=q, un=u, obs="EXEMPLO FICTÍCIO"))
        ilabels.append(f"IT-{k + 1:06d} | MTR {mt[m][0]} | {r}")
    rec = [  # (item idx, data, qtd, un, conferido)
        (0, D(2026, 5, 14), 500, "kg", "Sim"), (1, D(2026, 6, 5), 300, "kg", "Sim"), (2, D(2026, 7, 10), 800, "kg", "Sim"),
        (3, D(2026, 7, 10), 1.2, "t", "Sim"), (4, D(2026, 8, 7), 980, "kg", "Sim"), (5, D(2026, 8, 22), 1000, "kg", "Sim"),
        (8, D(2026, 9, 12), 650, "kg", "Sim"), (11, D(2026, 9, 30), 380, "kg", "Não"),
    ]
    REC = [dict(itemref=ilabels[i], dt=d, qty=q, un=u, sitsis="Recebido", conf=c,
                dtconf=d + dt.timedelta(days=2) if c == "Sim" else None,
                evid="Consulta ao sistema (exemplo)" if c == "Sim" else None, resp="Analista Exemplo",
                obs="EXEMPLO FICTÍCIO") for i, d, q, u, c in rec]
    cdf = [  # (num, emissão, período, destinador, modalidade, localizado, conferido)
        ("CDF-FICT-0101", D(2026, 6, 30), "maio a junho/2026", D1, "Reciclagem", "Sim", "Sim"),
        ("CDF-FICT-0202", D(2026, 8, 1), "julho/2026", D2, "Tratamento físico-químico", "Sim", "Não"),
        ("CDF-FICT-0303", D(2026, 9, 15), "agosto/2026", D3, "Rerrefino", "Sim", "Sim"),
    ]
    CDF, clabels = [], []
    for k, (num, de, per, d_, mod, loc, conf) in enumerate(cdf):
        CDF.append(dict(num=num, sis=S, dtem=de, periodo=per, dest=d_, modal=mod, sitsis="Emitido", loc=loc, conf=conf,
                        dtconf=D(2026, 10, 2) if conf == "Sim" else None, link=f"C:/Exemplo/CDF/{num}.pdf",
                        resp="Analista Exemplo", obs="EXEMPLO FICTÍCIO"))
        clabels.append(f"CDF-{k + 1:06d} | CDF {num} | {d_}")
    vin = [  # (cdf, item, qtd, un, conferido)
        (0, 0, 500, "kg", "Sim"), (0, 1, 300, "kg", "Sim"), (1, 2, 800, "kg", "Não"), (1, 3, 1.2, "t", "Não"),
        (2, 5, 600, "kg", "Sim"),
    ]
    VIN = [dict(cdfref=clabels[c], itemref=ilabels[i], qty=q, un=u, conf=cf, obs="EXEMPLO FICTÍCIO")
           for c, i, q, u, cf in vin]
    return dict(UN=UN, PA=PA, RS=RS, MTR=MTR, IT=IT, REC=REC, CDF=CDF, VIN=VIN)


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    build(f"{out}/Bonus02_Planilha_Controle_MTR_CDF.xlsx")
    build(f"{out}/Bonus02_Planilha_MTR_CDF_Exemplo.xlsx", demo_data())
    print("ok")
