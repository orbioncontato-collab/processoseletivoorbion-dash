"""Leitura dos resultados recalculados (cópia gerada pelo LibreOffice)."""
import datetime as dt
import openpyxl

ERROS = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:", "#ERR")


def _norm(v):
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0:
            return v.date()
    return v


class Resultado:
    def __init__(self, path, tabs):
        self.wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
        self.tabs = tabs

    def tabela(self, name, only_active=True):
        t = self.tabs[name]
        ws = self.wb[t.sheet]
        out = []
        for r in range(t.r0 + 1, t.r1 + 2):
            rec = {}
            for j, c in enumerate(t.cols):
                rec[c.h] = _norm(ws.cell(r, t.c0 + j + 1).value)
            rec["_row"] = r
            out.append(rec)
        if only_active:
            key = t.cols[0].h
            out = [x for x in out if any(x.get(c.h) not in (None, "") for c in t.cols if c.kind == "in")]
        return out

    def erros(self):
        found = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v.startswith(ERROS):
                        found.append((ws.title, c.coordinate, v))
        return found

    def valores_cache(self):
        """Dicionário (aba, linha base 0, coluna base 0) -> valor, para gravar como resultado inicial."""
        cache = {}
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None:
                        continue
                    if isinstance(v, dt.datetime):
                        v = (v - dt.datetime(1899, 12, 30)).total_seconds() / 86400.0
                    elif isinstance(v, dt.date):
                        v = (v - dt.date(1899, 12, 30)).days
                    elif isinstance(v, bool):
                        v = int(v)
                    cache[(ws.title, c.row - 1, c.column - 1)] = v
        return cache
