"""Conferência dos resultados recalculados contra o modelo independente e os valores esperados."""
import datetime as dt
import zipfile
import re

import modelo
import dados_teste

PER_COLS = {"Início do aquisitivo": "_ini_aq", "Término do aquisitivo": "_fim_aq", "Início do concessivo": "_ini_conc",
            "Término do concessivo": "_fim_conc", "Início padrão": "_ini_pad", "Término padrão": "_fim_pad",
            "Direito pela tabela": "_tabela", "Dias de direito": "_dir", "Gozo confirmado": "_gozo_conf",
            "Abono confirmado": "_abono_conf", "Gozo reservado": "_gozo_res", "Abono reservado": "_abono_res",
            "Saldo não gozado": "_sng", "Saldo disponível": "_sd", "Demanda solicitada": "_demanda",
            "Quantidade de parcelas": "_parcelas", "Téc frac": "_frac", "Téc fracionado": "_fracionado",
            "Téc abono qtd": "_abono_qtd", "Téc abono excesso": "_abono_exc", "Dias até o fim do concessivo": "_dias_fim",
            "Último início para o saldo": "_ult_ini"}
PROG_COLS = {"Término": "_fim", "Téc efeito": "_efeito", "Téc efeito abono": "_efeito_abono", "Téc ativa": "_ativa",
             "Téc cap": "_cap", "Téc conflito": "_conflito"}


def _eq(a, b):
    if a in ("", None) and b in ("", None):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-9
    return a == b


def linhas(R, name, n):
    t = R.tabs[name]
    ws = R.wb[t.sheet]
    out = []
    for i in range(n):
        r = t.r0 + 1 + i
        rec = {c.h: _val(ws.cell(r, t.c0 + j + 1).value) for j, c in enumerate(t.cols)}
        out.append(rec)
    return out


def _val(v):
    if isinstance(v, dt.datetime) and v.hour == 0 and v.minute == 0:
        return v.date()
    return v


def comparar_modelo(R, data, ref):
    M = modelo.Modelo(data, ref)
    falhas, n = [], 0
    per_x = linhas(R, "T_Per", len(data.get("T_Per", [])))
    for p, x in zip(M.per, per_x):
        for col, key in PER_COLS.items():
            n += 1
            if not _eq(x.get(col), p.get(key)):
                falhas.append(("T_Per", p["ID do período"], col, x.get(col), p.get(key)))
    prog_x = linhas(R, "T_Prog", len(data.get("T_Prog", [])))
    for g, x in zip(M.prog, prog_x):
        for col, key in PROG_COLS.items():
            n += 1
            if not _eq(x.get(col), g.get(key)):
                falhas.append(("T_Prog", g["ID da programação"], col, x.get(col), g.get(key)))
        if g.get("_vedado") is not None:
            n += 1
            ved = 1 if (x.get("Téc fer ini") or 0) > 0 or x.get("Téc dsr ini") == 1 else 0
            if ved != g["_vedado"]:
                falhas.append(("T_Prog", g["ID da programação"], "vedação de início", ved, g["_vedado"]))
        if g.get("_retorno") is not None:
            n += 1
            rx = x.get("Retorno previsto segundo escala")
            exp = g["_retorno"]
            # escala variável ou não validada é tratada antes pela planilha
            if not _eq(rx, exp) and rx not in ("Conferir escala",):
                falhas.append(("T_Prog", g["ID da programação"], "Retorno", rx, exp))
    return n, falhas, M


def checar_esperado(R, data):
    falhas, n = [], 0
    per_x = {x["ID do período"]: x for x in linhas(R, "T_Per", len(data["T_Per"]))}
    prog_x = {}
    for x in linhas(R, "T_Prog", len(data["T_Prog"])):
        prog_x.setdefault(x["ID da programação"], x)
    for (tb, key, col), exp in dados_teste.ESPERADO.items():
        n += 1
        x = (per_x if tb == "T_Per" else prog_x).get(key, {})
        if not _eq(x.get(col), exp):
            falhas.append((tb, key, col, x.get(col), exp))
    for key, trechos in dados_teste.ALERTA_CONTEM.items():
        txt = prog_x.get(key, {}).get("Alertas da parcela") or ""
        for t in trechos:
            n += 1
            if t not in txt:
                falhas.append(("T_Prog", key, "contém", txt, t))
    for key, trechos in dados_teste.ALERTA_NAO_CONTEM.items():
        txt = prog_x.get(key, {}).get("Alertas da parcela") or ""
        for t in trechos:
            n += 1
            if t in txt:
                falhas.append(("T_Prog", key, "não contém", txt, t))
    return n, falhas


def checar_vazias(R, data, amostra=25):
    """Linhas vazias não podem mostrar datas, alertas ou zeros nas colunas visíveis."""
    falhas, n = [], 0
    for name in ("T_Func", "T_Per", "T_Prog", "T_Prov", "T_Oc", "T_Est"):
        t = R.tabs[name]
        ws = R.wb[t.sheet]
        ini = len(data.get(name, []))
        for i in list(range(ini, min(t.n, ini + amostra))) + [t.n - 1]:
            r = t.r0 + 1 + i
            for j, c in enumerate(t.cols):
                if c.kind != "calc":
                    continue
                n += 1
                v = ws.cell(r, t.c0 + j + 1).value
                if v not in (None, ""):
                    falhas.append((name, r, c.h, v))
    return n, falhas


def checar_estrutura(path, tabs):
    """Confere, no XML do arquivo, colunas calculadas das tabelas, proteção e validações."""
    out = []
    z = zipfile.ZipFile(path)
    names = z.namelist()
    tables = [n for n in names if n.startswith("xl/tables/")]
    calc_ok = 0
    for tn in tables:
        x = z.read(tn).decode("utf8")
        calc_ok += x.count("<calculatedColumnFormula>")
    out.append(("Tabelas estruturadas no arquivo", len(tables)))
    out.append(("Colunas com fórmula de coluna calculada", calc_ok))
    sheets = [n for n in names if re.match(r"xl/worksheets/sheet\d+\.xml", n)]
    prot = sum(1 for s in sheets if b"<sheetProtection" in z.read(s)[:200000] or b"<sheetProtection" in z.read(s))
    out.append(("Abas protegidas", f"{prot} de {len(sheets)}"))
    dv = sum(z.read(s).count(b"<dataValidation ") for s in sheets)
    out.append(("Regras de validação de dados", dv))
    wbx = z.read("xl/workbook.xml").decode("utf8")
    out.append(("Nomes definidos", wbx.count("<definedName ")))
    out.append(("Áreas de impressão definidas", wbx.count("_xlnm.Print_Area")))
    out.append(("Links externos", sum(1 for n in names if "externalLink" in n)))
    out.append(("Macros (vbaProject)", sum(1 for n in names if "vbaProject" in n)))
    return out


def checar_protecao(path, tabs):
    """Amostra: células de entrada desbloqueadas e células calculadas bloqueadas."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=False)
    falhas, n = [], 0
    for name in ("T_Func", "T_Per", "T_Prog", "T_Prov", "T_Oc", "T_Est"):
        t = tabs[name]
        ws = wb[t.sheet]
        if not ws.protection.sheet:
            falhas.append((t.sheet, "aba sem proteção"))
        if ws.protection.autoFilter:
            falhas.append((t.sheet, "filtro bloqueado"))
        for r in (t.r0 + 1, t.r0 + 2, t.r1 + 1):
            for j, c in enumerate(t.cols):
                n += 1
                locked = ws.cell(r, t.c0 + j + 1).protection.locked
                if (c.kind == "in") == bool(locked):
                    falhas.append((t.sheet, r, c.h, "bloqueada" if locked else "desbloqueada"))
    return n, falhas


def checar_kpis(R, M):
    """Indicadores do painel conferidos com o modelo independente."""
    ws = R.wb["09 Painel"]
    vals = {}
    for row in ws.iter_rows(min_row=1, max_row=22, min_col=2, max_col=10):
        for c in row:
            if isinstance(c.value, str) and c.row + 1 <= ws.max_row:
                vals[c.value] = ws.cell(c.row + 1, c.column).value
    ref = M.ref
    ativos = sum(1 for f in M.func.values() if f.get("Situação do vínculo") == "Ativo")
    ferias = len({g["ID do vínculo"] for g in M.prog if g["_cap"] and g["Data de início"] <= ref <= g["_fim"]
                  and M.func[g["ID do vínculo"]].get("Situação do vínculo") != "Desligado"})
    aprov30 = sum(1 for g in M.prog if g["_S"] == 3 and g.get("Data de início") and ref < g["Data de início"] <= ref + dt.timedelta(days=30))
    confl = sum(1 for g in M.prog if g.get("_conflito") == 1)
    esperado = {"Vínculos ativos": ativos, "Pessoas em férias na data de referência": ferias,
                "Programações aprovadas para os próximos 30 dias": aprov30, "Programações com conflito de equipe": confl}
    falhas = [(k, vals.get(k), v) for k, v in esperado.items() if vals.get(k) != v]
    return len(esperado), falhas, esperado
