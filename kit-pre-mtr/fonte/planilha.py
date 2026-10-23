# -*- coding: utf-8 -*-
"""Gera a versão em planilha (.xlsx) do Kit Pré-MTR a partir de formularios.py.

Uso:  python3 planilha.py
Saída: ../Kit_Pre_MTR_15_Formularios.xlsx

Cada formulário vira uma aba (F01 a F15) com os mesmos blocos e campos do
HTML. A aba Início traz o índice com links e o andamento do preenchimento.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from formularios import CATEGORIAS, DATALISTS, FORMULARIOS

BASE = Path(__file__).resolve().parent
SAIDA = BASE.parent / "Kit_Pre_MTR_15_Formularios.xlsx"
VERSAO = "1.0"
LETRAS = "ABCDEFGH"
AVISO = ("Os formulários deste kit são modelos internos de organização e registro. Não substituem documentos "
         "oficiais, avaliações técnicas ou procedimentos exigidos pelos órgãos competentes.")

FONTE = "Arial"
C1, C2 = 2, 13          # área útil: colunas B a M (grade de 12 colunas)
PET, VERDE, GRAF, CINZA, OURO = "103C39", "17806C", "1E293B", "64748B", "E6B566"
CAMPO_FILL = PatternFill("solid", fgColor="FBFDF6")
BORDA_CAMPO = Side(style="thin", color="9DB5AE")
BORDA_LEVE = Side(style="thin", color="D5E2DD")

CAT = {c["num"]: c for c in CATEGORIAS}
for i, f in enumerate(FORMULARIOS, 1):
    f["n"], f["cod"] = i, f"F{i:02d}"

SIGLAS = {"tri": ["Conferido", "Pendente", "Não aplicável"],
          "vnc": ["Verificado", "Necessita conferência", "Não aplicável"]}


def hexcor(c):
    return c.lstrip("#").upper()


def fonte(size=10, bold=False, color=GRAF, italic=False):
    return Font(name=FONTE, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=hexcor(c))


# ------------------------------------------------------------------ listas

class Listas:
    """Aba oculta com as opções das listas suspensas (evita limites e vírgulas)."""

    def __init__(self, wb):
        self.ws = wb.create_sheet("Listas")
        self.col = 0
        self.cache = {}

    def ref(self, nome, opcoes):
        chave = tuple(opcoes)
        if chave in self.cache:
            return self.cache[chave]
        self.col += 1
        L = get_column_letter(self.col)
        self.ws.cell(1, self.col, nome).font = fonte(9, True)
        for i, o in enumerate(opcoes, 2):
            self.ws.cell(i, self.col, o).font = fonte(9)
        r = f"Listas!${L}$2:${L}${len(opcoes) + 1}"
        self.cache[chave] = r
        return r


# ------------------------------------------------------------------ desenho

def caixa(ws, r1, c1, r2, c2, borda=BORDA_CAMPO, preench=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cel = ws.cell(r, c)
            cel.border = Border(
                left=borda if c == c1 else None, right=borda if c == c2 else None,
                top=borda if r == r1 else None, bottom=borda if r == r2 else None)
            if preench:
                cel.fill = preench


def mesclar(ws, r, c1, c2, valor=None, f=None, al=None, r2=None):
    r2 = r2 or r
    if c2 > c1 or r2 > r:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r2, end_column=c2)
    cel = ws.cell(r, c1)
    if valor is not None:
        cel.value = valor
    if f:
        cel.font = f
    if al:
        cel.alignment = al
    return cel


AL_ROT = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
AL_CAMPO = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=0)
AL_AREA = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


class Folha:
    """Escreve um formulário em uma aba, acompanhando a linha atual."""

    def __init__(self, wb, f, listas, dvs):
        self.f, self.listas = f, listas
        self.cat = CAT[f["cat"]]
        self.cor, self.tint = hexcor(self.cat["cor"]), hexcor(self.cat["tint"])
        self.ws = wb.create_sheet(f["cod"])
        self.ws.sheet_properties.tabColor = self.cor
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_view.zoomScale = 110
        self.r = 1
        self.dvs = {}
        self.entradas = []   # faixas de células preenchíveis (para o andamento)
        self.n_entradas = 0
        self.quebras = []
        self.calc_ref = None
        for c in range(1, 15):
            self.ws.column_dimensions[get_column_letter(c)].width = 8.4
        self.ws.column_dimensions["A"].width = 1.6
        self.ws.column_dimensions["N"].width = 1.6

    # -------------------------------------------------- validações
    def dv(self, chave, **kw):
        if chave not in self.dvs:
            d = DataValidation(**kw)
            d.showInputMessage = True
            d.showErrorMessage = True
            self.ws.add_data_validation(d)
            self.dvs[chave] = d
        return self.dvs[chave]

    def validar(self, cel, tipo, opcoes=None, lista=None, rotulo=""):
        if tipo == "date":
            d = self.dv("date", type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2100,12,31)",
                        errorTitle="Data inválida", error="Digite uma data no formato dd/mm/aaaa, por exemplo 23/10/2026.",
                        promptTitle="Data", prompt="Formato dd/mm/aaaa")
            cel.number_format = "dd/mm/yyyy"
        elif tipo == "time":
            d = self.dv("time", type="time", operator="between", formula1="0", formula2="0.999988",
                        errorTitle="Horário inválido", error="Digite um horário no formato hh:mm, por exemplo 14:30.",
                        promptTitle="Horário", prompt="Formato hh:mm")
            cel.number_format = "hh:mm"
        elif tipo == "num":
            d = self.dv("num", type="decimal", operator="greaterThanOrEqual", formula1="0",
                        errorTitle="Valor inválido", error="Digite apenas números, por exemplo 1250,5. Informe a unidade no campo ao lado.",
                        promptTitle="Quantidade", prompt="Somente números. A unidade fica em outro campo.")
            cel.number_format = "#,##0.###"
        elif tipo in ("opcoes", "marca"):
            ref = self.listas.ref(rotulo or "opções", opcoes)
            d = self.dv("o:" + ref, type="list", formula1=ref, allow_blank=True,
                        errorTitle="Opção inválida", error="Escolha uma das opções da lista.",
                        promptTitle="Escolha na lista", prompt="Clique na seta para escolher.")
        elif tipo == "sugestao":
            ref = self.listas.ref(lista, DATALISTS[lista])
            d = self.dv("s:" + ref, type="list", formula1=ref, allow_blank=True, errorStyle="information",
                        errorTitle="Valor fora da lista", error="Este valor não está na lista de sugestões. Clique em OK para manter.",
                        promptTitle="Sugestões", prompt="Escolha na lista ou digite outro valor.")
        else:
            return
        d.add(cel.coordinate)

    def entrada(self, r, c1, c2):
        self.entradas.append(f"{get_column_letter(c1)}{r}:{get_column_letter(c2)}{r}")
        self.n_entradas += 1

    # -------------------------------------------------- cabeçalho
    def cabecalho(self):
        ws, f = self.ws, self.f
        ws.row_dimensions[1].height = 18
        cel = mesclar(ws, 1, 11, 13, "‹ Voltar ao índice", fonte(8, True, VERDE), Alignment(horizontal="right", vertical="center"))
        cel.hyperlink = "#'Início'!A1"
        ws.row_dimensions[2].height = 17
        ws.row_dimensions[3].height = 26
        ws.row_dimensions[4].height = 6
        for r in (2, 3, 4):
            for c in range(C1, C2 + 1):
                ws.cell(r, c).fill = fill(self.cor)
        mesclar(ws, 2, C1, 9, f"FORMULÁRIO {f['n']:02d} · CATEGORIA: {self.cat['nome'].upper()}", fonte(8, True, "FFFFFF"),
                Alignment(vertical="bottom", indent=1))
        mesclar(ws, 3, C1, 9, f["titulo"], fonte(15, True, "FFFFFF"), Alignment(vertical="center", indent=1, wrap_text=True))
        # bloco de identificação do modelo
        for c in range(10, 14):
            for r in (2, 3):
                ws.cell(r, c).fill = fill("FFFFFF")
        mesclar(ws, 2, 10, 10, "Código", fonte(7, True, CINZA), Alignment(vertical="bottom", indent=1))
        mesclar(ws, 2, 11, 11, "Versão", fonte(7, True, CINZA), Alignment(vertical="bottom"))
        mesclar(ws, 2, 12, 13, "Registro nº", fonte(7, True, CINZA), Alignment(vertical="bottom"))
        mesclar(ws, 3, 10, 10, f["cod"], fonte(11, True, self.cor), Alignment(vertical="center", indent=1))
        mesclar(ws, 3, 11, 11, VERSAO, fonte(11, True, self.cor), Alignment(vertical="center"))
        reg = mesclar(ws, 3, 12, 13, None, fonte(10), AL_CAMPO)
        caixa(ws, 3, 12, 3, 13, preench=CAMPO_FILL)
        self.entrada(3, 12, 13)
        self.r = 5
        if f.get("subtitulo"):
            ws.row_dimensions[self.r].height = 16
            mesclar(ws, self.r, C1, C2, f["subtitulo"], fonte(9, True, "9A6212"), Alignment(vertical="center"))
            self.r += 1
        ws.row_dimensions[self.r].height = 16
        cel = mesclar(ws, self.r, C1, C2, None, None, Alignment(vertical="center", wrap_text=True))
        cel.value = f"Finalidade de uso: {f['finalidade']}"
        cel.font = fonte(9, False, GRAF)
        self.r += 1
        ori = mesclar(ws, self.r, C1, C2, "ⓘ  " + f["orientacao"], fonte(8.5, False, "334155"),
                      Alignment(vertical="center", wrap_text=True, indent=1))
        ws.row_dimensions[self.r].height = 28 if len(f["orientacao"]) > 120 else 18
        for c in range(C1, C2 + 1):
            ws.cell(self.r, c).fill = fill(self.tint)
        self.r += 1
        if f.get("aviso"):
            ws.row_dimensions[self.r].height = 18
            mesclar(ws, self.r, C1, C2, "⚠  " + f["aviso"], fonte(8.5, True, "7A4A0B"), Alignment(vertical="center", indent=1))
            for c in range(C1, C2 + 1):
                ws.cell(self.r, c).fill = fill("FFF6E4")
            caixa(ws, self.r, C1, self.r, C2, borda=Side(style="thin", color="F0D9AE"), preench=fill("FFF6E4"))
            self.r += 1

    # -------------------------------------------------- blocos
    def bloco(self, i, titulo, campos, quebra=False):
        ws = self.ws
        if quebra:
            self.quebras.append(self.r - 1)
        ws.row_dimensions[self.r].height = 8
        self.r += 1
        ws.row_dimensions[self.r].height = 20
        ws.cell(self.r, C1).value = LETRAS[i]
        ws.cell(self.r, C1).font = fonte(10, True, "FFFFFF")
        ws.cell(self.r, C1).fill = fill(self.cor)
        ws.cell(self.r, C1).alignment = AL_CENTRO
        mesclar(ws, self.r, C1 + 1, C2, titulo, fonte(11, True, PET), Alignment(vertical="center", indent=1))
        for c in range(C1, C2 + 1):
            ws.cell(self.r, c).border = Border(bottom=Side(style="medium", color=self.tint))
        self.r += 1
        self.campos(campos, C1, 12)

    def campos(self, campos, col0, largura, escala=1.0):
        linha, soma = [], 0
        for x in campos:
            if x["t"] in ("tabela", "checklist", "sub", "calc", "nota", "checks"):
                self.flush(linha, col0)
                linha, soma = [], 0
                getattr(self, "w_" + x["t"])(x, col0, largura)
                continue
            s = max(1, round(x["s"] * escala))
            if soma + s > largura:
                self.flush(linha, col0)
                linha, soma = [], 0
            linha.append((x, soma, s))
            soma += s
        self.flush(linha, col0)

    def flush(self, linha, col0):
        if not linha:
            return
        ws = self.ws
        rl, ri = self.r, self.r + 1
        ws.row_dimensions[rl].height = 14
        altura = 21
        for x, off, s in linha:
            c1, c2 = col0 + off, col0 + off + s - 1
            rot = x["l"]
            if x["t"] == "radio":
                rot += "  (escolha na lista)"
            if x["t"] == "check":
                rot = ""
            mesclar(ws, rl, c1, c2, rot, fonte(7.5, True, "475569"), AL_ROT)
            if x["t"] == "area":
                altura = max(altura, 15 * x["rows"] + 6)
            elif x["t"] == "radio" and max(len(o) for o in x["o"]) > 2 * 8 * s:
                altura = max(altura, 30)
        ws.row_dimensions[ri].height = altura
        for x, off, s in linha:
            c1, c2 = col0 + off, col0 + off + s - 1
            if x["t"] == "check":
                # caixa de marcação + texto
                cel = ws.cell(ri, c1)
                cel.font = fonte(11, True, self.cor)
                cel.alignment = AL_CENTRO
                caixa(ws, ri, c1, ri, c1, preench=CAMPO_FILL)
                self.validar(cel, "marca", ["X"], rotulo="marcação")
                self.entrada(ri, c1, c1)
                mesclar(ws, ri, c1 + 1, c2, x["l"], fonte(9, False, "334155"), Alignment(vertical="center", indent=1, wrap_text=True))
                continue
            al = AL_AREA if x["t"] == "area" else AL_CAMPO
            cel = mesclar(ws, ri, c1, c2, None, fonte(10), al)
            caixa(ws, ri, c1, ri, c2, preench=CAMPO_FILL)
            self.entrada(ri, c1, c2)
            if x["t"] in ("date", "time", "num"):
                self.validar(cel, x["t"])
            elif x["t"] == "list":
                self.validar(cel, "sugestao", lista=x["list"])
            elif x["t"] == "radio":
                self.validar(cel, "opcoes", x["o"], rotulo=x["l"])
            if x["t"] in ("text", "list", "area") and x.get("ph"):
                cel.comment = None
            self.guarda(x, cel)
        self.r += 2

    def guarda(self, x, cel):
        k = x.get("k")
        if k in ("q1", "q2", "u1", "u2"):
            setattr(self, "ref_" + k, cel.coordinate)

    # -------------------------------------------------- componentes especiais
    def w_nota(self, x, col0, largura):
        self.ws.row_dimensions[self.r].height = 16
        mesclar(self.ws, self.r, col0, col0 + largura - 1, x["texto"], fonte(8, False, CINZA, True), Alignment(wrap_text=True, vertical="center"))
        self.r += 1

    def w_checks(self, x, col0, largura):
        ws = self.ws
        ws.row_dimensions[self.r].height = 14
        mesclar(ws, self.r, col0, col0 + largura - 1, x["l"] + "  (marque com X as opções que se aplicam)", fonte(7.5, True, "475569"), AL_ROT)
        self.r += 1
        meia = largura // 2
        for i in range(0, len(x["o"]), 2):
            ws.row_dimensions[self.r].height = 20
            for j, o in enumerate(x["o"][i:i + 2]):
                c = col0 + j * meia
                cel = ws.cell(self.r, c)
                cel.font = fonte(11, True, self.cor)
                cel.alignment = AL_CENTRO
                caixa(ws, self.r, c, self.r, c, preench=CAMPO_FILL)
                self.validar(cel, "marca", ["X"], rotulo="marcação")
                self.entrada(self.r, c, c)
                mesclar(ws, self.r, c + 1, c + meia - 1, o, fonte(9, False, "334155"), Alignment(vertical="center", indent=1, wrap_text=True))
            self.r += 1

    def w_sub(self, x, col0, largura):
        ws = self.ws
        ws.row_dimensions[self.r].height = 17
        mesclar(ws, self.r, col0, col0 + largura - 1, x["titulo"].upper(), fonte(8, True, self.cor), Alignment(vertical="bottom"))
        for c in range(col0, col0 + largura):
            ws.cell(self.r, c).border = Border(bottom=BORDA_LEVE)
        self.r += 1
        # subgrupos foram desenhados para meia largura: a grade interna é dividida por dois
        self.campos(x["campos"], col0, largura, escala=0.5)

    def w_calc(self, x, col0, largura):
        ws = self.ws
        q1, q2, u1, u2 = self.ref_q1, self.ref_q2, self.ref_u1, self.ref_u2
        ws.row_dimensions[self.r].height = 14
        mesclar(ws, self.r, col0, col0 + 5, "Diferença, quando as unidades forem comparáveis", fonte(7.5, True, PET), AL_ROT)
        mesclar(ws, self.r, col0 + 6, col0 + 7, "Unidade", fonte(7.5, True, PET), AL_ROT)
        mesclar(ws, self.r, col0 + 8, col0 + 11, "% sobre a quantidade informada", fonte(7.5, True, PET), AL_ROT)
        self.r += 1
        r = self.r
        ws.row_dimensions[r].height = 22
        dif = f"{get_column_letter(col0)}{r}"
        mesclar(ws, r, col0, col0 + 5, (
            f'=IF(OR({q1}="",{q2}=""),"",IF(OR(TRIM({u1})="",TRIM({u2})=""),"Não calculada: informe as duas unidades",'
            f'IF(LOWER(TRIM({u1}))<>LOWER(TRIM({u2})),"Não calculada: unidades diferentes",{q2}-{q1})))'),
            fonte(10, True), AL_CAMPO)
        ws[dif].number_format = '+#,##0.###;-#,##0.###;0'
        mesclar(ws, r, col0 + 6, col0 + 7, f'=IF(ISNUMBER({dif}),TRIM({u2}),"")', fonte(10, True), AL_CAMPO)
        pct = mesclar(ws, r, col0 + 8, col0 + 11, f'=IF(AND(ISNUMBER({dif}),{q1}<>0),{dif}/{q1},"")', fonte(10, True), AL_CAMPO)
        pct.number_format = '+0.0%;-0.0%;0.0%'
        for a, b in ((col0, col0 + 5), (col0 + 6, col0 + 7), (col0 + 8, col0 + 11)):
            caixa(ws, r, a, r, b, borda=Side(style="thin", color=self.cor), preench=fill(self.tint))
        self.r += 1
        ws.row_dimensions[self.r].height = 24
        mesclar(ws, self.r, col0, col0 + largura - 1,
                "Cálculo automático: quantidade verificada menos quantidade informada, e o percentual sobre a quantidade informada. "
                "Só é feito quando as duas unidades de medida são iguais. Peso, volume e número de volumes não são convertidos.",
                fonte(7.5, False, CINZA, True), Alignment(wrap_text=True, vertical="center"))
        self.r += 1
        self.calc_ref = dif

    def _spans(self, larguras, total, minimos):
        soma = sum(larguras)
        bruto = [w / soma * total for w in larguras]
        sp = [max(m, int(b)) for b, m in zip(bruto, minimos)]
        while sum(sp) < total:
            i = max(range(len(sp)), key=lambda k: bruto[k] - sp[k])
            sp[i] += 1
        while sum(sp) > total:
            i = max((k for k in range(len(sp)) if sp[k] > minimos[k]), key=lambda k: sp[k] - bruto[k])
            sp[i] -= 1
        return sp

    def w_tabela(self, x, col0, largura):
        ws = self.ws
        cols = x["cols"]
        minimos = [2 if t in SIGLAS else 1 for _, _, t, _ in cols]
        sp = self._spans([w for *_, w in cols], largura - 1, minimos)
        # cabeçalho
        rh = self.r
        altura_cab = 26 if any(len(c[1]) > 9 * s for c, s in zip(cols, sp)) else 18
        ws.row_dimensions[rh].height = altura_cab
        cel = ws.cell(rh, col0, "Nº")
        cel.font, cel.alignment, cel.fill = fonte(7.5, True, PET), AL_CENTRO, fill(self.tint)
        c = col0 + 1
        pos = []
        for (k, rot, t, _), s in zip(cols, sp):
            mesclar(ws, rh, c, c + s - 1, rot, fonte(7.5, True, PET), AL_CENTRO if t in ("check", "tri", "vnc", "date", "num") else Alignment(vertical="center", wrap_text=True, indent=0))
            for cc in range(c, c + s):
                ws.cell(rh, cc).fill = fill(self.tint)
            pos.append((c, c + s - 1, t))
            c += s
        caixa(ws, rh, col0, rh, col0 + largura - 1, borda=BORDA_CAMPO)
        for (a, b, _t) in pos:
            caixa(ws, rh, a, rh, b, borda=BORDA_CAMPO, preench=fill(self.tint))
        caixa(ws, rh, col0, rh, col0, borda=BORDA_CAMPO, preench=fill(self.tint))
        self.r += 1
        n = max(x["imp"], x["min"], len(x["pre"]))
        for i in range(n):
            r = self.r
            ws.row_dimensions[r].height = 20
            cel = ws.cell(r, col0, i + 1)
            cel.font, cel.alignment = fonte(8, True, CINZA), AL_CENTRO
            caixa(ws, r, col0, r, col0)
            pre = x["pre"][i] if i < len(x["pre"]) else None
            for j, (a, b, t) in enumerate(pos):
                v = pre if (j == 0 and pre) else None
                al = AL_CENTRO if t in ("check", "tri", "vnc") else AL_CAMPO
                cel = mesclar(ws, r, a, b, v, fonte(9, False, "334155" if v else GRAF), al)
                caixa(ws, r, a, r, b, preench=CAMPO_FILL)
                if t in SIGLAS:
                    self.validar(cel, "opcoes", SIGLAS[t], rotulo="situação")
                elif t == "check":
                    cel.font = fonte(11, True, self.cor)
                    self.validar(cel, "marca", ["X"], rotulo="marcação")
                elif t in ("date", "num"):
                    self.validar(cel, t)
                elif t == "unit":
                    self.validar(cel, "sugestao", lista="dl-unidades")
                elif t.startswith("list:"):
                    self.validar(cel, "sugestao", lista=t[5:])
                if not (j == 0 and pre):
                    self.entrada(r, a, a)
            self.r += 1
        nota = x.get("nota")
        rodape = "Para mais linhas, selecione a última linha da tabela, copie e use Inserir células copiadas."
        ws.row_dimensions[self.r].height = 24 if nota else 14
        mesclar(ws, self.r, col0, col0 + largura - 1, (nota + "  " if nota else "") + rodape, fonte(7.5, False, CINZA, True),
                Alignment(wrap_text=True, vertical="top"))
        self.r += 1

    def w_checklist(self, x, col0, largura):
        ws = self.ws
        rh = self.r
        ws.row_dimensions[rh].height = 18
        partes = [("Nº", col0, col0), ("Item de conferência", col0 + 1, col0 + 6), ("Situação", col0 + 7, col0 + 8), ("Observação", col0 + 9, col0 + 11)]
        for rot, a, b in partes:
            mesclar(ws, rh, a, b, rot, fonte(7.5, True, PET), AL_CENTRO if rot in ("Nº", "Situação") else Alignment(vertical="center"))
            caixa(ws, rh, a, rh, b, preench=fill(self.tint))
        self.r += 1
        itens = list(x["itens"]) + [None] * x["extras"]
        for i, item in enumerate(itens):
            r = self.r
            ws.row_dimensions[r].height = 28 if item and len(item) > 52 else 20
            cel = ws.cell(r, col0, i + 1)
            cel.font, cel.alignment = fonte(8, True, CINZA), AL_CENTRO
            caixa(ws, r, col0, r, col0)
            it = mesclar(ws, r, col0 + 1, col0 + 6, item, fonte(9, False, "334155"), Alignment(vertical="center", wrap_text=True))
            caixa(ws, r, col0 + 1, r, col0 + 6, preench=None if item else CAMPO_FILL)
            if not item:
                self.entrada(r, col0 + 1, col0 + 1)
            sit = mesclar(ws, r, col0 + 7, col0 + 8, None, fonte(9), AL_CENTRO)
            caixa(ws, r, col0 + 7, r, col0 + 8, preench=CAMPO_FILL)
            self.validar(sit, "opcoes", SIGLAS["tri"], rotulo="situação")
            self.entrada(r, col0 + 7, col0 + 7)
            mesclar(ws, r, col0 + 9, col0 + 11, None, fonte(9), AL_CAMPO)
            caixa(ws, r, col0 + 9, r, col0 + 11, preench=CAMPO_FILL)
            self.entrada(r, col0 + 9, col0 + 9)
            self.r += 1
        ws.row_dimensions[self.r].height = 14
        mesclar(ws, self.r, col0, col0 + largura - 1, "Escolha Conferido, Pendente ou Não aplicável em cada item. As últimas linhas ficam livres para itens próprios da sua empresa.",
                fonte(7.5, False, CINZA, True), Alignment(vertical="top", wrap_text=True))
        self.r += 1

    # -------------------------------------------------- rodapé e impressão
    def fechar(self):
        ws, f = self.ws, self.f
        ws.row_dimensions[self.r].height = 10
        self.r += 1
        ws.row_dimensions[self.r].height = 14
        mesclar(ws, self.r, C1, 8, f"{f['cod']} · {f['titulo']}", fonte(7.5, False, CINZA), Alignment(vertical="center"))
        mesclar(ws, self.r, 9, C2, f"Modelo interno de organização e registro · Versão {VERSAO}", fonte(7.5, False, CINZA), Alignment(horizontal="right", vertical="center"))
        for c in range(C1, C2 + 1):
            ws.cell(self.r, c).border = Border(top=BORDA_LEVE)
        ultima = self.r
        ws.print_area = f"B2:M{ultima}"
        ps = ws.page_setup
        ps.paperSize = ws.PAPERSIZE_A4
        ps.orientation = "portrait"
        if self.quebras:
            # o Excel ignora quebras manuais com "ajustar à página": usa escala fixa
            ws.sheet_properties.pageSetUpPr.fitToPage = False
            ps.scale = 92
            for q in self.quebras:
                ws.row_breaks.append(Break(id=q))
        else:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ps.fitToWidth = 1
            ps.fitToHeight = 1
        ws.page_margins.left = ws.page_margins.right = 0.45
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.6
        ws.page_margins.footer = 0.25
        ws.print_options.horizontalCentered = True
        ws.oddFooter.left.text = f"{f['cod']} · {f['titulo']} · Modelo interno v{VERSAO}"
        ws.oddFooter.left.size = 7
        ws.oddFooter.left.font = FONTE
        ws.oddFooter.right.text = "Página &P de &N"
        ws.oddFooter.right.size = 7
        ws.oddFooter.right.font = FONTE
        return ultima

    def gerar(self):
        self.cabecalho()
        for i, bloco in enumerate(self.f["blocos"]):
            titulo, campos = bloco[0], bloco[-1]
            opc = bloco[1] if len(bloco) == 3 else {}
            self.bloco(i, titulo, campos, quebra=opc.get("quebra", False))
        return self.fechar()


# ------------------------------------------------------------------ aba Início

def aba_inicio(wb, folhas):
    ws = wb["Início"]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PET
    larg = {"A": 1.6, "B": 7, "C": 34, "D": 29, "E": 12, "F": 14, "G": 8, "H": 1.6}
    for k, v in larg.items():
        ws.column_dimensions[k].width = v
    for r in range(1, 6):
        for c in range(2, 8):
            ws.cell(r, c).fill = fill(PET)
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 16
    mesclar(ws, 2, 2, 7, "KIT DE FORMULÁRIOS PARA A GESTÃO DE RESÍDUOS · VERSÃO " + VERSAO, fonte(8, True, "BFDCD4"), Alignment(vertical="bottom", indent=1))
    ws.row_dimensions[3].height = 40
    mesclar(ws, 3, 2, 7, "KIT PRÉ-MTR", fonte(28, True, "FFFFFF"), Alignment(vertical="center", indent=1))
    ws.row_dimensions[4].height = 22
    mesclar(ws, 4, 2, 7, "15 Formulários Prontos para Coletas e Movimentação de Resíduos", fonte(14, True, OURO), Alignment(vertical="center", indent=1))
    ws.row_dimensions[5].height = 26
    mesclar(ws, 5, 2, 7, "Modelos editáveis para organizar, conferir, registrar e acompanhar informações da rotina ambiental.  "
            "15 formulários | 3 categorias | Prontos para preencher", fonte(10, False, "D6E8E3"), Alignment(vertical="center", indent=1, wrap_text=True))
    r = 7
    ws.row_dimensions[r].height = 30
    mesclar(ws, r, 2, 7, "ⓘ  " + AVISO, fonte(9, False, "5B4220"), Alignment(vertical="center", wrap_text=True, indent=1))
    caixa(ws, r, 2, r, 7, borda=Side(style="thin", color="F0DDB8"), preench=fill("FFF8EA"))
    r = 9
    mesclar(ws, r, 2, 7, "COMO UTILIZAR", fonte(9, True, VERDE))
    passos = [
        ("1", "Escolha o formulário", "Clique em Abrir no índice abaixo ou use as abas F01 a F15."),
        ("2", "Preencha as informações", "Digite nas células em tom claro com borda. Nas células com seta, escolha na lista."),
        ("3", "Confira com os responsáveis", "Valide os dados com o setor, o transportador ou o destinador antes de usar."),
        ("4", "Salve ou imprima", "Salve uma cópia da planilha ou use Arquivo > Imprimir (cada aba já está ajustada para A4)."),
    ]
    for n, t, d in passos:
        r += 1
        ws.row_dimensions[r].height = 20
        c = ws.cell(r, 2, n)
        c.font, c.alignment, c.fill = fonte(11, True, OURO), AL_CENTRO, fill(PET)
        mesclar(ws, r, 3, 3, t, fonte(10, True, PET), Alignment(vertical="center", indent=1))
        mesclar(ws, r, 4, 7, d, fonte(9, False, "334155"), Alignment(vertical="center", wrap_text=True))
    r += 2
    mesclar(ws, r, 2, 7, "LEGENDA DAS CÉLULAS", fonte(9, True, VERDE))
    r += 1
    ws.row_dimensions[r].height = 20
    ex = ws.cell(r, 2, "")
    caixa(ws, r, 2, r, 2, preench=CAMPO_FILL)
    mesclar(ws, r, 3, 7, "Célula clara com borda: campo para preencher. As demais células são títulos e orientações.", fonte(9, False, "334155"), Alignment(vertical="center", indent=1))
    r += 1
    ws.row_dimensions[r].height = 20
    ex = ws.cell(r, 2, "X")
    ex.font, ex.alignment = fonte(11, True, VERDE), AL_CENTRO
    caixa(ws, r, 2, r, 2, preench=CAMPO_FILL)
    mesclar(ws, r, 3, 7, "Marcação: escolha X na lista. Situações (Conferido, Pendente, Não aplicável) e outras opções também abrem em lista.", fonte(9, False, "334155"), Alignment(vertical="center", indent=1, wrap_text=True))
    r += 1
    ws.row_dimensions[r].height = 20
    mesclar(ws, r, 2, 7, "Exemplo de formato: data 23/10/2026 · horário 14:30 · quantidade 1250,5 com a unidade kg no campo ao lado.",
            fonte(9, False, "334155", True), Alignment(vertical="center", indent=1))
    # índice
    r += 2
    mesclar(ws, r, 2, 7, "ÍNDICE DOS FORMULÁRIOS", fonte(9, True, VERDE))
    r += 1
    cab = ["Código", "Formulário", "Para que serve", "Registro nº", "Andamento", "Abrir"]
    ws.row_dimensions[r].height = 20
    for j, t in enumerate(cab, 2):
        c = ws.cell(r, j, t)
        c.font, c.fill = fonte(8.5, True, "FFFFFF"), fill(PET)
        c.alignment = Alignment(vertical="center", horizontal="center" if j in (2, 6, 7) else "left", indent=0 if j in (2, 6, 7) else 1)
    for cat in CATEGORIAS:
        r += 1
        ws.row_dimensions[r].height = 20
        mesclar(ws, r, 2, 7, f"Categoria {cat['num']} · {cat['nome']}", fonte(9, True, hexcor(cat["cor"])), Alignment(vertical="center", indent=1))
        for cc in range(2, 8):
            ws.cell(r, cc).fill = fill(cat["tint"])
        for f in FORMULARIOS:
            if f["cat"] != cat["num"]:
                continue
            fo = folhas[f["cod"]]
            r += 1
            ws.row_dimensions[r].height = 30
            c = ws.cell(r, 2, f["cod"])
            c.font, c.alignment = fonte(9, True, "FFFFFF"), AL_CENTRO
            c.fill = fill(cat["cor"])
            c = ws.cell(r, 3, f["titulo"])
            c.font, c.alignment = fonte(9.5, True, GRAF), Alignment(vertical="center", wrap_text=True, indent=1)
            c = ws.cell(r, 4, f["curta"])
            c.font, c.alignment = fonte(8, False, "475569"), Alignment(vertical="center", wrap_text=True)
            c = ws.cell(r, 5, f"=IF('{f['cod']}'!L3=\"\",\"\",'{f['cod']}'!L3)")
            c.font, c.alignment = fonte(9), Alignment(vertical="center", indent=1)
            faixas = ",".join(f"'{f['cod']}'!{x}" for x in fo.entradas)
            c = ws.cell(r, 6, f'=IF(COUNTA({faixas})=0,"Em branco",COUNTA({faixas})&" de {fo.n_entradas} campos")')
            c.font, c.alignment = fonte(8.5, False, "334155"), AL_CENTRO
            c = ws.cell(r, 7, "Abrir ›")
            c.hyperlink = f"#'{f['cod']}'!A1"
            c.font, c.alignment = fonte(9, True, VERDE, False), AL_CENTRO
            for cc in range(2, 8):
                ws.cell(r, cc).border = Border(bottom=BORDA_LEVE)
    r += 1
    c = ws.cell(r, 2, "Orientações finais de utilização ›")
    ws.merge_cells(start_row=r + 0, start_column=2, end_row=r, end_column=4)
    c.hyperlink = "#'Orientações'!A1"
    c.font = fonte(9, True, VERDE)
    ws.row_dimensions[r].height = 22
    c.alignment = Alignment(vertical="center")
    r += 1
    ws.row_dimensions[r].height = 26
    mesclar(ws, r, 2, 7, "Andamento conta as células preenchidas de cada formulário, incluindo o Registro nº. Os itens sugeridos já escritos nas tabelas não entram na contagem.",
            fonte(7.5, False, CINZA, True), Alignment(vertical="center", wrap_text=True))
    ws.print_area = f"B2:G{r}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = ws.page_margins.right = 0.45
    ws.oddFooter.left.text = "Kit Pré-MTR · Índice"
    ws.oddFooter.right.text = "Página &P de &N"
    for h in (ws.oddFooter.left, ws.oddFooter.right):
        h.size, h.font = 7, FONTE
    ws.freeze_panes = None


def aba_orientacoes(wb):
    ws = wb.create_sheet("Orientações")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PET
    for k, v in {"A": 1.6, "B": 6, "C": 32, "D": 62, "E": 1.6}.items():
        ws.column_dimensions[k].width = v
    cel = mesclar(ws, 1, 3, 4, "‹ Voltar ao índice", fonte(8, True, VERDE), Alignment(horizontal="right"))
    cel.hyperlink = "#'Início'!A1"
    ws.row_dimensions[2].height = 30
    mesclar(ws, 2, 2, 4, "Orientações finais de utilização", fonte(18, True, PET), Alignment(vertical="center"))
    itens = [
        ("Confira os dados antes do uso", "Compare as informações registradas com os documentos e com os responsáveis por cada etapa antes de lançar dados em qualquer sistema."),
        ("Adapte os campos à sua empresa", "Use os campos de observação e as linhas adicionais para incluir o que for próprio dos seus procedimentos internos."),
        ("Guarde cópias de forma organizada", "Salve as cópias com nomes claros, por exemplo F08_2026-10_Unidade-A, conforme a organização interna."),
        ("Consulte os procedimentos oficiais", "Quando houver dúvida sobre documentos exigidos, prazos ou classificação, siga os procedimentos oficiais e o sistema aplicável."),
    ]
    r = 3
    for i, (t, d) in enumerate(itens, 1):
        r += 1
        ws.row_dimensions[r].height = 34
        c = ws.cell(r, 2, i)
        c.font, c.alignment, c.fill = fonte(11, True, "FFFFFF"), AL_CENTRO, fill(VERDE)
        mesclar(ws, r, 3, 3, t, fonte(10, True, PET), Alignment(vertical="center", wrap_text=True, indent=1))
        mesclar(ws, r, 4, 4, d, fonte(9, False, "334155"), Alignment(vertical="center", wrap_text=True))
        for cc in (2, 3, 4):
            ws.cell(r, cc).border = Border(bottom=BORDA_LEVE)
    r += 2
    mesclar(ws, r, 2, 4, "Em que momento usar cada formulário", fonte(12, True, PET))
    fases = [
        ("Antes da coleta", ["F01", "F02", "F03", "F04", "F05", "F06", "F07"]),
        ("Preparação da saída", ["F08", "F09", "F10"]),
        ("Depois da saída", ["F11", "F12", "F13", "F14", "F15"]),
    ]
    por = {f["cod"]: f for f in FORMULARIOS}
    for nome, cods in fases:
        r += 1
        ws.row_dimensions[r].height = 20
        mesclar(ws, r, 2, 4, nome, fonte(10, True, "FFFFFF"), Alignment(vertical="center", indent=1))
        for cc in (2, 3, 4):
            ws.cell(r, cc).fill = fill(PET)
        for cod in cods:
            r += 1
            ws.row_dimensions[r].height = 18
            c = ws.cell(r, 2, cod)
            c.font, c.alignment = fonte(9, True, hexcor(CAT[por[cod]["cat"]]["cor"])), AL_CENTRO
            c.hyperlink = f"#'{cod}'!A1"
            mesclar(ws, r, 3, 4, por[cod]["titulo"], fonte(9, False, "334155"), Alignment(vertical="center", indent=1))
            for cc in (2, 3, 4):
                ws.cell(r, cc).border = Border(bottom=BORDA_LEVE)
    r += 2
    ws.row_dimensions[r].height = 46
    mesclar(ws, r, 2, 4, "O Kit Pré-MTR foi pensado para ser usado junto com os 70 Fluxos Visuais de MTR, CDF e Controle de Resíduos. "
            "Os fluxos ajudam a entender cada etapa, a planilha de controle acompanha os registros de MTR e CDF, e estes formulários "
            "organizam as informações de cada coleta e de cada situação específica.", fonte(9, False, "FFFFFF"),
            Alignment(vertical="center", wrap_text=True, indent=1))
    for cc in (2, 3, 4):
        ws.cell(r, cc).fill = fill(PET)
    ws.print_area = f"B2:D{r}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def main():
    wb = Workbook()
    wb.active.title = "Início"
    listas = Listas(wb)
    folhas = {}
    for f in FORMULARIOS:
        fo = Folha(wb, f, listas, None)
        fo.gerar()
        folhas[f["cod"]] = fo
    aba_orientacoes(wb)
    aba_inicio(wb, folhas)
    # Listas vai para o fim e fica oculta
    wb.move_sheet("Listas", offset=len(wb.sheetnames))
    wb["Listas"].sheet_state = "hidden"
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cel in row:
                if cel.font and cel.font.name != FONTE:
                    cel.font = Font(name=FONTE, size=cel.font.size or 10, bold=cel.font.bold, italic=cel.font.italic, color=cel.font.color)
    wb.active = 0
    wb.save(SAIDA)
    print(f"Planilha gerada: {SAIDA} ({SAIDA.stat().st_size // 1024} KB, {len(wb.sheetnames)} abas)")


if __name__ == "__main__":
    main()
